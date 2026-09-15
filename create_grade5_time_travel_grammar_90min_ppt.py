"""Grade 5 grammar lesson - 90 minutes, 40 slides, no speaker notes.

"The Time Travel Grammar Mission" - the student is a Grammar Time Traveler who
repairs a time machine by solving challenges in Present, Past and Future Simple.
Each explanation is followed immediately by a game, a speaking turn or a challenge.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x14, 0x1E, 0x46)
SLATE = RGBColor(0x2E, 0x3A, 0x5C)
PRESENT = RGBColor(0x2E, 0x9E, 0x5B)
PAST = RGBColor(0x2E, 0x6F, 0xD9)
FUTURE = RGBColor(0x8B, 0x4F, 0xC8)
TEAL = RGBColor(0x00, 0x9B, 0x95)
AMBER = RGBColor(0xEE, 0xA2, 0x1A)
GOLD = RGBColor(0xFF, 0xC5, 0x30)
CORAL = RGBColor(0xE5, 0x55, 0x4C)
CREAM = RGBColor(0xFC, 0xFC, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x2A, 0x36)
SOFT = RGBColor(0x69, 0x76, 0x88)
L_PRESENT = RGBColor(0xE3, 0xF4, 0xE9)
L_PAST = RGBColor(0xE4, 0xEE, 0xFC)
L_FUTURE = RGBColor(0xF0, 0xE9, 0xFA)
L_TEAL = RGBColor(0xDD, 0xF4, 0xF3)
L_AMBER = RGBColor(0xFF, 0xF3, 0xD8)
L_CORAL = RGBColor(0xFD, 0xEA, 0xE8)
L_GREY = RGBColor(0xF0, 0xF3, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TOTAL = 40
_counter = {"n": 0}

TENSES = [("PRESENT", "🟢", PRESENT, L_PRESENT), ("PAST", "🔵", PAST, L_PAST),
          ("FUTURE", "🟣", FUTURE, L_FUTURE)]

# ------------------------------------------------------------------ content

ICEBREAKERS = [
    ("🏠", "What do you usually do after school?", "NOW", PRESENT),
    ("⏪", "What did you do yesterday?", "BEFORE", PAST),
    ("⏩", "What will you do this weekend?", "LATER", FUTURE),
]

NOW_BEFORE_LATER = [
    ("BEFORE", "PAST", "🔵", "It already happened.", "I played soccer yesterday.", PAST, L_PAST),
    ("NOW", "PRESENT", "🟢", "It happens again and again.",
     "I play soccer every Saturday.", PRESENT, L_PRESENT),
    ("LATER", "FUTURE", "🟣", "It has not happened yet.",
     "I will play soccer tomorrow.", FUTURE, L_FUTURE),
]

TIME_DETECTIVE_A = [
    ("I visited my grandma yesterday.", "yesterday"),
    ("I will visit my grandma tomorrow.", "tomorrow"),
    ("I visit my grandma every Sunday.", "every Sunday"),
    ("She plays soccer every Saturday.", "every Saturday"),
]
TIME_DETECTIVE_B = [
    ("We watched a movie last night.", "last night"),
    ("They will go to the park next week.", "next week"),
    ("He reads books every night.", "every night"),
    ("I ate pizza two days ago.", "two days ago"),
]

PRESENT_USES = [
    ("🔁", "Habits", "I brush my teeth every morning."),
    ("📅", "Routines", "School starts at 8 o'clock."),
    ("🌍", "Facts", "The sun rises in the east."),
    ("⭐", "Repeated actions", "She plays soccer on Saturdays."),
]

PRESENT_EXAMPLES = [
    ("I", "play", "I play soccer."), ("You", "read", "You read books."),
    ("We", "go", "We go to school."), ("She", "plays", "She plays soccer."),
    ("He", "reads", "He reads books."), ("It", "runs", "The dog runs fast."),
]

ROUTINE_PICS = [
    ("🌅", "wake up", "She wakes up early."), ("🍳", "eat breakfast", "She eats breakfast."),
    ("🏫", "go to school", "She goes to school."), ("⚽", "play soccer", "She plays soccer."),
    ("📚", "read", "She reads books."), ("😴", "sleep", "She sleeps at 9."),
]

S_RULE = [
    ("I / You / We / They", "NO -s", ["I play.", "They read.", "We go."], PRESENT, L_PRESENT),
    ("He / She / It", "ADD -s", ["He plays.", "She reads.", "The dog runs."], CORAL, L_CORAL),
]
ES_RULE = [("watch", "watches"), ("go", "goes"), ("fix", "fixes"), ("wash", "washes")]

MISSING_S_A = [
    ("She ___ to school every day.", ["go", "goes"]),
    ("He ___ football after class.", ["play", "plays"]),
    ("The dog ___ very fast.", ["run", "runs"]),
    ("I ___ books every night.", ["read", "reads"]),
]
MISSING_S_B = [
    ("She ___ TV after dinner.", ["watch", "watches"]),
    ("They ___ soccer on Sunday.", ["play", "plays"]),
    ("My mom ___ to work at 8.", ["go", "goes"]),
    ("We ___ our homework every day.", ["do", "does"]),
]

ACTION_GAME = ["walk", "dance", "read", "run", "eat"]
ACTION_STEPS = [
    ("🟢 PRESENT", "Do the action now!", "Walk in place.", PRESENT, L_PRESENT),
    ("🔵 PAST", "Say it happened before.", "\"I walked yesterday.\"", PAST, L_PAST),
    ("🟣 FUTURE", "Say it will happen later.", "\"I will walk tomorrow.\"", FUTURE, L_FUTURE),
]

PAST_TIME_WORDS = ["yesterday", "last night", "last week", "last Sunday", "two days ago"]
PAST_EXAMPLES = [
    ("I played soccer yesterday.", "played"),
    ("She watched a movie last night.", "watched"),
    ("We visited the zoo last Sunday.", "visited"),
]

REGULAR_VERBS = [("play", "played"), ("walk", "walked"), ("watch", "watched"),
                 ("visit", "visited"), ("cook", "cooked"), ("help", "helped")]
IRREGULAR_VERBS = [("go", "went"), ("eat", "ate"), ("see", "saw"),
                   ("come", "came"), ("have", "had"), ("run", "ran")]

YESTERDAY_A = [
    ("📚", "What did Mia do yesterday?",
     ["She reads a book.", "She read a book.", "She will read a book."]),
    ("⚽", "What did Tom do yesterday?",
     ["He played soccer.", "He plays soccer.", "He will play soccer."]),
    ("🎬", "What did they do last night?",
     ["They watch a movie.", "They will watch a movie.", "They watched a movie."]),
]
YESTERDAY_B = [
    ("🦁", "What did the class do last Sunday?",
     ["They visited the zoo.", "They visit the zoo.", "They will visit the zoo."]),
    ("🍕", "What did Sam do two days ago?",
     ["He eats pizza.", "He ate pizza.", "He will eat pizza."]),
    ("🏫", "What did Ana do yesterday?",
     ["She goes to school.", "She went to school.", "She will go to school."]),
]

FUTURE_FORMULA = [("I", "will", "play"), ("She", "will", "read"), ("They", "will", "go")]
FUTURE_TIME_WORDS = ["tomorrow", "next week", "next month", "later", "soon"]
FUTURE_EXAMPLES = ["I will visit my friend.", "She will play soccer.",
                   "We will watch a movie.", "They will go to the park."]

PREDICTOR = [
    ("🎒", "Tomorrow is Saturday.", "I will ______."),
    ("🎂", "Next week is your birthday.", "I will ______."),
    ("🌧️", "It will rain later today.", "I will ______."),
]

FUTURE_CHALLENGE = [
    ("Tomorrow, I ___ visit the park.", ["will", "did", "am"]),
    ("Next week we ___ go to the zoo.", ["will", "went", "going"]),
    ("She will ___ a movie tonight.", ["watches", "watch", "watched"]),
    ("They will ___ to school by bus.", ["go", "goes", "went"]),
    ("I will ___ my homework later.", ["did", "do", "does"]),
]

TRANSFORM_DEMO = ("I play soccer.", "I played soccer.", "I will play soccer.")
TRANSFORMS = [
    ("She reads a book.", "She read a book.", "She will read a book."),
    ("We watch TV.", "We watched TV.", "We will watch TV."),
    ("He goes to the park.", "He went to the park.", "He will go to the park."),
    ("They visit grandma.", "They visited grandma.", "They will visit grandma."),
    ("I eat pizza.", "I ate pizza.", "I will eat pizza."),
]

MISTAKES_A = [
    ("She go to school every day.", "She goes to school every day.", "He / She / It needs -s"),
    ("Yesterday I go to the park.", "Yesterday I went to the park.", "\"Yesterday\" = past"),
    ("Tomorrow I went to the zoo.", "Tomorrow I will go to the zoo.", "\"Tomorrow\" = future"),
]
MISTAKES_B = [
    ("I will goes home now.", "I will go home now.", "After WILL use the base verb"),
    ("He watch TV every night.", "He watches TV every night.", "watch → watches"),
    ("Last week we visit the museum.", "Last week we visited the museum.",
     "\"Last week\" = past"),
]

VERB_TABLE = [
    ("⚽", "play", "play / plays", "played", "will play"),
    ("📺", "watch", "watch / watches", "watched", "will watch"),
    ("📚", "read", "read / reads", "read", "will read"),
    ("🏫", "go", "go / goes", "went", "will go"),
    ("🍕", "eat", "eat / eats", "ate", "will eat"),
    ("👵", "visit", "visit / visits", "visited", "will visit"),
]

STORY_P1 = [
    ("Maya is ten years old.", PRESENT),
    ("Every Saturday, she plays soccer in the park.", PRESENT),
    ("She wakes up early and eats breakfast with her family.", PRESENT),
    ("Her best friend Zoe always comes with her.", PRESENT),
    ("Maya loves soccer because it makes her happy.", PRESENT),
    ("Yesterday was different.", PAST),
    ("Maya played a big game against another school.", PAST),
    ("She ran fast and scored two goals.", PAST),
    ("Her team won! After the game, Zoe and Maya ate pizza together.", PAST),
]
STORY_P2 = [
    ("Tomorrow, Maya will visit her grandmother.", FUTURE),
    ("They will bake cookies and watch a movie.", FUTURE),
    ("Maya will tell her all about the big game.", FUTURE),
    ("She is very excited.", PRESENT),
    ("Next week, Maya will play another match.", FUTURE),
    ("She will practice every day.", FUTURE),
    ("Her coach says she will become a great player.", FUTURE),
    ("Maya smiles. Every day brings something new.", PRESENT),
]

HUNT_SENTENCES = [
    "Every Saturday, she plays soccer.",
    "Maya played a big game.",
    "Tomorrow, Maya will visit her grandmother.",
    "She wakes up early.",
    "She ran fast and scored two goals.",
    "She will practice every day.",
]

STORY_QUESTIONS = [
    ("What does Maya do every Saturday?", "🟢 Present"),
    ("What did Maya do yesterday?", "🔵 Past"),
    ("What will Maya do tomorrow?", "🟣 Future"),
    ("Which tense is \"She will practice every day.\"?", "🟣 Future"),
]

ROLL_VERBS = [("1", "play", "⚽"), ("2", "watch", "📺"), ("3", "go", "🏫"),
              ("4", "eat", "🍕"), ("5", "visit", "👵"), ("6", "read", "📚")]

TALK_PROMPTS = [
    ("🟢 PRESENT", "Every day I ______.", PRESENT, L_PRESENT),
    ("🔵 PAST", "Yesterday I ______.", PAST, L_PAST),
    ("🟣 FUTURE", "Tomorrow I will ______.", FUTURE, L_FUTURE),
]

CHAMPION = [
    ("Every day I ___ to school.", ["go", "went", "will go"]),
    ("Yesterday I ___ a movie.", ["watch", "watched", "will watch"]),
    ("Tomorrow I ___ my friend.", ["visit", "visited", "will visit"]),
    ("She ___ soccer every Saturday.", ["play", "plays", "played"]),
    ("Last week we ___ the zoo.", ["visit", "visited", "will visit"]),
]

HINTS = [
    ("Hint 1", "Look at the TIME WORD.", "yesterday · every day · tomorrow", AMBER),
    ("Hint 2", "Now, before, or later?", "When does it happen?", TEAL),
    ("Hint 3", "Look at the VERB.", "go · went · will go", PRESENT),
    ("Hint 4", "Let's say it together.", "Then she tries alone.", FUTURE),
]

SUPPORT_LEVELS = [
    ("FULL SUPPORT", "Teacher models the whole sentence.", "\"I played soccer yesterday.\""),
    ("PARTIAL SUPPORT", "Give two choices only.", "go  or  went ?"),
    ("LESS SUPPORT", "Give the sentence with a blank.", "Yesterday I ___ soccer."),
    ("INDEPENDENT", "Student builds her own sentence.", "Your turn — any verb!"),
]

RUBRIC_SKILLS = ["Identifies tense", "Present Simple", "Past Simple", "Future Simple",
                 "Uses correct verbs", "Corrects mistakes", "Creates sentences",
                 "Speaking confidence"]
RUBRIC_NOTES = ["Grammar areas needing more practice", "Words / verbs she struggled with",
                "Recommended next lesson"]

# ------------------------------------------------------------------ helpers


def set_run(run, size=20, bold=False, color=DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    tree = slide.shapes._spTree
    el = s._element
    tree.remove(el)
    tree.insert(2, el)


def add_rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_round(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_oval(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def tb(slide, l, t, w, h, text, size=20, bold=False, color=DARK,
       align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size, bold, color, font)
    return box


def bullets(slide, l, t, w, h, items, size=12, color=DARK, sp=5):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(sp)
        r = p.add_run()
        # An item starting with spaces is a wrapped continuation, not a new bullet.
        r.text = ("      " + item.strip()) if item.startswith(" ") else ("•  " + item)
        set_run(r, size, False, color)
    return box


def fade(slide):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", "med")
    tr.set("advClick", "1")
    etree.SubElement(tr, qn("p:fade"))


def footer(slide, n, timing="", step=""):
    add_rect(slide, Inches(0), Inches(7.02), prs.slide_width, Inches(0.08),
             RGBColor(0xE4, 0xEA, 0xF0))
    add_rect(slide, Inches(0), Inches(7.02), int(prs.slide_width * n / TOTAL), Inches(0.08), GOLD)
    add_rect(slide, Inches(0), Inches(7.10), prs.slide_width, Inches(0.40), NAVY)
    msg = "🕰️ The Time Travel Grammar Mission  |  Grade 5  |  90 min"
    if step:
        msg += f"  |  {step}"
    if timing:
        msg += f"  |  {timing}"
    tb(slide, Inches(0.35), Inches(7.14), Inches(11.3), Inches(0.32), msg, size=10, color=WHITE)
    tb(slide, Inches(11.85), Inches(7.14), Inches(1.2), Inches(0.32), f"{n} / {TOTAL}",
       size=10, color=WHITE, align=PP_ALIGN.RIGHT)


def chip(slide, text, color):
    add_round(slide, Inches(0.4), Inches(0.26), Inches(3.4), Inches(0.42), color)
    tb(slide, Inches(0.4), Inches(0.31), Inches(3.4), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def time_chip(slide, text):
    add_round(slide, Inches(10.3), Inches(0.26), Inches(2.65), Inches(0.42), CORAL)
    tb(slide, Inches(10.3), Inches(0.31), Inches(2.65), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def new_slide(title, tag, timing, step, accent=TEAL, bg=CREAM):
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_rect(slide, Inches(0), Inches(0), Inches(0.16), prs.slide_height, accent)
    chip(slide, tag, accent)
    if timing:
        time_chip(slide, timing)
    tb(slide, Inches(0.4), Inches(0.8), Inches(12.4), Inches(0.5), title, size=28, bold=True,
       color=NAVY, font="Georgia")
    footer(slide, n, timing, step)
    fade(slide)
    return slide, n


def subtitle(slide, text, color=CORAL, top=1.36, size=16):
    return tb(slide, Inches(0.45), Inches(top), Inches(11.5), Inches(0.4), text,
              size=size, bold=True, color=color)


def tense_legend(slide, top=6.45):
    for i, (name, dot, color, light) in enumerate(TENSES):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(top), Inches(3.9), Inches(0.46), light)
        tb(slide, left, Inches(top + 0.06), Inches(3.9), Inches(0.36), f"{dot}  {name}",
           size=14, bold=True, color=color, align=PP_ALIGN.CENTER)


def answer_choices(slide, top, options, light, color, left_start=5.3, width=2.3, gap=2.55):
    for j, opt in enumerate(options):
        left = Inches(left_start + j * gap)
        add_round(slide, left, top, Inches(width), Inches(0.58), light)
        tb(slide, left, top + Inches(0.05), Inches(width), Inches(0.46),
           f"{chr(65 + j)}.  {opt}", size=16, bold=True, color=color, align=PP_ALIGN.CENTER)


def quiz_rows(slide, items, accent, light, start_index=1, top_start=1.9, gap=1.0,
              height=0.88, prompt_width=3.5, prompt_size=17, opt_left=5.3, opt_width=2.3,
              opt_gap=2.55, opt_size=16):
    for i, (prompt, options) in enumerate(items):
        top = Inches(top_start + i * gap)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(height), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.19), Inches(0.5), Inches(0.5), accent)
        tb(slide, Inches(0.72), top + Inches(0.25), Inches(0.5), Inches(0.4),
           str(start_index + i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.45), top + Inches(0.22), Inches(prompt_width), Inches(0.5), prompt,
           size=prompt_size, bold=True, color=NAVY)
        for j, opt in enumerate(options):
            left = Inches(opt_left + j * opt_gap)
            add_round(slide, left, top + Inches(0.15), Inches(opt_width), Inches(0.58), light)
            tb(slide, left, top + Inches(0.2), Inches(opt_width), Inches(0.46),
               f"{chr(65 + j)}.  {opt}", size=opt_size, bold=True, color=NAVY,
               align=PP_ALIGN.CENTER)


def tense_choice_rows(slide, items, start_index, top_start=1.9, gap=1.16, height=1.04):
    for i, (sentence, clue) in enumerate(items):
        top = Inches(top_start + i * gap)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(height), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.28), Inches(0.5), Inches(0.5), TEAL)
        tb(slide, Inches(0.72), top + Inches(0.34), Inches(0.5), Inches(0.4),
           str(start_index + i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.45), top + Inches(0.14), Inches(5.3), Inches(0.5), sentence,
           size=19, bold=True, color=NAVY)
        tb(slide, Inches(1.45), top + Inches(0.6), Inches(5.3), Inches(0.36),
           f"⏱ time word: {clue}", size=12, color=SOFT)
        for j, (name, dot, color, light) in enumerate(TENSES):
            left = Inches(7.0 + j * 1.98)
            add_round(slide, left, top + Inches(0.22), Inches(1.8), Inches(0.6), light)
            tb(slide, left, top + Inches(0.29), Inches(1.8), Inches(0.44), f"{dot} {name}",
               size=13, bold=True, color=color, align=PP_ALIGN.CENTER)


# ------------------------------------------------------------------ slides


def s01_title():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.65, 0.8, PRESENT), (11.95, 0.85, PAST), (0.85, 5.8, FUTURE),
                    (11.9, 5.75, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(1.15), Inches(12), Inches(1.05), "🕰️", size=58,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(2.18), Inches(12), Inches(0.85),
       "The Time Travel Grammar Mission", size=40, bold=True, color=WHITE,
       align=PP_ALIGN.CENTER, font="Georgia")
    tb(slide, Inches(0.7), Inches(3.08), Inches(12), Inches(0.5),
       "Welcome, Grammar Time Traveler!  •  Grade 5  •  90 Minutes",
       size=18, color=RGBColor(0xC7, 0xD6, 0xF2), align=PP_ALIGN.CENTER)
    add_round(slide, Inches(3.15), Inches(3.85), Inches(7.0), Inches(1.25), TEAL)
    tb(slide, Inches(3.35), Inches(4.1), Inches(6.6), Inches(0.8),
       "The time machine is broken.\nSolve the grammar challenges to repair it!",
       size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, (name, dot, color, _light) in enumerate(TENSES):
        left = Inches(2.6 + i * 2.9)
        add_round(slide, left, Inches(5.5), Inches(2.6), Inches(0.72), color)
        tb(slide, left, Inches(5.66), Inches(2.6), Inches(0.45), f"{dot}  {name}", size=17,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(slide, n, "0–7 min", "Welcome")
    fade(slide)


def s02_time_machine():
    slide, n = new_slide("🕰️ Meet the Time Machine", "MEET", "0–7 min", "Time machine", TEAL)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(5.4), Inches(4.95), L_TEAL)
    tb(slide, Inches(0.5), Inches(2.25), Inches(5.4), Inches(1.9), "🕰️", size=96,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(4.35), Inches(5.4), Inches(0.7), "CHRONO-5", size=36,
       bold=True, color=TEAL, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(0.8), Inches(5.15), Inches(4.8), Inches(0.6),
       "Your time machine for today", size=16, color=SOFT, align=PP_ALIGN.CENTER)
    facts = [("⚠️", "Chrono-5 is broken. It cannot travel."),
             ("🔧", "Each grammar challenge repairs one part."),
             ("🎯", "Three zones to unlock: Present, Past, Future."),
             ("🏆", "Finish all three and you earn your badge.")]
    for i, (icon, text) in enumerate(facts):
        top = Inches(1.6 + i * 1.22)
        add_round(slide, Inches(6.25), top, Inches(6.6), Inches(1.05), WHITE)
        add_oval(slide, Inches(6.5), top + Inches(0.24), Inches(0.58), Inches(0.58), L_AMBER)
        tb(slide, Inches(6.5), top + Inches(0.3), Inches(0.58), Inches(0.45), icon, size=16,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(7.3), top + Inches(0.3), Inches(5.3), Inches(0.5), text, size=16,
           bold=True, color=NAVY)


def s03_mission():
    slide, n = new_slide("🗺️ Today's Mission — 3 Time Zones", "MISSION", "0–7 min",
                         "Mission", AMBER)
    subtitle(slide, "Repair one zone at a time. Each zone has lessons, games and a challenge.")
    zones = [
        ("🟢", "ZONE 1 — PRESENT", "What happens again and again", "22–42 min",
         PRESENT, L_PRESENT),
        ("🔵", "ZONE 2 — PAST", "What already happened", "42–58 min", PAST, L_PAST),
        ("🟣", "ZONE 3 — FUTURE", "What has not happened yet", "58–70 min",
         FUTURE, L_FUTURE),
    ]
    for i, (dot, title, blurb, when, color, light) in enumerate(zones):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(1.9), Inches(3.9), Inches(3.1), light)
        tb(slide, left, Inches(2.15), Inches(3.9), Inches(0.75), dot, size=34,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), Inches(2.95), Inches(3.6), Inches(0.55), title,
           size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.25), Inches(3.55), Inches(3.4), Inches(0.65), blurb,
           size=14, color=DARK, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(1.15), Inches(4.3), Inches(1.6), Inches(0.45), WHITE)
        tb(slide, left + Inches(1.15), Inches(4.36), Inches(1.6), Inches(0.36), when, size=12,
           bold=True, color=SOFT, align=PP_ALIGN.CENTER)
    extras = [("🔄", "Tense Transformation", "70–77 min"), ("🕵️", "Grammar Detective", "77–82 min"),
              ("📖", "Maya's Story", "82–87 min"), ("🏆", "Tense Champion", "87–90 min")]
    for i, (icon, label, when) in enumerate(extras):
        left = Inches(0.5 + i * 3.16)
        add_round(slide, left, Inches(5.2), Inches(2.96), Inches(1.15), WHITE)
        tb(slide, left + Inches(0.2), Inches(5.35), Inches(0.6), Inches(0.5), icon, size=18,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.85), Inches(5.35), Inches(2.0), Inches(0.42), label,
           size=13, bold=True, color=NAVY)
        tb(slide, left + Inches(0.85), Inches(5.76), Inches(2.0), Inches(0.36), when, size=11,
           color=SOFT)


def s04_icebreaker():
    slide, n = new_slide("💬 Quick Talk — Tell Me About You", "SPEAKING", "0–7 min",
                         "Icebreaker", CORAL)
    subtitle(slide, "No grammar rules yet. Just answer in your own words!")
    for i, (icon, question, tag, color) in enumerate(ICEBREAKERS):
        top = Inches(2.0 + i * 1.45)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.25), WHITE)
        add_oval(slide, Inches(0.8), top + Inches(0.32), Inches(0.62), Inches(0.62), L_AMBER)
        tb(slide, Inches(0.8), top + Inches(0.4), Inches(0.62), Inches(0.46), icon, size=17,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.7), top + Inches(0.36), Inches(8.2), Inches(0.6), question,
           size=22, bold=True, color=NAVY)
        add_round(slide, Inches(10.4), top + Inches(0.35), Inches(2.2), Inches(0.58), color)
        tb(slide, Inches(10.4), top + Inches(0.42), Inches(2.2), Inches(0.44), tag, size=15,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.35), Inches(12.35), Inches(0.55), L_TEAL)
    tb(slide, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.4),
       "🎤 Teacher answers first — it makes speaking feel safe. Do not correct yet.",
       size=15, bold=True, color=NAVY)


def s05_now_before_later():
    slide, n = new_slide("⏳ Now, Before, Later", "CONCEPT", "7–15 min", "Now/Before/Later",
                         TEAL)
    subtitle(slide, "You already used all three tenses when you answered. Look!")
    for i, (when, tense, dot, blurb, example, color, light) in enumerate(NOW_BEFORE_LATER):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(1.9), Inches(3.9), Inches(4.35), light)
        tb(slide, left, Inches(2.1), Inches(3.9), Inches(0.6), dot, size=28,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), Inches(2.75), Inches(3.6), Inches(0.5), when, size=16,
           bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), Inches(3.2), Inches(3.6), Inches(0.62), tense, size=26,
           bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left + Inches(0.25), Inches(3.9), Inches(3.4), Inches(0.5), blurb, size=14,
           color=DARK, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.2), Inches(4.55), Inches(3.5), Inches(1.5), WHITE)
        tb(slide, left + Inches(0.35), Inches(4.8), Inches(3.2), Inches(1.0), example,
           size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def s06_what_is_tense():
    slide, n = new_slide("📚 What Is a Tense?", "CONCEPT", "7–15 min", "What is a tense", TEAL)
    add_round(slide, Inches(0.5), Inches(1.45), Inches(12.35), Inches(0.85), L_AMBER)
    tb(slide, Inches(0.8), Inches(1.62), Inches(11.7), Inches(0.55),
       "A tense tells us WHEN something happens.", size=26, bold=True, color=NAVY)
    add_rect(slide, Inches(1.2), Inches(3.08), Inches(10.9), Inches(0.09), SLATE)
    for i, (label, dot, color) in enumerate([("PAST", "🔵", PAST), ("NOW", "🟢", PRESENT),
                                             ("FUTURE", "🟣", FUTURE)]):
        cx = Inches(1.35 + i * 5.05)
        add_oval(slide, cx, Inches(2.82), Inches(0.62), Inches(0.62), color)
        tb(slide, cx - Inches(0.95), Inches(2.3), Inches(2.5), Inches(0.45),
           f"{dot} {label}", size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.9), Inches(2.76), Inches(0.6), Inches(0.5), "←", size=22, bold=True,
       color=SLATE)
    tb(slide, Inches(11.95), Inches(2.76), Inches(0.6), Inches(0.5), "→", size=22, bold=True,
       color=SLATE)
    cards = [("🔵 PAST", "I played soccer yesterday.", "played", PAST, L_PAST),
             ("🟢 PRESENT", "I play soccer every Saturday.", "play", PRESENT, L_PRESENT),
             ("🟣 FUTURE", "I will play soccer tomorrow.", "will play", FUTURE, L_FUTURE)]
    for i, (label, sentence, verb, color, light) in enumerate(cards):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(3.85), Inches(3.9), Inches(2.5), light)
        tb(slide, left + Inches(0.15), Inches(4.05), Inches(3.6), Inches(0.45), label,
           size=15, bold=True, color=color, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.25), Inches(4.6), Inches(3.4), Inches(1.0), sentence,
           size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(1.05), Inches(5.65), Inches(1.8), Inches(0.5), WHITE)
        tb(slide, left + Inches(1.05), Inches(5.73), Inches(1.8), Inches(0.38), verb, size=15,
           bold=True, color=color, align=PP_ALIGN.CENTER)


def s07_time_detective_a():
    slide, n = new_slide("🕵️ Game: Time Detective", "GAME", "15–22 min", "Time detective", AMBER)
    subtitle(slide, "Read the sentence. When did it happen? Find the time word first!")
    tense_choice_rows(slide, TIME_DETECTIVE_A, 1)


def s08_time_detective_b():
    slide, n = new_slide("🕵️ Time Detective — Round 2", "GAME", "15–22 min",
                         "Time detective", AMBER)
    subtitle(slide, "Four more. The time word is your biggest clue.")
    tense_choice_rows(slide, TIME_DETECTIVE_B, 5)


def s09_present_intro():
    slide, n = new_slide("🟢 ZONE 1 — Present Simple", "PRESENT", "22–32 min",
                         "Present simple", PRESENT)
    subtitle(slide, "We use Present Simple for things that happen again and again.")
    for i, (icon, label, example) in enumerate(PRESENT_USES):
        col, row = i % 2, i // 2
        left = Inches(0.5 + col * 6.25)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(5.95), Inches(2.2), L_PRESENT)
        add_oval(slide, left + Inches(0.3), top + Inches(0.65), Inches(0.85), Inches(0.85),
                 WHITE)
        tb(slide, left + Inches(0.3), top + Inches(0.8), Inches(0.85), Inches(0.55), icon,
           size=20, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.45), top + Inches(0.42), Inches(4.2), Inches(0.55), label,
           size=22, bold=True, color=PRESENT)
        add_round(slide, left + Inches(1.45), top + Inches(1.08), Inches(4.25), Inches(0.75),
                  WHITE)
        tb(slide, left + Inches(1.6), top + Inches(1.2), Inches(3.95), Inches(0.55), example,
           size=14, bold=True, color=NAVY)


def s10_present_examples():
    slide, n = new_slide("🟢 Present Simple — Look at the Verb", "PRESENT", "22–32 min",
                         "Present examples", PRESENT)
    subtitle(slide, "Same verb, but it changes for He / She / It. Watch the ending!")
    for i, (subject, verb, sentence) in enumerate(PRESENT_EXAMPLES):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.95 + row * 2.4)
        highlight = row == 1
        add_round(slide, left, top, Inches(3.9), Inches(2.2),
                  L_CORAL if highlight else L_PRESENT)
        add_round(slide, left + Inches(0.25), top + Inches(0.22), Inches(1.5), Inches(0.5),
                  CORAL if highlight else PRESENT)
        tb(slide, left + Inches(0.25), top + Inches(0.29), Inches(1.5), Inches(0.38), subject,
           size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.95), top + Inches(0.24), Inches(1.7), Inches(0.5), verb,
           size=22, bold=True, color=CORAL if highlight else PRESENT)
        add_round(slide, left + Inches(0.25), top + Inches(0.9), Inches(3.4), Inches(1.05),
                  WHITE)
        tb(slide, left + Inches(0.4), top + Inches(1.15), Inches(3.1), Inches(0.6), sentence,
           size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def s11_routine_detective():
    slide, n = new_slide("🎯 Game: Routine Detective", "GAME", "22–32 min", "Routine detective",
                         PRESENT)
    subtitle(slide, "What does she do EVERY DAY? Make a full sentence for each picture.")
    for i, (emoji, label, _answer) in enumerate(ROUTINE_PICS):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(3.9), Inches(2.2), L_PRESENT)
        tb(slide, left, top + Inches(0.2), Inches(3.9), Inches(0.85), emoji, size=40,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(1.1), Inches(3.6), Inches(0.5), label,
           size=19, bold=True, color=PRESENT, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.4), top + Inches(1.62), Inches(3.1), Inches(0.45),
                  WHITE)
        tb(slide, left + Inches(0.4), top + Inches(1.69), Inches(3.1), Inches(0.36),
           "She ______ .", size=14, bold=True, color=SOFT, align=PP_ALIGN.CENTER)


def s12_he_she_it():
    slide, n = new_slide("🟢 The -s Rule: He / She / It", "PRESENT", "32–38 min",
                         "He/She/It", CORAL)
    subtitle(slide, "This is the mistake almost everyone makes. Let's beat it!")
    for i, (group, rule, examples, color, light) in enumerate(S_RULE):
        left = Inches(0.5 + i * 6.25)
        add_round(slide, left, Inches(1.9), Inches(5.95), Inches(2.9), light)
        tb(slide, left + Inches(0.2), Inches(2.1), Inches(5.55), Inches(0.5), group, size=20,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(2.1), Inches(2.65), Inches(1.75), Inches(0.5), color)
        tb(slide, left + Inches(2.1), Inches(2.73), Inches(1.75), Inches(0.38), rule, size=15,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, example in enumerate(examples):
            top = Inches(3.3 + j * 0.46)
            add_round(slide, left + Inches(0.5), top, Inches(4.95), Inches(0.4), WHITE)
            tb(slide, left + Inches(0.5), top + Inches(0.04), Inches(4.95), Inches(0.32),
               example, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(5.0), Inches(12.35), Inches(1.4), L_AMBER)
    tb(slide, Inches(0.8), Inches(5.12), Inches(11.7), Inches(0.42),
       "Ends in ch, sh, s, x or o?  Add -es instead:", size=16, bold=True, color=NAVY)
    for i, (base, changed) in enumerate(ES_RULE):
        left = Inches(0.8 + i * 3.0)
        add_round(slide, left, Inches(5.62), Inches(2.8), Inches(0.6), WHITE)
        tb(slide, left, Inches(5.7), Inches(2.8), Inches(0.44), f"{base} → {changed}",
           size=16, bold=True, color=CORAL, align=PP_ALIGN.CENTER)


def s13_missing_s_a():
    slide, n = new_slide("🏃 Game: Add the Missing S", "GAME", "32–38 min", "Missing S", CORAL)
    subtitle(slide, "Who is doing the action? That tells you if you need the -s.")
    quiz_rows(slide, MISSING_S_A, CORAL, L_CORAL, 1, top_start=2.0, gap=1.2, height=1.0,
              prompt_width=4.2, prompt_size=18, opt_left=6.2, opt_width=2.9, opt_gap=3.15,
              opt_size=18)


def s14_missing_s_b():
    slide, n = new_slide("🏃 Add the Missing S — Round 2", "GAME", "32–38 min",
                         "Present challenge", CORAL)
    subtitle(slide, "Careful — two of these do NOT need the -s.")
    quiz_rows(slide, MISSING_S_B, CORAL, L_CORAL, 5, top_start=2.0, gap=1.2, height=1.0,
              prompt_width=4.2, prompt_size=18, opt_left=6.2, opt_width=2.9, opt_gap=3.15,
              opt_size=18)


def s15_brain_break():
    slide, n = new_slide("🎮 Brain Break: Tense Action Game", "BREAK", "38–42 min",
                         "Action game", GOLD)
    subtitle(slide, "I call a tense — you do the action or say the sentence!")
    for i, (label, instruction, example, color, light) in enumerate(ACTION_STEPS):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(1.95), Inches(3.9), Inches(2.85), light)
        add_round(slide, left + Inches(0.9), Inches(2.18), Inches(2.1), Inches(0.55), color)
        tb(slide, left + Inches(0.9), Inches(2.27), Inches(2.1), Inches(0.4), label, size=14,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.25), Inches(2.95), Inches(3.4), Inches(0.6), instruction,
           size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.3), Inches(3.65), Inches(3.3), Inches(0.95), WHITE)
        tb(slide, left + Inches(0.45), Inches(3.85), Inches(3.0), Inches(0.6), example,
           size=15, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(5.1), Inches(12.35), Inches(1.3), L_AMBER)
    tb(slide, Inches(0.8), Inches(5.25), Inches(11.7), Inches(0.42),
       "Do all three tenses for each verb:", size=16, bold=True, color=NAVY)
    for i, verb in enumerate(ACTION_GAME):
        left = Inches(0.8 + i * 2.4)
        add_round(slide, left, Inches(5.72), Inches(2.2), Inches(0.55), WHITE)
        tb(slide, left, Inches(5.8), Inches(2.2), Inches(0.4), verb, size=17, bold=True,
           color=NAVY, align=PP_ALIGN.CENTER)


def s16_past_intro():
    slide, n = new_slide("🔵 ZONE 2 — Past Simple", "PAST", "42–52 min", "Past simple", PAST)
    subtitle(slide, "Past Simple = the action already finished. It is over.")
    add_round(slide, Inches(0.5), Inches(1.9), Inches(12.35), Inches(1.3), L_AMBER)
    tb(slide, Inches(0.8), Inches(2.03), Inches(11.7), Inches(0.42),
       "⏱ These time words tell you it is PAST:", size=16, bold=True, color=NAVY)
    for i, word in enumerate(PAST_TIME_WORDS):
        left = Inches(0.8 + i * 2.4)
        add_round(slide, left, Inches(2.5), Inches(2.2), Inches(0.55), WHITE)
        tb(slide, left, Inches(2.58), Inches(2.2), Inches(0.4), word, size=15, bold=True,
           color=PAST, align=PP_ALIGN.CENTER)
    for i, (sentence, verb) in enumerate(PAST_EXAMPLES):
        top = Inches(3.45 + i * 1.08)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.95), L_PAST)
        tb(slide, Inches(0.9), top + Inches(0.22), Inches(8.4), Inches(0.55), sentence,
           size=22, bold=True, color=NAVY)
        add_round(slide, Inches(9.9), top + Inches(0.2), Inches(2.6), Inches(0.55), PAST)
        tb(slide, Inches(9.9), top + Inches(0.28), Inches(2.6), Inches(0.4), verb, size=16,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def s17_regular_verbs():
    slide, n = new_slide("🔵 Regular Verbs — Just Add -ed", "PAST", "42–52 min",
                         "Regular verbs", PAST)
    subtitle(slide, "These are the easy ones. The ending is always -ed.")
    for i, (base, past) in enumerate(REGULAR_VERBS):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(3.9), Inches(2.2), L_PAST)
        tb(slide, left + Inches(0.2), top + Inches(0.55), Inches(1.5), Inches(0.7), base,
           size=26, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.75), top + Inches(0.6), Inches(0.5), Inches(0.6), "→",
           size=22, bold=True, color=PAST, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(2.3), top + Inches(0.5), Inches(1.4), Inches(0.8), WHITE)
        tb(slide, left + Inches(2.3), top + Inches(0.62), Inches(1.4), Inches(0.6), past,
           size=22, bold=True, color=PAST, align=PP_ALIGN.CENTER, font="Arial Black")
        add_round(slide, left + Inches(0.3), top + Inches(1.5), Inches(3.3), Inches(0.5), WHITE)
        tb(slide, left + Inches(0.3), top + Inches(1.57), Inches(3.3), Inches(0.38),
           f"I {past} yesterday.", size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def s18_irregular_verbs():
    slide, n = new_slide("🔵 Irregular Verbs — They Change Completely", "PAST", "42–52 min",
                         "Irregular verbs", CORAL)
    subtitle(slide, "No -ed here. These six just have to be remembered.")
    for i, (base, past) in enumerate(IRREGULAR_VERBS):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(3.9), Inches(2.2), L_CORAL)
        tb(slide, left + Inches(0.2), top + Inches(0.55), Inches(1.5), Inches(0.7), base,
           size=26, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.75), top + Inches(0.6), Inches(0.5), Inches(0.6), "→",
           size=22, bold=True, color=CORAL, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(2.3), top + Inches(0.5), Inches(1.4), Inches(0.8), WHITE)
        tb(slide, left + Inches(2.3), top + Inches(0.62), Inches(1.4), Inches(0.6), past,
           size=22, bold=True, color=CORAL, align=PP_ALIGN.CENTER, font="Arial Black")
        add_round(slide, left + Inches(0.3), top + Inches(1.5), Inches(3.3), Inches(0.5), WHITE)
        tb(slide, left + Inches(0.3), top + Inches(1.57), Inches(3.3), Inches(0.38),
           f"I {past} yesterday.", size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def _yesterday_slide(items, title, timing, start_index):
    slide, n = new_slide(title, "GAME", timing, "What happened", PAST)
    subtitle(slide, "The question is about the PAST. Which answer matches?")
    for i, (emoji, question, options) in enumerate(items):
        top = Inches(1.95 + i * 1.65)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.45), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.45), Inches(0.5), Inches(0.5), PAST)
        tb(slide, Inches(0.72), top + Inches(0.51), Inches(0.5), Inches(0.4),
           str(start_index + i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.4), top + Inches(0.42), Inches(0.9), Inches(0.62), emoji, size=26,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.4), top + Inches(0.45), Inches(3.0), Inches(0.6), question,
           size=15, bold=True, color=NAVY)
        for j, opt in enumerate(options):
            left = Inches(5.6 + j * 2.45)
            add_round(slide, left, top + Inches(0.32), Inches(2.25), Inches(0.8), L_PAST)
            tb(slide, left + Inches(0.1), top + Inches(0.42), Inches(2.05), Inches(0.6),
               f"{chr(65 + j)}. {opt}", size=13, bold=True, color=NAVY,
               align=PP_ALIGN.CENTER)
    return slide, n


def s19_yesterday_a():
    _yesterday_slide(YESTERDAY_A, "🕵️ Game: What Happened Yesterday?", "52–58 min", 1)


def s20_yesterday_b():
    _yesterday_slide(YESTERDAY_B, "🕵️ Past Tense Challenge", "52–58 min", 4)


def s21_future_intro():
    slide, n = new_slide("🟣 ZONE 3 — Future Simple", "FUTURE", "58–65 min",
                         "Future simple", FUTURE)
    subtitle(slide, "One little word does all the work: WILL")
    add_round(slide, Inches(0.5), Inches(1.9), Inches(12.35), Inches(1.75), L_FUTURE)
    tb(slide, Inches(0.8), Inches(2.05), Inches(11.7), Inches(0.45),
       "THE FORMULA:   SUBJECT  +  WILL  +  BASE VERB", size=20, bold=True, color=FUTURE)
    for i, (subject, will, verb) in enumerate(FUTURE_FORMULA):
        left = Inches(0.8 + i * 4.05)
        for j, (part, color) in enumerate([(subject, SLATE), (will, FUTURE), (verb, PRESENT)]):
            box_left = left + Inches(j * 1.25)
            add_round(slide, box_left, Inches(2.65), Inches(1.05), Inches(0.72), WHITE)
            tb(slide, box_left, Inches(2.78), Inches(1.05), Inches(0.5), part, size=17,
               bold=True, color=color, align=PP_ALIGN.CENTER)
            if j < 2:
                tb(slide, box_left + Inches(1.07), Inches(2.8), Inches(0.17), Inches(0.45),
                   "+", size=15, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(3.85), Inches(6.1), Inches(2.5), L_AMBER)
    tb(slide, Inches(0.8), Inches(4.0), Inches(5.5), Inches(0.42), "⏱ Future time words:",
       size=16, bold=True, color=NAVY)
    for i, word in enumerate(FUTURE_TIME_WORDS):
        top = Inches(4.5 + i * 0.36)
        tb(slide, Inches(1.0), top, Inches(5.2), Inches(0.32), f"•  {word}", size=14,
           color=DARK)
    add_round(slide, Inches(6.75), Inches(3.85), Inches(6.1), Inches(2.5), WHITE)
    tb(slide, Inches(7.05), Inches(4.0), Inches(5.5), Inches(0.42), "Examples:", size=16,
       bold=True, color=FUTURE)
    for i, example in enumerate(FUTURE_EXAMPLES):
        top = Inches(4.5 + i * 0.44)
        tb(slide, Inches(7.25), top, Inches(5.3), Inches(0.4), f"•  {example}", size=15,
           bold=True, color=NAVY)


def s22_future_predictor():
    slide, n = new_slide("🔮 Game: Future Predictor", "GAME", "65–70 min",
                         "Future predictor", FUTURE)
    subtitle(slide, "What will YOU do? Answer in a full sentence using WILL.")
    for i, (emoji, situation, starter) in enumerate(PREDICTOR):
        top = Inches(1.95 + i * 1.45)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.25), L_FUTURE)
        tb(slide, Inches(0.85), top + Inches(0.3), Inches(0.9), Inches(0.65), emoji, size=26,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.95), top + Inches(0.34), Inches(5.2), Inches(0.6), situation,
           size=20, bold=True, color=NAVY)
        add_round(slide, Inches(7.4), top + Inches(0.3), Inches(5.2), Inches(0.65), WHITE)
        tb(slide, Inches(7.55), top + Inches(0.42), Inches(4.9), Inches(0.45), starter,
           size=18, bold=True, color=FUTURE)
    add_round(slide, Inches(0.5), Inches(6.35), Inches(12.35), Inches(0.55), L_AMBER)
    tb(slide, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.4),
       "🎤 Sentence starter if she is stuck: \"I will ___ with my ___.\"",
       size=15, bold=True, color=NAVY)


def s23_future_challenge():
    slide, n = new_slide("🟣 Future Challenge", "GAME", "65–70 min", "Future challenge", FUTURE)
    subtitle(slide, "Remember: after WILL, the verb never changes.")
    quiz_rows(slide, FUTURE_CHALLENGE, FUTURE, L_FUTURE, 1, top_start=1.9, gap=1.02,
              height=0.9, prompt_width=4.0, prompt_size=17, opt_left=5.9, opt_width=2.2,
              opt_gap=2.4, opt_size=16)
    add_round(slide, Inches(0.5), Inches(7.0), Inches(12.35), Inches(0.01), CREAM)


def s24_transformation():
    slide, n = new_slide("🔄 Tense Transformation", "CONCEPT", "70–77 min",
                         "Transformation", TEAL)
    subtitle(slide, "One sentence, three time zones. Only the verb changes!")
    present, past, future = TRANSFORM_DEMO
    rows = [("🟢 PRESENT", present, "play", PRESENT, L_PRESENT),
            ("🔵 PAST", past, "played", PAST, L_PAST),
            ("🟣 FUTURE", future, "will play", FUTURE, L_FUTURE)]
    for i, (label, sentence, verb, color, light) in enumerate(rows):
        top = Inches(1.95 + i * 1.5)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.3), light)
        add_round(slide, Inches(0.8), top + Inches(0.38), Inches(2.1), Inches(0.55), color)
        tb(slide, Inches(0.8), top + Inches(0.46), Inches(2.1), Inches(0.4), label, size=14,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(3.3), top + Inches(0.34), Inches(6.2), Inches(0.62), sentence,
           size=26, bold=True, color=NAVY)
        add_round(slide, Inches(10.1), top + Inches(0.38), Inches(2.5), Inches(0.55), WHITE)
        tb(slide, Inches(10.1), top + Inches(0.46), Inches(2.5), Inches(0.4), verb, size=16,
           bold=True, color=color, align=PP_ALIGN.CENTER)
        if i < 2:
            tb(slide, Inches(6.3), top + Inches(1.24), Inches(0.7), Inches(0.3), "⬇",
               size=13, bold=True, color=SOFT, align=PP_ALIGN.CENTER)


def s25_transform_challenge():
    slide, n = new_slide("🔄 Transformation Challenge", "GAME", "70–77 min",
                         "Transform it", TEAL)
    subtitle(slide, "I give you PRESENT. You give me PAST and FUTURE.")
    add_round(slide, Inches(0.5), Inches(1.85), Inches(12.35), Inches(0.5), NAVY)
    for j, (label, width, left) in enumerate([("🟢 PRESENT (given)", 4.6, 0.7),
                                              ("🔵 PAST (you say)", 3.6, 5.5),
                                              ("🟣 FUTURE (you say)", 3.4, 9.3)]):
        tb(slide, Inches(left), Inches(1.94), Inches(width), Inches(0.35), label, size=13,
           bold=True, color=WHITE)
    for i, (present, _past, _future) in enumerate(TRANSFORMS):
        top = Inches(2.5 + i * 0.85)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.74), WHITE)
        add_oval(slide, Inches(0.68), top + Inches(0.17), Inches(0.42), Inches(0.42), TEAL)
        tb(slide, Inches(0.68), top + Inches(0.21), Inches(0.42), Inches(0.34), str(i + 1),
           size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(1.25), top + Inches(0.1), Inches(4.05), Inches(0.54),
                  L_PRESENT)
        tb(slide, Inches(1.4), top + Inches(0.16), Inches(3.75), Inches(0.42), present,
           size=15, bold=True, color=NAVY)
        add_round(slide, Inches(5.5), top + Inches(0.1), Inches(3.6), Inches(0.54), L_PAST)
        tb(slide, Inches(5.5), top + Inches(0.16), Inches(3.6), Inches(0.42), "___________",
           size=15, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(9.3), top + Inches(0.1), Inches(3.4), Inches(0.54), L_FUTURE)
        tb(slide, Inches(9.3), top + Inches(0.16), Inches(3.4), Inches(0.42), "___________",
           size=15, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.85), Inches(12.35), Inches(0.05), CREAM)


def _mistake_slide(items, title, timing, start_index):
    slide, n = new_slide(title, "GAME", timing, "Find the mistake", CORAL)
    subtitle(slide, "Every sentence has ONE mistake. Find it and fix it out loud!")
    for i, (wrong, _right, hint) in enumerate(items):
        top = Inches(1.95 + i * 1.62)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.42), WHITE)
        add_oval(slide, Inches(0.75), top + Inches(0.45), Inches(0.52), Inches(0.52), CORAL)
        tb(slide, Inches(0.75), top + Inches(0.51), Inches(0.52), Inches(0.4),
           str(start_index + i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(1.5), top + Inches(0.22), Inches(7.6), Inches(0.68), L_CORAL)
        tb(slide, Inches(1.7), top + Inches(0.3), Inches(7.2), Inches(0.52), f"❌  {wrong}",
           size=20, bold=True, color=NAVY)
        add_round(slide, Inches(1.5), top + Inches(0.96), Inches(7.6), Inches(0.36), WHITE)
        tb(slide, Inches(1.7), top + Inches(0.98), Inches(7.2), Inches(0.32),
           f"✅  ______________________", size=14, bold=True, color=SOFT)
        add_round(slide, Inches(9.35), top + Inches(0.38), Inches(3.3), Inches(0.66), L_AMBER)
        tb(slide, Inches(9.5), top + Inches(0.5), Inches(3.0), Inches(0.45), f"💡 {hint}",
           size=12, bold=True, color=NAVY)
    return slide, n


def s26_mistakes_a():
    _mistake_slide(MISTAKES_A, "🕵️ Catch the Grammar Mistake!", "77–82 min", 1)


def s27_mistakes_b():
    _mistake_slide(MISTAKES_B, "🕵️ Fix the Sentence — Round 2", "77–82 min", 4)


def s28_verb_table():
    slide, n = new_slide("📋 One Verb, Three Tenses", "REVIEW", "77–82 min",
                         "3-tense verbs", AMBER)
    subtitle(slide, "Your cheat sheet. Same verb, three different jobs.")
    header = [("VERB", 0.5, 3.0), ("🟢 PRESENT", 3.65, 3.0), ("🔵 PAST", 6.8, 3.0),
              ("🟣 FUTURE", 9.95, 2.9)]
    for label, left, width in header:
        add_round(slide, Inches(left), Inches(1.88), Inches(width), Inches(0.5), NAVY)
        tb(slide, Inches(left), Inches(1.97), Inches(width), Inches(0.36), label, size=14,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, (emoji, base, present, past, future) in enumerate(VERB_TABLE):
        top = Inches(2.5 + i * 0.72)
        cells = [(f"{emoji}  {base}", 0.5, 3.0, L_GREY, NAVY),
                 (present, 3.65, 3.0, L_PRESENT, PRESENT),
                 (past, 6.8, 3.0, L_PAST, PAST),
                 (future, 9.95, 2.9, L_FUTURE, FUTURE)]
        for text, left, width, light, color in cells:
            add_round(slide, Inches(left), top, Inches(width), Inches(0.62), light)
            tb(slide, Inches(left), top + Inches(0.1), Inches(width), Inches(0.44), text,
               size=16, bold=True, color=color, align=PP_ALIGN.CENTER)


def s29_story_intro():
    slide, n = new_slide("📖 Story Mission: Maya's Three-Day Adventure", "STORY", "82–87 min",
                         "Story intro", PRESENT)
    subtitle(slide, "All three tenses are hiding in this story. Your job is to spot them.")
    add_round(slide, Inches(0.5), Inches(1.9), Inches(5.6), Inches(4.4), L_AMBER)
    tb(slide, Inches(0.5), Inches(2.5), Inches(5.6), Inches(1.5), "⚽", size=80,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(4.2), Inches(5.6), Inches(0.7), "MAYA", size=34, bold=True,
       color=AMBER, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(0.8), Inches(4.95), Inches(5.0), Inches(0.9),
       "She is 10. She loves soccer.", size=16, color=SOFT, align=PP_ALIGN.CENTER)
    jobs = [("🟢", "Find the PRESENT sentences", "What she does every week", PRESENT,
             L_PRESENT),
            ("🔵", "Find the PAST sentences", "What happened yesterday", PAST, L_PAST),
            ("🟣", "Find the FUTURE sentences", "What she will do next", FUTURE, L_FUTURE)]
    for i, (dot, title, blurb, color, light) in enumerate(jobs):
        top = Inches(1.9 + i * 1.5)
        add_round(slide, Inches(6.45), top, Inches(6.4), Inches(1.3), light)
        tb(slide, Inches(6.75), top + Inches(0.35), Inches(0.6), Inches(0.55), dot, size=22,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(7.55), top + Inches(0.26), Inches(5.0), Inches(0.5), title, size=19,
           bold=True, color=color)
        tb(slide, Inches(7.55), top + Inches(0.76), Inches(5.0), Inches(0.42), blurb, size=14,
           color=DARK)


def _story_slide(lines, part, timing):
    slide, n = new_slide(f"📖 Maya's Three-Day Adventure — {part}", "STORY", timing, part,
                         AMBER)
    add_round(slide, Inches(0.5), Inches(1.45), Inches(12.35), Inches(4.75), WHITE)
    gap = 0.5 if len(lines) > 8 else 0.55
    for i, (line, color) in enumerate(lines):
        top = Inches(1.65 + i * gap)
        add_oval(slide, Inches(0.8), top + Inches(0.12), Inches(0.22), Inches(0.22), color)
        tb(slide, Inches(1.25), top, Inches(11.4), Inches(0.5), line, size=20, bold=True,
           color=NAVY)
    tense_legend(slide, 6.5)
    return slide, n


def s30_story_1():
    _story_slide(STORY_P1, "Part 1", "82–87 min")


def s31_story_2():
    _story_slide(STORY_P2, "Part 2", "82–87 min")


def s32_grammar_hunt():
    slide, n = new_slide("🔍 Story Grammar Hunt", "GAME", "82–87 min", "Grammar hunt", TEAL)
    subtitle(slide, "Which tense is each sentence? Point to your answer.")
    for i, sentence in enumerate(HUNT_SENTENCES):
        top = Inches(1.9 + i * 0.84)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.72), WHITE)
        add_oval(slide, Inches(0.7), top + Inches(0.16), Inches(0.42), Inches(0.42), TEAL)
        tb(slide, Inches(0.7), top + Inches(0.2), Inches(0.42), Inches(0.34), str(i + 1),
           size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.3), top + Inches(0.16), Inches(5.3), Inches(0.44), sentence,
           size=16, bold=True, color=NAVY)
        for j, (name, dot, color, light) in enumerate(TENSES):
            left = Inches(6.9 + j * 2.0)
            add_round(slide, left, top + Inches(0.11), Inches(1.85), Inches(0.5), light)
            tb(slide, left, top + Inches(0.16), Inches(1.85), Inches(0.4), f"{dot} {name}",
               size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.95), Inches(12.35), Inches(0.02), CREAM)


def s33_story_questions():
    slide, n = new_slide("💬 Story Questions — Say It in a Sentence", "SPEAKING", "82–87 min",
                         "Story questions", CORAL)
    subtitle(slide, "Answer in a FULL sentence, using the right tense.")
    for i, (question, tense) in enumerate(STORY_QUESTIONS):
        top = Inches(1.95 + i * 1.25)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.08), WHITE)
        add_oval(slide, Inches(0.75), top + Inches(0.28), Inches(0.52), Inches(0.52), CORAL)
        tb(slide, Inches(0.75), top + Inches(0.34), Inches(0.52), Inches(0.4), str(i + 1),
           size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.5), top + Inches(0.3), Inches(7.6), Inches(0.55), question,
           size=19, bold=True, color=NAVY)
        add_round(slide, Inches(9.6), top + Inches(0.28), Inches(3.0), Inches(0.55), L_AMBER)
        tb(slide, Inches(9.6), top + Inches(0.36), Inches(3.0), Inches(0.4), tense, size=14,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.95), Inches(12.35), Inches(0.02), CREAM)


def s34_roll_and_talk():
    slide, n = new_slide("🎲 Roll & Create  +  🎤 3-Tense Talk", "SPEAKING", "87–89 min",
                         "Roll & talk", FUTURE)
    subtitle(slide, "Pick a number, then make all THREE sentences with that verb.")
    for i, (num, verb, emoji) in enumerate(ROLL_VERBS):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.9 + row * 1.5)
        add_round(slide, left, top, Inches(3.9), Inches(1.3), L_GREY)
        add_oval(slide, left + Inches(0.22), top + Inches(0.34), Inches(0.62), Inches(0.62),
                 FUTURE)
        tb(slide, left + Inches(0.22), top + Inches(0.42), Inches(0.62), Inches(0.45), num,
           size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.05), top + Inches(0.36), Inches(0.7), Inches(0.58), emoji,
           size=20, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.9), top + Inches(0.38), Inches(1.8), Inches(0.55), verb,
           size=24, bold=True, color=NAVY, font="Arial Black")
    for i, (label, prompt, color, light) in enumerate(TALK_PROMPTS):
        left = Inches(0.5 + i * 4.15)
        add_round(slide, left, Inches(4.95), Inches(3.9), Inches(1.5), light)
        add_round(slide, left + Inches(1.1), Inches(5.15), Inches(1.7), Inches(0.45), color)
        tb(slide, left + Inches(1.1), Inches(5.21), Inches(1.7), Inches(0.36), label,
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.2), Inches(5.75), Inches(3.5), Inches(0.55), prompt,
           size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def s35_champion():
    slide, n = new_slide("🏆 Final Game: Tense Champion", "CHALLENGE", "87–89 min",
                         "Tense champion", GOLD)
    subtitle(slide, "Five rapid-fire challenges. Celebrate every single one!")
    quiz_rows(slide, CHAMPION, AMBER, L_AMBER, 1, top_start=1.9, gap=1.02, height=0.9,
              prompt_width=4.0, prompt_size=17, opt_left=5.9, opt_width=2.2, opt_gap=2.4,
              opt_size=16)


def s36_exit_ticket():
    slide, n = new_slide("🎟️ Exit Ticket — Make Your Own Sentences", "RECAP", "89–90 min",
                         "Exit ticket", TEAL)
    subtitle(slide, "Three sentences, all your own. Any verb you like!")
    prompts = [("🟢 PRESENT", "I ______ every day.", PRESENT, L_PRESENT),
               ("🔵 PAST", "Yesterday, I ______.", PAST, L_PAST),
               ("🟣 FUTURE", "Tomorrow, I will ______.", FUTURE, L_FUTURE)]
    for i, (label, prompt, color, light) in enumerate(prompts):
        top = Inches(1.95 + i * 1.6)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.4), light)
        add_round(slide, Inches(0.85), top + Inches(0.42), Inches(2.2), Inches(0.58), color)
        tb(slide, Inches(0.85), top + Inches(0.51), Inches(2.2), Inches(0.42), label,
           size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(3.5), top + Inches(0.3), Inches(9.1), Inches(0.8), WHITE)
        tb(slide, Inches(3.8), top + Inches(0.45), Inches(8.5), Inches(0.55), prompt,
           size=24, bold=True, color=NAVY)


def s37_badge():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.7, 0.85, PRESENT), (11.9, 0.85, PAST), (0.9, 5.75, FUTURE),
                    (11.85, 5.7, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(0.95), Inches(12), Inches(0.8),
       "⭐ GRAMMAR TIME TRAVELER ⭐", size=36, bold=True, color=GOLD,
       align=PP_ALIGN.CENTER, font="Georgia")
    add_oval(slide, Inches(5.42), Inches(1.85), Inches(2.5), Inches(2.5), GOLD)
    tb(slide, Inches(5.42), Inches(2.4), Inches(2.5), Inches(1.4), "🕰️", size=52,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(4.6), Inches(12), Inches(0.6),
       "You completed the mission!", size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
       font="Georgia")
    tb(slide, Inches(0.7), Inches(5.28), Inches(12), Inches(0.5),
       "Chrono-5 is repaired. All three time zones unlocked.", size=19, color=WHITE,
       align=PP_ALIGN.CENTER)
    for i, (name, dot, color, _light) in enumerate(TENSES):
        left = Inches(2.6 + i * 2.9)
        add_round(slide, left, Inches(5.95), Inches(2.6), Inches(0.62), color)
        tb(slide, left, Inches(6.07), Inches(2.6), Inches(0.42), f"{dot}  {name}  ✓",
           size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(slide, n, "89–90 min", "Badge")
    fade(slide)


def s38_support():
    slide, n = new_slide("🧰 Support Ladder — If She Gets Stuck", "TEACHER ONLY", "",
                         "Support", SLATE)
    tb(slide, Inches(0.45), Inches(1.34), Inches(12.4), Inches(0.35),
       "TEACHER ONLY — do not reveal the answer. Climb these hints one step at a time.",
       size=13, bold=True, color=CORAL)
    for i, (label, title, detail, color) in enumerate(HINTS):
        left = Inches(0.45 + i * 3.18)
        add_round(slide, left, Inches(1.85), Inches(3.0), Inches(2.1), WHITE)
        add_round(slide, left + Inches(0.2), Inches(2.05), Inches(1.3), Inches(0.45), color)
        tb(slide, left + Inches(0.2), Inches(2.12), Inches(1.3), Inches(0.34), label,
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.2), Inches(2.62), Inches(2.6), Inches(0.62), title,
           size=16, bold=True, color=NAVY)
        tb(slide, left + Inches(0.2), Inches(3.3), Inches(2.6), Inches(0.5), detail, size=12,
           color=SOFT)
    tb(slide, Inches(0.45), Inches(4.18), Inches(12.4), Inches(0.4),
       "Then lower the difficulty until she succeeds:", size=16, bold=True, color=NAVY)
    for i, (level, what, example) in enumerate(SUPPORT_LEVELS):
        left = Inches(0.45 + i * 3.18)
        add_round(slide, left, Inches(4.7), Inches(3.0), Inches(2.0), L_TEAL)
        tb(slide, left + Inches(0.18), Inches(4.88), Inches(2.64), Inches(0.42), level,
           size=13, bold=True, color=TEAL)
        tb(slide, left + Inches(0.18), Inches(5.35), Inches(2.64), Inches(0.6), what, size=13,
           color=DARK)
        add_round(slide, left + Inches(0.18), Inches(5.98), Inches(2.64), Inches(0.52), WHITE)
        tb(slide, left + Inches(0.28), Inches(6.08), Inches(2.44), Inches(0.38), example,
           size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def s39_answer_key():
    slide, n = new_slide("🔒 TEACHER ONLY — Answer Key", "TEACHER ONLY", "", "Answer key",
                         SLATE)
    tb(slide, Inches(0.45), Inches(1.34), Inches(12.4), Inches(0.35),
       "Hide this slide before presenting, or keep it on a second screen.",
       size=13, bold=True, color=CORAL)
    add_round(slide, Inches(0.45), Inches(1.78), Inches(6.1), Inches(5.05), WHITE)
    tb(slide, Inches(0.7), Inches(1.94), Inches(5.6), Inches(0.38),
       "🕵️ Time Detective  &  🟢 Present", size=14, bold=True, color=PRESENT)
    bullets(slide, Inches(0.7), Inches(2.38), Inches(5.6), Inches(4.3), [
        "Time Detective 1–4: Past, Future, Present, Present",
        "Time Detective 5–8: Past, Future, Present, Past",
        "Missing S 1–4: goes, plays, runs, read",
        "Missing S 5–8: watches, play, goes, do",
        "Routine: wakes / eats / goes / plays / reads / sleeps",
        "-es after ch, sh, s, x, o (watches, goes, fixes,",
        "     washes)",
        "Irregular past: went, ate, saw, came, had, ran",
        "Yesterday 1–3: B, A, C",
        "Yesterday 4–6: A, B, B",
    ], size=12, sp=7)
    add_round(slide, Inches(6.75), Inches(1.78), Inches(6.1), Inches(5.05), WHITE)
    tb(slide, Inches(7.0), Inches(1.94), Inches(5.6), Inches(0.38),
       "🟣 Future · 🔄 Transform · 🕵️ Mistakes · 📖 Story", size=14, bold=True, color=FUTURE)
    bullets(slide, Inches(7.0), Inches(2.38), Inches(5.6), Inches(4.3), [
        "Future Challenge 1–5: will, will, watch, go, do",
        "Transform: She read / She will read",
        "     We watched / We will watch",
        "     He went / He will go",
        "     They visited / They will visit",
        "     I ate / I will eat",
        "Mistakes: She goes · Yesterday I went ·",
        "     Tomorrow I will go · I will go ·",
        "     He watches · Last week we visited",
        "Grammar Hunt: Present, Past, Future,",
        "     Present, Past, Future",
        "Champion 1–5: go, watched, will visit,",
        "     plays, visited",
    ], size=12, sp=7)


def s40_assessment():
    slide, n = new_slide("📋 TEACHER ONLY — End-of-Class Assessment", "TEACHER ONLY", "",
                         "Assessment", SLATE)
    tb(slide, Inches(0.45), Inches(1.34), Inches(12.4), Inches(0.35),
       "Record progress, not failure. Complete this within five minutes of finishing.",
       size=13, bold=True, color=CORAL)
    cols = [(0.45, 4.6, "SKILL"), (5.25, 2.35, "Needs Support"), (7.85, 2.35, "Developing"),
            (10.45, 2.4, "Confident")]
    header_y = Inches(1.8)
    for left, width, label in cols:
        add_round(slide, Inches(left), header_y, Inches(width), Inches(0.46), NAVY)
        tb(slide, Inches(left), header_y + Inches(0.08), Inches(width), Inches(0.32), label,
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, skill in enumerate(RUBRIC_SKILLS):
        top = header_y + Inches(0.54 + i * 0.44)
        band = WHITE if i % 2 == 0 else L_GREY
        add_round(slide, Inches(0.45), top, Inches(4.6), Inches(0.39), band)
        tb(slide, Inches(0.65), top + Inches(0.05), Inches(4.2), Inches(0.3), skill, size=13,
           bold=True, color=NAVY)
        for left, width, _l in cols[1:]:
            add_round(slide, Inches(left), top, Inches(width), Inches(0.39), band)
            tb(slide, Inches(left), top + Inches(0.03), Inches(width), Inches(0.33), "☐",
               size=14, color=SOFT, align=PP_ALIGN.CENTER)
    for i, label in enumerate(RUBRIC_NOTES):
        top = Inches(5.92 + i * 0.36)
        add_round(slide, Inches(0.45), top, Inches(12.4), Inches(0.31), L_AMBER)
        tb(slide, Inches(0.65), top + Inches(0.02), Inches(12.0), Inches(0.26), f"{label}:",
           size=12, bold=True, color=NAVY)


BUILDERS = [
    s01_title, s02_time_machine, s03_mission, s04_icebreaker, s05_now_before_later,
    s06_what_is_tense, s07_time_detective_a, s08_time_detective_b, s09_present_intro,
    s10_present_examples, s11_routine_detective, s12_he_she_it, s13_missing_s_a,
    s14_missing_s_b, s15_brain_break, s16_past_intro, s17_regular_verbs,
    s18_irregular_verbs, s19_yesterday_a, s20_yesterday_b, s21_future_intro,
    s22_future_predictor, s23_future_challenge, s24_transformation,
    s25_transform_challenge, s26_mistakes_a, s27_mistakes_b, s28_verb_table,
    s29_story_intro, s30_story_1, s31_story_2, s32_grammar_hunt, s33_story_questions,
    s34_roll_and_talk, s35_champion, s36_exit_ticket, s37_badge, s38_support,
    s39_answer_key, s40_assessment,
]

for build in BUILDERS:
    build()

assert len(prs.slides) == TOTAL, f"expected {TOTAL} slides, built {len(prs.slides)}"

# This deck ships without speaker notes.
for slide in prs.slides:
    if slide.has_notes_slide:
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype.endswith("/notesSlide"):
                slide.part.drop_rel(rid)

out = r"C:\Users\bhushaja\Downloads\shaip\Grade5_Time_Travel_Grammar_90min.pptx"
try:
    prs.save(out)
except PermissionError:
    out = out.replace(".pptx", "_new.pptx")
    prs.save(out)

story_words = sum(len(line.split()) for line, _c in STORY_P1 + STORY_P2)
with_notes = [i + 1 for i, s in enumerate(prs.slides) if s.has_notes_slide]
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
print(f"Story word count: {story_words}")
print(f"Slides carrying notes: {with_notes if with_notes else 'none'}")


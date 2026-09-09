"""Remove speaker notes from a deck, keeping a local copy that still has them."""
import shutil
import sys

from pptx import Presentation


def strip(path, keep_master=True):
    if keep_master:
        master = path.replace(".pptx", "_TeacherNotes.pptx")
        shutil.copyfile(path, master)
        print(f"kept notes copy: {master}")

    prs = Presentation(path)
    cleared = 0
    for slide in prs.slides:
        if not slide.has_notes_slide:
            continue
        tf = slide.notes_slide.notes_text_frame
        if tf.text.strip():
            tf.clear()
            tf.text = ""
            cleared += 1
    prs.save(path)

    check = Presentation(path)
    remaining = [i + 1 for i, s in enumerate(check.slides)
                 if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()]
    print(f"{path}: cleared {cleared} slides, "
          f"remaining with notes: {remaining if remaining else 'none'}")
    return not remaining


if __name__ == "__main__":
    ok = all(strip(p) for p in sys.argv[1:])
    sys.exit(0 if ok else 1)

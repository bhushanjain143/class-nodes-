"""Flag overlapping text boxes and off-canvas shapes in a generated deck."""
import sys

from pptx import Presentation
from pptx.util import Emu

OVERLAP_LIMIT = 0.15


def safe(text):
    return text.encode("ascii", "replace").decode("ascii").replace("\n", " ")


def text_boxes(slide):
    found = []
    for shape in slide.shapes:
        if shape.left is None or shape.top is None or not shape.has_text_frame:
            continue
        body = shape.text_frame.text.strip()
        if body:
            found.append((shape.left, shape.top, shape.width, shape.height, body))
    return found


def overlap_ratio(a, b):
    ax, ay, aw, ah, _ = a
    bx, by, bw, bh, _ = b
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    if ix <= 0 or iy <= 0:
        return 0.0
    return (ix * iy) / min(aw * ah, bw * bh)


def check(path):
    prs = Presentation(path)
    width, height = prs.slide_width, prs.slide_height
    problems = 0

    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            off = (shape.left < -1000 or shape.top < -1000
                   or shape.left + shape.width > width + 1000
                   or shape.top + shape.height > height + 1000)
            if off:
                print(f"slide {index}: OUT OF BOUNDS "
                      f"L={Emu(shape.left).inches:.2f} T={Emu(shape.top).inches:.2f} "
                      f"R={Emu(shape.left + shape.width).inches:.2f} "
                      f"B={Emu(shape.top + shape.height).inches:.2f}")
                problems += 1

        boxes = text_boxes(slide)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                ratio = overlap_ratio(boxes[i], boxes[j])
                if ratio > OVERLAP_LIMIT:
                    print(f"slide {index}: OVERLAP {ratio:.0%} -> "
                          f"[{safe(boxes[i][4])[:34]}] vs [{safe(boxes[j][4])[:34]}]")
                    problems += 1

    notes_missing = [i + 1 for i, s in enumerate(prs.slides)
                     if not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip()]
    print(f"slides: {len(prs.slides)}")
    print(f"slides missing teacher notes: {notes_missing if notes_missing else 'none'}")
    print(f"total geometry problems: {problems}")
    return problems


if __name__ == "__main__":
    sys.exit(1 if check(sys.argv[1]) else 0)

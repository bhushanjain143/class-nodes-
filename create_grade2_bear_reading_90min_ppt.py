"""Grade 2 reading intervention - 90 minutes, 35 slides.

"The Bear's Reading Adventure" - a one-to-one lesson for a struggling reader.
Builds sounds -> letters -> blending -> CVC words -> sight words -> sentences
-> a short story, using an I DO / WE DO / YOU DO cycle throughout.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x16, 0x36, 0x6B)
FOREST = RGBColor(0x2E, 0x7D, 0x53)
TEAL = RGBColor(0x00, 0x9C, 0x8F)
SKY = RGBColor(0x2E, 0xA8, 0xE8)
AMBER = RGBColor(0xF2, 0xA1, 0x1B)
GOLD = RGBColor(0xFF, 0xC4, 0x2E)
CORAL = RGBColor(0xE8, 0x5A, 0x4F)
PLUM = RGBColor(0x7B, 0x4F, 0xB8)
BROWN = RGBColor(0x8B, 0x5E, 0x3C)
CREAM = RGBColor(0xFF, 0xFB, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x24, 0x2C, 0x38)
SOFT = RGBColor(0x6A, 0x77, 0x88)
L_SKY = RGBColor(0xE2, 0xF3, 0xFD)
L_TEAL = RGBColor(0xDE, 0xF5, 0xF2)
L_AMBER = RGBColor(0xFF, 0xF2, 0xD6)
L_CORAL = RGBColor(0xFD, 0xEA, 0xE8)
L_PLUM = RGBColor(0xEF, 0xE8, 0xFA)
L_FOREST = RGBColor(0xE4, 0xF3, 0xEA)
L_BROWN = RGBColor(0xF4, 0xEA, 0xE1)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TOTAL = 35
_counter = {"n": 0}

# ------------------------------------------------------------------ content

LETTERS = [
    ("A", "/a/", "🍎", "Apple", CORAL, L_CORAL),
    ("B", "/b/", "🐻", "Bear", BROWN, L_BROWN),
    ("C", "/k/", "🐱", "Cat", SKY, L_SKY),
    ("D", "/d/", "🐶", "Dog", FOREST, L_FOREST),
    ("E", "/e/", "🥚", "Egg", AMBER, L_AMBER),
    ("S", "/s/", "☀️", "Sun", GOLD, L_AMBER),
    ("M", "/m/", "🌙", "Moon", PLUM, L_PLUM),
    ("R", "/r/", "🔴", "Red", CORAL, L_CORAL),
]

MYSTERY = [
    ("I make the sound /m/.", "You hear me in Moon.", "M", "🌙"),
    ("I make the sound /b/.", "You hear me in Bear.", "B", "🐻"),
    ("I make the sound /s/.", "You hear me in Sun.", "S", "☀️"),
    ("I make the sound /t/.", "You hear me at the END of Cat.", "T", "🐱"),
]

BEGINNING_SOUNDS = [
    ("🐻", "bear", "B"),
    ("🐱", "cat", "C"),
    ("☀️", "sun", "S"),
    ("🌙", "moon", "M"),
]

DIGRAPHS = [("sh", "🤫", "shhh"), ("ch", "🚂", "ch-ch"), ("th", "👍", "thumb"), ("wh", "🐋", "whale")]

CVC_WORDS = [
    ("CAT", ["C", "A", "T"], "🐱", SKY, L_SKY),
    ("SUN", ["S", "U", "N"], "☀️", GOLD, L_AMBER),
    ("DOG", ["D", "O", "G"], "🐶", FOREST, L_FOREST),
    ("MAP", ["M", "A", "P"], "🗺️", PLUM, L_PLUM),
    ("PIG", ["P", "I", "G"], "🐷", CORAL, L_CORAL),
    ("HAT", ["H", "A", "T"], "🎩", TEAL, L_TEAL),
    ("RUN", ["R", "U", "N"], "🏃", AMBER, L_AMBER),
    ("BIG", ["B", "I", "G"], "🐘", BROWN, L_BROWN),
]

FEED_BEAR = [
    ("🐱", "CAT", ["CAT", "SUN", "DOG"]),
    ("☀️", "SUN", ["MAP", "SUN", "PIG"]),
    ("🐶", "DOG", ["HAT", "RUN", "DOG"]),
    ("🐘", "BIG", ["BIG", "RED", "CAT"]),
]

HUNT_GRID = [
    ["CAT", "SUN", "DOG", "MAP", "HAT"],
    ["RUN", "BIG", "CAT", "PIG", "SUN"],
    ["RED", "DOG", "HAT", "BIG", "RUN"],
]
HUNT_ROUNDS = [("CAT", 2), ("SUN", 2), ("DOG", 2), ("BIG", 2), ("RUN", 2), ("HAT", 2)]

SIGHT_SETS = [
    ("Round 1", ["I", "see", "the"], TEAL, L_TEAL),
    ("Round 2", ["can", "my", "is"], CORAL, L_CORAL),
    ("Round 3", ["we", "go", "and", "you"], PLUM, L_PLUM),
]

DETECTIVE = [
    ("the", ["and", "the", "you", "see"]),
    ("can", ["can", "cat", "car", "cap"]),
    ("my", ["me", "my", "may", "man"]),
    ("see", ["saw", "sea", "see", "so"]),
]

MEMORY_PAIRS = [("🐱", "CAT"), ("☀️", "SUN"), ("🐶", "DOG"), ("🐻", "BEAR")]

SIMON_SAYS = [
    "Point to the word CAT.",
    "Point to the word SUN.",
    "Touch the word CAN.",
    "Show me the letter B.",
    "Point to the word THE.",
    "Show me the letter M.",
]
SIMON_BOARD = ["CAT", "SUN", "CAN", "THE", "B", "M"]

SENTENCES = [
    ("I see a cat.", "🐱", SKY, L_SKY),
    ("The cat is big.", "🐱", TEAL, L_TEAL),
    ("I can run.", "🏃", FOREST, L_FOREST),
    ("We see the sun.", "☀️", AMBER, L_AMBER),
    ("My dog can run.", "🐶", CORAL, L_CORAL),
]

SCRAMBLES = [
    (["cat", "I", "a", "see"], "I see a cat.", "🐱"),
    (["big", "The", "is", "cat"], "The cat is big.", "🐱"),
    (["run", "I", "can"], "I can run.", "🏃"),
    (["sun", "the", "We", "see"], "We see the sun.", "☀️"),
    (["can", "dog", "My", "run"], "My dog can run.", "🐶"),
    (["red", "a", "has", "Ben", "ball"], "Ben has a red ball.", "🔴"),
]

VOCAB = [
    ("🐻", "bear", "b - ear", "Ben is a bear."),
    ("🔴", "ball", "b - all", "I see a ball."),
    ("🌳", "tree", "t - ree", "The tree is big."),
    ("🎨", "red", "r - e - d", "My ball is red."),
    ("🏃", "run", "r - u - n", "I can run."),
    ("🐘", "big", "b - i - g", "The dog is big."),
    ("😀", "happy", "hap - py", "Ben is happy."),
    ("🐣", "little", "lit - tle", "A little bear."),
]

STORY = [
    ("Part 1", "🐻", ["Ben is a little bear.", "Ben has a red ball.",
                      "Ben can run and play.", "Ben is happy."], BROWN, L_BROWN),
    ("Part 2", "❓", ["One day, the ball is gone.", "Ben can not see the ball.",
                      "Ben looks and looks.", "Ben sees a big tree."], SKY, L_SKY),
    ("Part 3", "🌳", ["The ball is under the tree.", "A cat sits on the ball.",
                      "\"My ball!\" says Ben.", "The cat runs to Ben."], FOREST, L_FOREST),
    ("Part 4", "😀", ["Ben gets the red ball.", "Ben and the cat play.",
                      "The sun is big and hot.", "Ben is happy."], AMBER, L_AMBER),
]

COMPREHENSION = [
    ("Who is the story about?", [("🐻", "Bear"), ("🐶", "Dog"), ("🐱", "Cat")], "Bear"),
    ("What did Ben have?", [("🔴", "Ball"), ("🍎", "Apple"), ("📚", "Book")], "Ball"),
    ("Where was the ball?", [("🌳", "Under the tree"), ("🏠", "In the house"),
                             ("🪑", "On the chair")], "Under the tree"),
    ("What happened to the ball?", [("❓", "It was gone"), ("💥", "It broke"),
                                    ("🔵", "It turned blue")], "It was gone"),
    ("How did Ben feel at the end?", [("😀", "Happy"), ("😢", "Sad"), ("😠", "Mad")], "Happy"),
]

PICTURE_MATCH = [
    ("Ben has a red ball.", ["🔴", "🍎", "📚"], "🔴"),
    ("The cat sits on the ball.", ["🐶", "🐱", "🐻"], "🐱"),
    ("The ball is under the tree.", ["🌳", "🏠", "🚗"], "🌳"),
]

FINAL_WORDS = ["CAT", "SUN", "BIG", "RUN", "RED"]
FINAL_SENTENCES = ["I see a cat.", "The sun is big.", "Ben can run."]

RUBRIC_SKILLS = [
    "Letter-sound recognition", "Phonics / blending", "Word recognition",
    "Sight words", "Sentence reading", "Comprehension",
    "Speaking", "Confidence", "Participation",
]

HOMEWORK = [
    ("🔤", "Say the sounds", "Read CAT, SUN, DOG out loud", "3 min"),
    ("👀", "Sight word hunt", "Find THE and CAN in any book", "5 min"),
    ("📖", "Read to someone", "Read \"I see a cat.\" to a grown-up", "3 min"),
    ("🎨", "Draw Ben", "Draw the bear and his red ball", "5 min"),
]

# ------------------------------------------------------------------ helpers


def set_run(run, size=20, bold=False, color=DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    tree = slide.shapes._spTree
    el = s._element
    tree.remove(el)
    tree.insert(2, el)


def add_rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_round(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_oval(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def tb(slide, l, t, w, h, text, size=20, bold=False, color=DARK,
       align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size, bold, color, font)
    return box


def bullets(slide, l, t, w, h, items, size=15, color=DARK, sp=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(sp)
        r = p.add_run()
        r.text = "•  " + item
        set_run(r, size, False, color)
    return box


def fade(slide):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", "med")
    tr.set("advClick", "1")
    etree.SubElement(tr, qn("p:fade"))


def footer(slide, n, timing="", step=""):
    add_rect(slide, Inches(0), Inches(7.02), prs.slide_width, Inches(0.08),
             RGBColor(0xE4, 0xEA, 0xF0))
    add_rect(slide, Inches(0), Inches(7.02), int(prs.slide_width * n / TOTAL), Inches(0.08), GOLD)
    add_rect(slide, Inches(0), Inches(7.10), prs.slide_width, Inches(0.40), NAVY)
    msg = "🐻 The Bear's Reading Adventure  |  Grade 2  |  90 min"
    if step:
        msg += f"  |  {step}"
    if timing:
        msg += f"  |  {timing}"
    tb(slide, Inches(0.35), Inches(7.14), Inches(11.3), Inches(0.32), msg, size=10, color=WHITE)
    tb(slide, Inches(11.85), Inches(7.14), Inches(1.2), Inches(0.32), f"{n} / {TOTAL}",
       size=10, color=WHITE, align=PP_ALIGN.RIGHT)


def chip(slide, text, color):
    add_round(slide, Inches(0.4), Inches(0.26), Inches(3.2), Inches(0.42), color)
    tb(slide, Inches(0.4), Inches(0.31), Inches(3.2), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def time_chip(slide, text):
    add_round(slide, Inches(10.3), Inches(0.26), Inches(2.65), Inches(0.42), CORAL)
    tb(slide, Inches(10.3), Inches(0.31), Inches(2.65), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def star_badge(slide, left=Inches(10.3), top=Inches(6.42), text="⭐ +1 Reading Star"):
    add_round(slide, left, top, Inches(2.65), Inches(0.5), GOLD)
    tb(slide, left, top + Inches(0.08), Inches(2.65), Inches(0.38), text, size=13, bold=True,
       color=NAVY, align=PP_ALIGN.CENTER)


def new_slide(title, tag, timing, step, accent=SKY, bg=CREAM):
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_rect(slide, Inches(0), Inches(0), Inches(0.16), prs.slide_height, accent)
    chip(slide, tag, accent)
    if timing:
        time_chip(slide, timing)
    tb(slide, Inches(0.4), Inches(0.8), Inches(12.4), Inches(0.6), title, size=28, bold=True,
       color=NAVY, font="Georgia")
    footer(slide, n, timing, step)
    fade(slide)
    return slide, n


def notes(slide, say, instructions, expected, help_step, easier, challenge, praise,
          observe, timing):
    tf = slide.notes_slide.notes_text_frame
    tf.text = (
        f"⏱ TIMING: {timing}\n\n"
        f"SAY: {say}\n\n"
        f"INSTRUCTIONS: {instructions}\n\n"
        f"EXPECTED RESPONSE: {expected}\n\n"
        f"IF THE STUDENT STRUGGLES: {help_step}\n\n"
        f"EASIER ALTERNATIVE: {easier}\n\n"
        f"OPTIONAL CHALLENGE: {challenge}\n\n"
        f"ENCOURAGEMENT: {praise}\n\n"
        f"OBSERVE (assessment): {observe}\n\n"
        "NEVER say \"Wrong.\" Say \"Almost! Let's try it together.\" "
        "Give 5 seconds of thinking time before helping."
    )


def ido_wedo_youdo(slide, top, accent, labels=("I read it", "We read it", "YOU read it!")):
    steps = [("1️⃣", labels[0], "Listen and follow", L_AMBER, NAVY),
             ("2️⃣", labels[1], "Say it with me", L_TEAL, NAVY),
             ("3️⃣", labels[2], "Your turn ⭐", accent, WHITE)]
    for i, (num, title, sub, bg, fg) in enumerate(steps):
        left = Inches(0.55 + i * 4.15)
        add_round(slide, left, top, Inches(3.9), Inches(1.5), bg)
        tb(slide, left + Inches(0.15), top + Inches(0.14), Inches(3.6), Inches(0.45), num,
           size=22, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(0.62), Inches(3.6), Inches(0.45), title,
           size=18, bold=True, color=fg, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(1.06), Inches(3.6), Inches(0.36), sub,
           size=13, color=fg, align=PP_ALIGN.CENTER)


# ------------------------------------------------------------------ slides


def s01_welcome():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.65, 0.75, SKY), (11.95, 0.8, CORAL), (0.85, 5.85, TEAL), (11.9, 5.8, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(1.25), Inches(12), Inches(1.1), "🐻", size=64,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(2.35), Inches(12), Inches(0.85),
       "Welcome, Reading Explorer!", size=42, bold=True, color=WHITE,
       align=PP_ALIGN.CENTER, font="Georgia")
    tb(slide, Inches(0.7), Inches(3.3), Inches(12), Inches(0.5),
       "The Bear's Reading Adventure  •  Grade 2  •  90 Minutes",
       size=18, color=L_SKY, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(3.5), Inches(4.05), Inches(6.3), Inches(1.25), FOREST)
    tb(slide, Inches(3.7), Inches(4.3), Inches(5.9), Inches(0.8),
       "Ben the Bear needs a reading buddy.\nToday, that's YOU! ⭐",
       size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(5.65), Inches(12), Inches(0.5),
       "🔤 Sounds  →  🧩 Words  →  📖 Sentences  →  📚 A Story",
       size=18, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    footer(slide, n, "0–5 min", "Welcome")
    fade(slide)
    notes(slide,
          "Hi! I am so glad you are here. Today you are a Reading Explorer, and this is Ben "
          "the Bear. Ben needs a reading buddy to help him on an adventure. That's you!",
          "Smile and wave. Say your name and Ben's name. Do not ask the student to read "
          "anything on this slide. Keep this to about 45 seconds.",
          "The student says hello, smiles, or waves. Any response is a success.",
          "If the student is quiet, do not push. Wave and say 'You can just nod today.'",
          "Simply say 'Wave to Ben!' and wave together.",
          "Ask which animal they would pick as a reading buddy.",
          "I am so happy you are here. You are already being a great explorer!",
          "Note comfort level, eye contact, and willingness to speak at the start.",
          "0–5 min")


def s02_mission():
    slide, n = new_slide("🗺️ Today's Mission", "MISSION", "0–5 min", "Mission", AMBER)
    stops = [
        ("🔤", "Letter Sounds", "5–15 min", CORAL),
        ("🎵", "Blending", "15–25 min", TEAL),
        ("🐻", "Feed the Bear", "25–35 min", BROWN),
        ("👀", "Sight Words", "35–45 min", PLUM),
        ("📖", "Sentences", "50–68 min", SKY),
        ("📚", "Ben's Story", "68–84 min", FOREST),
        ("🎯", "Big Challenge", "84–87 min", AMBER),
        ("🏆", "Your Badge", "87–90 min", GOLD),
    ]
    for i, (icon, label, when, color) in enumerate(stops):
        col, row = i % 4, i // 4
        left = Inches(0.45 + col * 3.18)
        top = Inches(1.6 + row * 2.4)
        add_round(slide, left, top, Inches(3.0), Inches(2.15), WHITE)
        add_oval(slide, left + Inches(1.15), top + Inches(0.2), Inches(0.7), Inches(0.7), color)
        tb(slide, left + Inches(1.15), top + Inches(0.31), Inches(0.7), Inches(0.5), icon,
           size=20, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.02), Inches(2.76), Inches(0.65), label,
           size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.62), Inches(2.76), Inches(0.35), when,
           size=12, color=SOFT, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.45), Inches(6.45), Inches(9.5), Inches(0.45),
       "⭐ Earn a Reading Star at every stop. Collect all 8!", size=16, bold=True, color=CORAL)
    notes(slide,
          "Here is our map for today. Eight stops, and you earn a Reading Star at every one. "
          "Nothing lasts long — if something feels hard, we move on quickly.",
          "Point to each stop as you name it. Do not read the small times aloud. "
          "Keep this under 45 seconds.",
          "The student looks at the map and may point to a stop they like.",
          "Just point to the bear stop and say 'This one is my favorite.'",
          "Name only three stops instead of eight.",
          "Ask which stop looks the most fun and why.",
          "Eight stars is a lot — and I think you can get them all!",
          "Note attention span and whether the student engages with the plan.",
          "0–5 min")


def s03_feelings():
    slide, n = new_slide("💛 How Are You Today?", "CONNECT", "0–5 min", "Feelings", TEAL)
    tb(slide, Inches(0.45), Inches(1.42), Inches(11.5), Inches(0.45),
       "Point to your answer — no reading needed!", size=17, bold=True, color=CORAL)
    faces = [("😀", "Great"), ("🙂", "Good"), ("😐", "Okay"), ("😴", "Tired")]
    for i, (emoji, label) in enumerate(faces):
        left = Inches(0.55 + i * 3.15)
        add_round(slide, left, Inches(2.0), Inches(2.95), Inches(2.0), WHITE)
        tb(slide, left + Inches(0.1), Inches(2.25), Inches(2.75), Inches(0.9), emoji, size=44,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.1), Inches(3.3), Inches(2.75), Inches(0.45), label, size=17,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    questions = [("🐾", "What is your favorite animal?"), ("🎨", "What is your favorite color?"),
                 ("📚", "Do you like stories?")]
    for i, (icon, q) in enumerate(questions):
        top = Inches(4.35 + i * 0.72)
        add_round(slide, Inches(0.55), top, Inches(9.3), Inches(0.6), L_TEAL)
        tb(slide, Inches(0.75), top + Inches(0.11), Inches(0.5), Inches(0.4), icon, size=16)
        tb(slide, Inches(1.35), top + Inches(0.12), Inches(8.3), Inches(0.4), q, size=17,
           bold=True, color=NAVY)
    star_badge(slide, Inches(10.3), Inches(4.35), "⭐ Star #1 earned!")
    notes(slide,
          "How are you feeling today? Just point to a face. There is no wrong answer. "
          "Now tell me — what is your favorite animal?",
          "Let the student point to a face, then ask the three questions one at a time. "
          "Answer each one yourself first so they hear a model.",
          "A pointed face and short spoken answers such as 'dog', 'blue', 'yes'.",
          "Answer for yourself first: 'My favorite animal is a bear. What about you?'",
          "Offer two choices: 'Dogs or cats?'",
          "Ask them to say why they like that animal.",
          "Thank you for telling me! You just earned your very first Reading Star.",
          "Note speaking confidence and whether answers are single words or phrases.",
          "0–5 min")


def s04_letter_sounds():
    slide, n = new_slide("🔤 Letter Sound Warm-Up", "PHONICS", "5–15 min", "Letter sounds", CORAL)
    tb(slide, Inches(0.45), Inches(1.4), Inches(11.5), Inches(0.42),
       "I say it → You say it. Then I point → You tell me!", size=16, bold=True, color=CORAL)
    for i, (letter, sound, emoji, word, color, light) in enumerate(LETTERS):
        col, row = i % 4, i // 4
        left = Inches(0.45 + col * 3.18)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(3.0), Inches(2.2), light)
        tb(slide, left + Inches(0.12), top + Inches(0.12), Inches(1.3), Inches(0.9), letter,
           size=44, bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left + Inches(1.55), top + Inches(0.32), Inches(1.3), Inches(0.5), sound,
           size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.15), top + Inches(1.15), Inches(2.7), Inches(0.88), WHITE)
        tb(slide, left + Inches(0.3), top + Inches(1.3), Inches(0.7), Inches(0.6), emoji,
           size=22, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.1), top + Inches(1.34), Inches(1.65), Inches(0.5), word,
           size=17, bold=True, color=NAVY)
    notes(slide,
          "Let's warm up our sounds. I will say the letter and the sound, then you say it back. "
          "B says /b/... like Bear. Your turn!",
          "Round 1: you say the letter, sound and picture word; the student repeats. "
          "Round 2: you point to a letter and the student gives the sound alone. "
          "Do only 5 or 6 letters if attention drops.",
          "The student repeats each sound, then names sounds independently in round 2.",
          "Say the sound with the picture: 'Bear... /b/ /b/ /b/. What sound?' Exaggerate it.",
          "Work with just A, B and C. Mastery of a few beats exposure to many.",
          "Ask for another word that starts with the same sound.",
          "Your sounds are getting stronger every time. Well done!",
          "Note which letter sounds are automatic and which need support. Record for next lesson.",
          "5–15 min")


def s05_mystery_letter():
    slide, n = new_slide("🎁 Mystery Letter Game", "GAME", "5–15 min", "Mystery letter", PLUM)
    tb(slide, Inches(0.45), Inches(1.4), Inches(11.5), Inches(0.42),
       "I am thinking of a letter... Who am I?", size=17, bold=True, color=CORAL)
    for i, (clue1, clue2, answer, emoji) in enumerate(MYSTERY):
        col, row = i % 2, i // 2
        left = Inches(0.5 + col * 6.25)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(5.95), Inches(2.2), WHITE)
        add_oval(slide, left + Inches(0.25), top + Inches(0.6), Inches(1.0), Inches(1.0), L_PLUM)
        tb(slide, left + Inches(0.25), top + Inches(0.78), Inches(1.0), Inches(0.65), emoji,
           size=26, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.5), top + Inches(0.42), Inches(4.2), Inches(0.5), clue1,
           size=18, bold=True, color=NAVY)
        tb(slide, left + Inches(1.5), top + Inches(0.95), Inches(4.2), Inches(0.5), clue2,
           size=16, color=SOFT)
        add_round(slide, left + Inches(1.5), top + Inches(1.5), Inches(2.2), Inches(0.55), L_AMBER)
        tb(slide, left + Inches(1.5), top + Inches(1.58), Inches(2.2), Inches(0.4),
           "Who am I?  ___", size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(slide,
          "Here comes a mystery! I am thinking of a letter. It makes the sound /m/. "
          "You hear me in Moon. Who am I?",
          "Read both clues slowly, then wait. Give at least five seconds of silence before "
          "helping. Do one mystery at a time.",
          "M  •  B  •  S  •  T",
          "Say the sound again and stretch it: 'mmmmm - Moon.' Then write the letter in the air.",
          "Show two letter cards and ask which one makes that sound.",
          "Ask them to think of another word starting with that letter.",
          "You figured it out! You are a real letter detective.",
          "Note whether the student connects sound to letter name without a visual prompt.",
          "5–15 min")


def s06_sound_game():
    slide, n = new_slide("🔍 Letter Detective: First Sounds", "GAME", "5–15 min",
                         "First sounds", TEAL)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "What sound do you hear FIRST?", size=17, bold=True, color=CORAL)
    for i, (emoji, word, letter) in enumerate(BEGINNING_SOUNDS):
        left = Inches(0.5 + i * 3.15)
        add_round(slide, left, Inches(1.9), Inches(2.95), Inches(2.35), WHITE)
        tb(slide, left + Inches(0.1), Inches(2.12), Inches(2.75), Inches(0.85), emoji, size=40,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.1), Inches(3.1), Inches(2.75), Inches(0.5), word, size=24,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.75), Inches(3.68), Inches(1.45), Inches(0.45), L_TEAL)
        tb(slide, left + Inches(0.75), Inches(3.74), Inches(1.45), Inches(0.35), "/ ? /",
           size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(4.5), Inches(12.35), Inches(1.85), L_AMBER)
    tb(slide, Inches(0.75), Inches(4.68), Inches(11.8), Inches(0.42),
       "🎈 BONUS (only if the student is ready) — Two letters, one sound:",
       size=15, bold=True, color=NAVY)
    for i, (pair, emoji, hint) in enumerate(DIGRAPHS):
        left = Inches(0.75 + i * 3.05)
        add_round(slide, left, Inches(5.2), Inches(2.8), Inches(0.95), WHITE)
        tb(slide, left + Inches(0.15), Inches(5.32), Inches(0.8), Inches(0.6), emoji, size=22,
           align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.0), Inches(5.3), Inches(1.6), Inches(0.42), pair, size=24,
           bold=True, color=CORAL)
        tb(slide, left + Inches(1.0), Inches(5.72), Inches(1.6), Inches(0.34), hint, size=12,
           color=SOFT)
    notes(slide,
          "Listen closely. Bbbb-ear. What sound do you hear at the very beginning of bear?",
          "Say each word slowly and stretch the first sound. The student gives the sound, "
          "not the letter name. Skip the bonus digraph row entirely unless the four "
          "beginning sounds were easy.",
          "bear = /b/  •  cat = /k/  •  sun = /s/  •  moon = /m/. "
          "Bonus: sh, ch, th, wh each make one sound.",
          "Stretch the sound much longer and say it three times before asking again.",
          "Give two sounds to choose from: 'Is it /b/ or /s/?'",
          "Ask for the LAST sound in the word instead.",
          "You heard it! That is exactly what good readers do.",
          "Note if the student confuses letter names with letter sounds — very common at "
          "this stage and worth tracking.",
          "5–15 min")


def s07_blending_intro():
    slide, n = new_slide("🎵 Blending: Sounds Become Words", "PHONICS", "15–25 min",
                         "Blending", FOREST)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(12.35), Inches(2.5), WHITE)
    parts = [("C", "/k/", SKY), ("A", "/a/", CORAL), ("T", "/t/", FOREST)]
    for i, (letter, sound, color) in enumerate(parts):
        left = Inches(1.35 + i * 2.5)
        add_round(slide, left, Inches(1.85), Inches(1.9), Inches(1.8), L_SKY)
        tb(slide, left, Inches(2.0), Inches(1.9), Inches(0.9), letter, size=48, bold=True,
           color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left, Inches(2.95), Inches(1.9), Inches(0.5), sound, size=20, bold=True,
           color=NAVY, align=PP_ALIGN.CENTER)
        if i < 2:
            tb(slide, left + Inches(1.95), Inches(2.35), Inches(0.5), Inches(0.6), "+",
               size=28, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
    tb(slide, Inches(8.9), Inches(2.35), Inches(0.7), Inches(0.6), "=", size=28, bold=True,
       color=SOFT, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(9.7), Inches(1.85), Inches(2.75), Inches(1.8), L_AMBER)
    tb(slide, Inches(9.7), Inches(2.0), Inches(2.75), Inches(0.9), "CAT", size=42, bold=True,
       color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(9.7), Inches(2.95), Inches(2.75), Inches(0.5), "🐱", size=24,
       align=PP_ALIGN.CENTER)
    ido_wedo_youdo(slide, Inches(4.3), FOREST,
                   ("I blend it", "We blend it", "YOU blend it!"))
    tb(slide, Inches(0.55), Inches(6.05), Inches(9.0), Inches(0.42),
       "🐢 Slow:  /k/ ... /a/ ... /t/       🐇 Fast:  CAT!", size=17, bold=True, color=CORAL)
    notes(slide,
          "Watch what happens when sounds hold hands. /k/ ... /a/ ... /t/. Now faster: "
          "/k/-/a/-/t/... CAT! The sounds became a word.",
          "Say the sounds slowly first with a pause, then speed up and blend. Sweep your "
          "finger under the letters as you blend so the student sees left-to-right movement.",
          "The student blends /k/ /a/ /t/ into CAT, first with you and then alone.",
          "Blend just the first two sounds: /k/ /a/ = 'ca'. Then add /t/. "
          "Two sounds are much easier than three.",
          "You blend it and let the student only say the final word.",
          "Ask them to blend it with their eyes closed, listening only.",
          "You made the sounds turn into a word. That is exactly what reading is!",
          "Note whether the student can hold sounds in memory long enough to blend them.",
          "15–25 min")


def s08_sound_train():
    slide, n = new_slide("🚂 Game: Sound Train", "GAME", "15–25 min", "Sound train", SKY)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Each car is a sound. Push them together and the train says a word!",
       size=16, bold=True, color=CORAL)
    trains = CVC_WORDS[:4]
    for i, (word, sounds, emoji, color, light) in enumerate(trains):
        top = Inches(1.9 + i * 1.28)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.12), WHITE)
        tb(slide, Inches(0.7), top + Inches(0.28), Inches(0.7), Inches(0.55), "🚂", size=22,
           align=PP_ALIGN.CENTER)
        for j, s in enumerate(sounds):
            left = Inches(1.6 + j * 1.55)
            add_round(slide, left, top + Inches(0.2), Inches(1.35), Inches(0.72), light)
            tb(slide, left, top + Inches(0.3), Inches(1.35), Inches(0.5), s, size=24, bold=True,
               color=color, align=PP_ALIGN.CENTER, font="Arial Black")
            if j < len(sounds) - 1:
                tb(slide, left + Inches(1.36), top + Inches(0.32), Inches(0.2), Inches(0.45),
                   "›", size=20, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, Inches(6.5), top + Inches(0.3), Inches(0.8), Inches(0.5), "➡️", size=18,
           align=PP_ALIGN.CENTER)
        add_round(slide, Inches(7.5), top + Inches(0.2), Inches(2.6), Inches(0.72), L_AMBER)
        tb(slide, Inches(7.5), top + Inches(0.28), Inches(2.6), Inches(0.55), word, size=26,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, Inches(10.4), top + Inches(0.26), Inches(0.9), Inches(0.6), emoji, size=22,
           align=PP_ALIGN.CENTER)
        add_round(slide, Inches(11.5), top + Inches(0.3), Inches(1.15), Inches(0.5), L_TEAL)
        tb(slide, Inches(11.5), top + Inches(0.37), Inches(1.15), Inches(0.36), "⭐", size=14,
           align=PP_ALIGN.CENTER)
    notes(slide,
          "All aboard the Sound Train! Each car holds one sound. Say each car... now push "
          "them together... CAT! You get a star for every train you finish.",
          "Point to each car as the student says the sound, then sweep your finger fast "
          "across all three while they blend. Do all four words unless attention drops.",
          "CAT  •  SUN  •  DOG  •  MAP",
          "Cover the last car with your hand. Blend the first two, then reveal the last sound.",
          "Do one word only, and blend it together with the student every time.",
          "Ask the student to build a train for their own name.",
          "Choo choo! That train read perfectly. Another star for you!",
          "Note blending speed and whether the student needs the visual cars or can blend by ear.",
          "15–25 min")


def s09_cvc_practice():
    slide, n = new_slide("🧩 CVC Word Practice", "PHONICS", "15–25 min", "CVC words", TEAL)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Sound it out, then say the whole word.", size=16, bold=True, color=CORAL)
    for i, (word, sounds, emoji, color, light) in enumerate(CVC_WORDS):
        col, row = i % 4, i // 4
        left = Inches(0.45 + col * 3.18)
        top = Inches(1.9 + row * 2.45)
        add_round(slide, left, top, Inches(3.0), Inches(2.25), light)
        tb(slide, left + Inches(0.1), top + Inches(0.15), Inches(2.8), Inches(0.6), emoji,
           size=26, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.1), top + Inches(0.8), Inches(2.8), Inches(0.7), word,
           size=34, bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        add_round(slide, left + Inches(0.25), top + Inches(1.58), Inches(2.5), Inches(0.5), WHITE)
        tb(slide, left + Inches(0.25), top + Inches(1.65), Inches(2.5), Inches(0.38),
           " - ".join(sounds), size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(slide,
          "Let's read these together. Look at the picture first — that is a clue. "
          "Now sound it out: /s/ /u/ /n/... SUN!",
          "Work down the list. Say the sounds with the student, then let them say the whole "
          "word alone. Choose only 4 or 5 words if the student is tiring.",
          "CAT, SUN, DOG, MAP, PIG, HAT, RUN, BIG — each blended correctly.",
          "Cover the word and show only the picture, then reveal one letter at a time.",
          "Do the four words with pictures the student already knows best.",
          "Ask which two words rhyme (CAT and HAT).",
          "Look how many words you just read! You are reading real words.",
          "Note which short vowels are secure and which need review next time.",
          "15–25 min")


def s10_feed_bear():
    slide, n = new_slide("🐻 Game: Feed the Bear", "GAME", "25–35 min", "Feed the bear", BROWN)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Ben is hungry! Find the right word and feed it to him.",
       size=16, bold=True, color=CORAL)
    for i, (emoji, target, options) in enumerate(FEED_BEAR):
        top = Inches(1.9 + i * 1.28)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.12), WHITE)
        tb(slide, Inches(0.7), top + Inches(0.26), Inches(0.8), Inches(0.6), "🐻", size=22,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.6), top + Inches(0.3), Inches(0.8), Inches(0.55), emoji, size=20,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.5), top + Inches(0.34), Inches(2.5), Inches(0.45),
           f"Find  {target}", size=18, bold=True, color=BROWN)
        for j, opt in enumerate(options):
            left = Inches(5.3 + j * 2.55)
            add_round(slide, left, top + Inches(0.2), Inches(2.3), Inches(0.72), L_BROWN)
            tb(slide, left, top + Inches(0.3), Inches(2.3), Inches(0.5), opt, size=22,
               bold=True, color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
    notes(slide,
          "Ben the Bear is hungry, and he only eats the right word! Look at the picture, "
          "then find the word that matches. Point to it and read it out loud.",
          "One row at a time. The student points and reads the word before you move on. "
          "Make a happy munching sound when they get it right.",
          "CAT  •  SUN  •  DOG  •  BIG",
          "Ask what sound the word starts with, then look for that letter in the choices.",
          "Cover one wrong option so there are only two words to choose from.",
          "Ask them to use the fed word in a short sentence.",
          "Nom nom nom! Ben loved that word. Great reading!",
          "Note whether the student decodes the word or guesses from the picture alone.",
          "25–35 min")


def s11_word_hunt():
    slide, n = new_slide("🎯 Game: Word Hunt", "GAME", "25–35 min", "Word hunt", AMBER)
    tb(slide, Inches(0.45), Inches(1.38), Inches(9.0), Inches(0.42),
       "Point to the word I say — then read it!", size=16, bold=True, color=CORAL)
    for r, row in enumerate(HUNT_GRID):
        for c, word in enumerate(row):
            left = Inches(0.5 + c * 1.92)
            top = Inches(1.9 + r * 1.32)
            add_round(slide, left, top, Inches(1.78), Inches(1.15), WHITE)
            tb(slide, left, top + Inches(0.3), Inches(1.78), Inches(0.6), word, size=24,
               bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(10.25), Inches(1.9), Inches(2.6), Inches(3.79), L_AMBER)
    tb(slide, Inches(10.45), Inches(2.08), Inches(2.2), Inches(0.42), "Find these:", size=16,
       bold=True, color=NAVY)
    for i, (word, _count) in enumerate(HUNT_ROUNDS):
        tb(slide, Inches(10.55), Inches(2.6 + i * 0.48), Inches(2.1), Inches(0.42),
           f"{i + 1}.  {word}", size=16, bold=True, color=DARK)
    star_badge(slide, Inches(10.25), Inches(6.05), "⭐ +1 Star")
    notes(slide,
          "Can you find the word SUN? Point to it. Now read it in a big voice. "
          "Can you find another SUN?",
          "Call one word at a time. Every word appears exactly twice, so ask 'Can you find "
          "the other one?' after each success. Six rounds.",
          "Each of CAT, SUN, DOG, BIG, RUN and HAT appears 2 times in the grid.",
          "Give the first sound as a clue, then cover two rows so there is less to scan.",
          "Search one row at a time instead of the whole grid.",
          "Time them: can they find both copies in ten seconds?",
          "Your eyes are getting so quick at spotting words!",
          "Note visual tracking and whether the student reads or matches word shapes.",
          "25–35 min")


def s12_sight_words():
    slide, n = new_slide("👀 Sight Words — Look and Say", "SIGHT WORDS", "35–45 min",
                         "Sight words", PLUM)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "These words we do NOT sound out. We just know them by sight!",
       size=16, bold=True, color=CORAL)
    top = Inches(1.9)
    for label, words, color, light in SIGHT_SETS:
        height = Inches(1.5)
        add_round(slide, Inches(0.5), top, Inches(12.35), height, light)
        add_round(slide, Inches(0.72), top + Inches(0.45), Inches(1.5), Inches(0.55), color)
        tb(slide, Inches(0.72), top + Inches(0.54), Inches(1.5), Inches(0.4), label, size=13,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, word in enumerate(words):
            left = Inches(2.6 + j * 2.55)
            add_round(slide, left, top + Inches(0.3), Inches(2.3), Inches(0.88), WHITE)
            tb(slide, left, top + Inches(0.45), Inches(2.3), Inches(0.58), word, size=30,
               bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        top = top + Inches(1.62)
    notes(slide,
          "These are special words. We do not sound them out — we just learn them by sight, "
          "like recognizing a friend's face. This one says THE. Your turn: THE.",
          "Cover Rounds 2 and 3 with a sheet of paper. Teach only Round 1 first: point, say, "
          "student repeats, three times each. Reveal the next round only when Round 1 is solid.",
          "The student reads I, see, the — then can, my, is — then we, go, and, you.",
          "Trace the word in the air together while saying it. Say it three times in a row.",
          "Teach only two words this lesson. Two words truly known beats ten half-known.",
          "Ask them to find the word in a book or on the screen.",
          "You just KNEW that word without sounding it out. That is real sight reading!",
          "Note which sight words are instant and which need more repetition. "
          "Reuse the weak ones in the sentence section later.",
          "35–45 min")


def s13_word_detective():
    slide, n = new_slide("🔍 Game: Word Detective", "GAME", "35–45 min", "Word detective", TEAL)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Careful — the other words look almost the same!", size=16, bold=True, color=CORAL)
    for i, (target, options) in enumerate(DETECTIVE):
        top = Inches(1.9 + i * 1.3)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.14), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.3), Inches(0.55), Inches(0.55), TEAL)
        tb(slide, Inches(0.72), top + Inches(0.37), Inches(0.55), Inches(0.42), str(i + 1),
           size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.5), top + Inches(0.33), Inches(2.5), Inches(0.5),
           f"Find  {target.upper()}", size=19, bold=True, color=TEAL)
        for j, opt in enumerate(options):
            left = Inches(4.3 + j * 2.15)
            add_round(slide, left, top + Inches(0.22), Inches(1.95), Inches(0.7), L_TEAL)
            tb(slide, left, top + Inches(0.32), Inches(1.95), Inches(0.5), opt, size=21,
               bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(slide,
          "Detective time! Can you find the word THE? Look carefully — some of these look "
          "very similar. Point to it when you find it.",
          "One row at a time. Read the four choices aloud yourself if the student needs it, "
          "then let them point. Confirm by having them say the word.",
          "1. the  •  2. can  •  3. my  •  4. see",
          "Say the target word again slowly, then cover two wrong choices with your fingers.",
          "Use two choices instead of four.",
          "Ask what the other three words say.",
          "You spotted it! Those tricky words did not fool you.",
          "Note confusion between visually similar words such as can/cat and see/saw.",
          "35–45 min")


def s14_memory_match():
    slide, n = new_slide("🧠 Game: Memory Match", "GAME", "35–45 min", "Memory match", CORAL)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Match each picture to its word. Say the word out loud!",
       size=16, bold=True, color=CORAL)
    for i, (emoji, _word) in enumerate(MEMORY_PAIRS):
        left = Inches(0.75 + i * 3.05)
        add_round(slide, left, Inches(1.95), Inches(2.8), Inches(1.95), L_CORAL)
        tb(slide, left, Inches(2.3), Inches(2.8), Inches(1.1), emoji, size=44,
           align=PP_ALIGN.CENTER)
    shuffled = ["DOG", "BEAR", "CAT", "SUN"]
    for i, word in enumerate(shuffled):
        left = Inches(0.75 + i * 3.05)
        add_round(slide, left, Inches(4.35), Inches(2.8), Inches(1.15), WHITE)
        tb(slide, left, Inches(4.6), Inches(2.8), Inches(0.65), word, size=28, bold=True,
           color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(0.75), Inches(5.75), Inches(9.0), Inches(0.45),
       "Draw a line in the air from each picture to its word! ✏️", size=16, bold=True,
       color=TEAL)
    notes(slide,
          "Look at the pictures on top and the words on the bottom. They got mixed up! "
          "Point to the cat, then point to the word CAT.",
          "Let the student match one pair at a time. Have them say the word after each match. "
          "Keep it quick and playful — under three minutes.",
          "🐱 = CAT (3rd word)  •  ☀️ = SUN (4th word)  •  🐶 = DOG (1st word)  •  "
          "🐻 = BEAR (2nd word)",
          "Ask what sound the picture word starts with, then find that letter in the words.",
          "Match only two pairs instead of four.",
          "Cover the pictures and ask them to read all four words alone.",
          "Four out of four! Your word memory is getting strong.",
          "Note whether the student decodes the word or relies on first-letter matching.",
          "35–45 min")


def s15_brain_break():
    slide, n = new_slide("🧠 Brain Break: Simon Says — Reading Edition", "BREAK", "45–50 min",
                         "Brain break", GOLD)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Stay in your seat — just point and touch! 👉", size=16, bold=True, color=CORAL)
    for i, item in enumerate(SIMON_BOARD):
        col, row = i % 3, i // 3
        left = Inches(0.6 + col * 3.05)
        top = Inches(1.9 + row * 1.5)
        add_round(slide, left, top, Inches(2.85), Inches(1.32), WHITE)
        tb(slide, left, top + Inches(0.35), Inches(2.85), Inches(0.62), item, size=30,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
    add_round(slide, Inches(9.9), Inches(1.9), Inches(2.95), Inches(2.92), L_AMBER)
    tb(slide, Inches(10.1), Inches(2.08), Inches(2.55), Inches(0.42), "Simon says...", size=16,
       bold=True, color=NAVY)
    for i, cmd in enumerate(SIMON_SAYS[:4]):
        tb(slide, Inches(10.15), Inches(2.58 + i * 0.55), Inches(2.5), Inches(0.5), cmd,
           size=12, color=DARK)
    tb(slide, Inches(0.6), Inches(5.1), Inches(9.0), Inches(0.45),
       "🎵 Fast and fun — no reading out loud needed!", size=16, bold=True, color=TEAL)
    notes(slide,
          "Let's give our brains a little break — but we are still playing with words! "
          "Simon says: point to the word CAT. Simon says: show me the letter B.",
          "Fast pace, roughly five seconds per command. The student only points or touches — "
          "no reading aloud required. This is a rest, so keep it light and silly.",
          "The student points to the correct word or letter on the board.",
          "Point to it together the first time, then let them try the next one alone.",
          "Give only three commands and always name the word clearly.",
          "Add the classic twist: if you do not say 'Simon says', they should not move.",
          "You are quick! Great pointing. Ready for the next part?",
          "Note recognition speed without decoding pressure, and re-engagement after the break.",
          "45–50 min")


def s16_sentence_intro():
    slide, n = new_slide("📖 Now We Read Sentences!", "READING", "50–60 min", "Sentences", SKY)
    sentence, emoji, color, light = SENTENCES[0]
    add_round(slide, Inches(0.5), Inches(1.5), Inches(12.35), Inches(2.55), WHITE)
    tb(slide, Inches(0.85), Inches(1.95), Inches(1.5), Inches(1.7), emoji, size=54,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(2.6), Inches(2.15), Inches(9.9), Inches(1.3), sentence, size=54, bold=True,
       color=NAVY, align=PP_ALIGN.CENTER, font="Arial Black")
    ido_wedo_youdo(slide, Inches(4.35), SKY)
    tb(slide, Inches(0.55), Inches(6.1), Inches(9.2), Inches(0.42),
       "👉 Point under each word as you read it.", size=17, bold=True, color=CORAL)
    notes(slide,
          "Now we put our words together into a sentence. Watch my finger. "
          "I... see... a... cat. Now let's read it together.",
          "Point under each word as you read, one word at a time. Then read together. "
          "Only then invite the student to try alone — and it is fine if they are not ready.",
          "I see a cat.",
          "Read the first three words and let the student finish with just 'cat'. "
          "Build up backwards from the last word.",
          "Read the sentence together every time, with no independent turn yet.",
          "Ask them to read it with a happy voice, then a sleepy voice.",
          "You read a whole sentence! That is a big jump from words.",
          "Note left-to-right tracking, whether they point, and where they pause.",
          "50–60 min")


def _sentence_pair(idx_a, idx_b, title, timing):
    slide, n = new_slide(title, "READING", timing, "Sentences", TEAL)
    for slot, idx in enumerate((idx_a, idx_b)):
        sentence, emoji, color, light = SENTENCES[idx]
        top = Inches(1.6 + slot * 2.6)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(2.35), light)
        tb(slide, Inches(0.85), top + Inches(0.6), Inches(1.4), Inches(1.2), emoji, size=44,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.5), top + Inches(0.55), Inches(8.0), Inches(1.15), sentence,
           size=42, bold=True, color=NAVY, font="Arial Black")
        add_round(slide, Inches(10.6), top + Inches(0.75), Inches(2.0), Inches(0.75), color)
        tb(slide, Inches(10.6), top + Inches(0.9), Inches(2.0), Inches(0.5), "YOU read ⭐",
           size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return slide, n


def s17_read_with_me():
    slide, n = _sentence_pair(1, 2, "🤝 Read With Me", "50–60 min")
    notes(slide,
          "My turn first, then we read together, then it is your turn. "
          "The cat is big. Now with me: The cat is big.",
          "Use the full I DO / WE DO / YOU DO cycle on each sentence before moving to the "
          "next. Point under every word. Do not correct mid-sentence — wait until the end.",
          "The cat is big.  •  I can run.",
          "Cover all but the first word, then reveal one word at a time as they read.",
          "Read it together twice and skip the independent turn.",
          "Ask them to change one word: 'The DOG is big.'",
          "Take your time — there is no rush at all. That was excellent reading!",
          "Note self-correction, sight word automaticity, and reading pace.",
          "50–60 min")


def s18_sentence_practice():
    slide, n = _sentence_pair(3, 4, "⭐ Your Turn to Read", "50–60 min")
    notes(slide,
          "These two are yours. Take your time. If a word is tricky, we will sound it out "
          "together — that is what readers do.",
          "Let the student lead. Stay silent for five seconds before offering help. "
          "Praise first, then work on at most one word.",
          "We see the sun.  •  My dog can run.",
          "Use the support ladder: first sound, then vowel, then blend, then say the word "
          "and have the student repeat it successfully.",
          "Read the sentence first and let them echo it back.",
          "Ask them to make up a new sentence using 'my'.",
          "You read that all by yourself! I am so proud of you.",
          "Note independence level and which sight words still need support.",
          "50–60 min")


def _build_slide(items, title, timing, start_index):
    slide, n = new_slide(title, "GAME", timing, "Build it", PLUM)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "The words got mixed up! Put them back in order, then read it.",
       size=16, bold=True, color=CORAL)
    for i, (tiles, _answer, emoji) in enumerate(items):
        top = Inches(1.9 + i * 1.65)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.45), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.42), Inches(0.6), Inches(0.6), PLUM)
        tb(slide, Inches(0.72), top + Inches(0.5), Inches(0.6), Inches(0.45),
           str(start_index + i), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.5), top + Inches(0.45), Inches(0.75), Inches(0.55), emoji, size=22,
           align=PP_ALIGN.CENTER)
        for j, tile in enumerate(tiles):
            left = Inches(2.4 + j * 1.95)
            add_round(slide, left, top + Inches(0.35), Inches(1.78), Inches(0.75), L_PLUM)
            tb(slide, left, top + Inches(0.48), Inches(1.78), Inches(0.5), tile, size=20,
               bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    return slide, n


def s19_build_a():
    slide, n = _build_slide(SCRAMBLES[:3], "🧩 Game: Build the Sentence", "60–68 min", 1)
    notes(slide,
          "Uh oh, these words fell out of order! Which word goes first? "
          "Here is a clue: it has a capital letter.",
          "Have the student say the order out loud, or point to the tiles in sequence. "
          "Then they read the finished sentence. One row at a time.",
          "1. I see a cat.   2. The cat is big.   3. I can run.",
          "Point out the capital letter for the first word and the period for the last word. "
          "That solves both ends immediately.",
          "Give the first word yourself, then let them order the remaining two or three.",
          "Ask them to build a brand new sentence using the same tiles.",
          "You put the whole sentence back together. That is smart thinking!",
          "Note understanding of sentence structure, capitals and end punctuation.",
          "60–68 min")


def s20_build_b():
    slide, n = _build_slide(SCRAMBLES[3:], "🧩 Build the Sentence — Round 2", "60–68 min", 4)
    notes(slide,
          "Three more! These have a few extra words, so read all the tiles first before "
          "you choose.",
          "Same routine. Sentence 6 introduces Ben from the story coming next, so pause and "
          "say 'You will meet Ben in a minute!'",
          "4. We see the sun.   5. My dog can run.   6. Ben has a red ball.",
          "Read all the tiles aloud yourself first so the student hears every word.",
          "Do sentence 4 only, and build it together.",
          "Ask what Ben's story might be about, based on sentence 6.",
          "Six sentences built! You are ready for a real story now.",
          "Note stamina late in the game and whether accuracy is holding up.",
          "60–68 min")


def s21_story_words():
    slide, n = new_slide("📚 Story Words — Meet Ben the Bear", "VOCAB", "68–78 min",
                         "Story words", FOREST)
    tb(slide, Inches(0.45), Inches(1.36), Inches(11.5), Inches(0.4),
       "These words are in Ben's story. Let's meet them first!", size=16, bold=True, color=CORAL)
    for i, (emoji, word, sounds, sentence) in enumerate(VOCAB):
        col, row = i % 4, i // 4
        left = Inches(0.45 + col * 3.18)
        top = Inches(1.85 + row * 2.5)
        add_round(slide, left, top, Inches(3.0), Inches(2.3), L_FOREST)
        tb(slide, left + Inches(0.1), top + Inches(0.12), Inches(2.8), Inches(0.6), emoji,
           size=26, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.1), top + Inches(0.75), Inches(2.8), Inches(0.6), word,
           size=28, bold=True, color=FOREST, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left + Inches(0.1), top + Inches(1.35), Inches(2.8), Inches(0.35), sounds,
           size=13, color=SOFT, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.15), top + Inches(1.72), Inches(2.7), Inches(0.45), WHITE)
        tb(slide, left + Inches(0.15), top + Inches(1.79), Inches(2.7), Inches(0.35), sentence,
           size=12, color=DARK, align=PP_ALIGN.CENTER)
    notes(slide,
          "Before we read Ben's story, let's meet the words we will see. This one is BEAR. "
          "Ben is a bear. Your turn: BEAR.",
          "Say the word, break it into the sounds shown, then read the tiny sentence. "
          "Student repeats. Roughly 20 seconds per word. Skip 'happy' and 'little' if time "
          "is short — they can be handled inside the story.",
          "The student repeats each word and can read at least four of them independently.",
          "Cover the word and show only the picture, then reveal the word and blend together.",
          "Preview only four words: bear, ball, tree, red. Those carry the story.",
          "Ask them to predict what happens to Ben and the ball.",
          "You already know so many of these words. The story is going to be easy for you!",
          "Note which two-syllable words (happy, little) need extra support.",
          "68–78 min")


def _story_slide(index, timing):
    part, emoji, lines, color, light = STORY[index]
    slide, n = new_slide(f"📚 The Little Bear — {part}", "STORY", timing, part, color)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(8.9), Inches(4.95), WHITE)
    for i, line in enumerate(lines):
        tb(slide, Inches(0.95), Inches(1.95 + i * 1.12), Inches(8.1), Inches(0.95), line,
           size=32, bold=True, color=NAVY)
    add_round(slide, Inches(9.65), Inches(1.5), Inches(3.2), Inches(4.95), light)
    tb(slide, Inches(9.65), Inches(2.1), Inches(3.2), Inches(1.4), emoji, size=64,
       align=PP_ALIGN.CENTER)
    steps = ["1️⃣  I read it", "2️⃣  We read it", "3️⃣  You pick ONE line"]
    for i, step in enumerate(steps):
        top = Inches(3.95 + i * 0.72)
        add_round(slide, Inches(9.9), top, Inches(2.7), Inches(0.58), WHITE)
        tb(slide, Inches(9.9), top + Inches(0.11), Inches(2.7), Inches(0.4), step, size=14,
           bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    return slide, n


def s22_story_1():
    slide, n = _story_slide(0, "68–78 min")
    notes(slide,
          "Meet Ben. Ben is a little bear, and he has a red ball. I will read it first, "
          "then we read it together, then you pick just one line to read by yourself.",
          "Read with warmth and expression. Point under each word. Let the student CHOOSE "
          "which single line to read — choice reduces anxiety a lot.",
          "Ben is a little bear. Ben has a red ball. Ben can run and play. Ben is happy.",
          "Offer the shortest line: 'Ben is happy.' Sound out any tricky word together.",
          "Read every line together with no solo turn.",
          "Ask what color Ben's ball is, then ask them to find that word.",
          "You picked a line and read it. That is being a real reader!",
          "Note willingness to attempt independent reading and which line they choose.",
          "68–78 min")


def s23_story_2():
    slide, n = _story_slide(1, "68–78 min")
    notes(slide,
          "Oh no! Something happened to Ben's ball. Let's find out. "
          "One day, the ball is gone...",
          "Build a little suspense. After reading, ask 'Where do you think the ball is?' "
          "before moving on — prediction keeps engagement high.",
          "One day, the ball is gone. Ben can not see the ball. Ben looks and looks. "
          "Ben sees a big tree.",
          "The word 'gone' may be new — sound it out together and reread the whole line.",
          "Read all four lines yourself and ask the student only to point along.",
          "Ask them to predict where the ball is before turning the page.",
          "Great predicting! Good readers always think about what comes next.",
          "Note comprehension of the problem in the story and prediction ability.",
          "68–78 min")


def s24_story_3():
    slide, n = _story_slide(2, "68–78 min")
    notes(slide,
          "There it is! The ball is under the tree — but look who is sitting on it. A cat!",
          "Use a different voice for Ben's line, 'My ball!' Ask the student to try that line "
          "with an excited voice — dialogue is motivating.",
          "The ball is under the tree. A cat sits on the ball. \"My ball!\" says Ben. "
          "The cat runs to Ben.",
          "Explain that quotation marks mean somebody is talking, then read that line together.",
          "Have them read only the two-word part: 'My ball!'",
          "Ask how Ben feels right now and why.",
          "I loved how you used Ben's voice. That was excellent reading!",
          "Note expression, use of punctuation, and engagement with dialogue.",
          "68–78 min")


def s25_story_4():
    slide, n = _story_slide(3, "68–78 min")
    notes(slide,
          "And here is the happy ending. Ben gets his ball back, and now he has a new friend.",
          "Celebrate finishing the whole story. Then ask the student to retell it in their "
          "own words in two or three sentences.",
          "Ben gets the red ball. Ben and the cat play. The sun is big and hot. Ben is happy.",
          "Retell it together, one sentence per picture, using the four story slides as prompts.",
          "Ask only 'Is Ben happy or sad at the end?'",
          "Ask them to invent one more sentence for the story.",
          "You read a whole story today! Not one word — a whole story.",
          "Note overall comprehension and whether they can retell the sequence.",
          "68–78 min")


def s26_story_challenge():
    slide, n = new_slide("⭐ Story Reading Challenge", "CHALLENGE", "68–78 min",
                         "Story challenge", AMBER)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Pick a line from Ben's story and read it. You choose which one!",
       size=16, bold=True, color=CORAL)
    picks = [("⭐", "Ben is happy.", TEAL, L_TEAL),
             ("⭐⭐", "Ben has a red ball.", AMBER, L_AMBER),
             ("⭐⭐⭐", "The ball is under the tree.", CORAL, L_CORAL)]
    for i, (stars, line, color, light) in enumerate(picks):
        top = Inches(1.95 + i * 1.62)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.42), light)
        add_round(slide, Inches(0.75), top + Inches(0.42), Inches(1.7), Inches(0.58), color)
        tb(slide, Inches(0.75), top + Inches(0.5), Inches(1.7), Inches(0.42), stars, size=15,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.8), top + Inches(0.36), Inches(9.5), Inches(0.75), line, size=32,
           bold=True, color=NAVY, font="Arial Black")
    tb(slide, Inches(0.5), Inches(6.85), Inches(9.0), Inches(0.3), "", size=10)
    notes(slide,
          "You choose! One star, two stars, or three. Any one you pick is a win. "
          "Which one do you want to read?",
          "Let the student choose freely. If they pick three stars, support them fully so "
          "they succeed. Reading choice builds ownership and confidence.",
          "Ben is happy.  •  Ben has a red ball.  •  The ball is under the tree.",
          "Use the ladder: first sound, vowel, blend, then give the word and have them "
          "repeat it successfully. Always end on a success.",
          "Read the one-star line together, then let them read it alone.",
          "Invite them to try all three lines.",
          "You chose a hard one and you did it. That took courage!",
          "Note risk-taking: does the student choose an easy or a stretch option?",
          "68–78 min")


def _comprehension_slide(items, title, timing, start_index):
    slide, n = new_slide(title, "THINK", timing, "Comprehension", PLUM)
    for i, (question, options, _answer) in enumerate(items):
        top = Inches(1.55 + i * 1.78)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.6), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.5), Inches(0.6), Inches(0.6), PLUM)
        tb(slide, Inches(0.72), top + Inches(0.58), Inches(0.6), Inches(0.45),
           str(start_index + i), size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.55), top + Inches(0.52), Inches(3.9), Inches(0.6), question,
           size=18, bold=True, color=NAVY)
        for j, (emoji, label) in enumerate(options):
            left = Inches(5.7 + j * 2.4)
            add_round(slide, left, top + Inches(0.28), Inches(2.2), Inches(1.05), L_PLUM)
            tb(slide, left, top + Inches(0.38), Inches(2.2), Inches(0.45), emoji, size=20,
               align=PP_ALIGN.CENTER)
            tb(slide, left + Inches(0.1), top + Inches(0.86), Inches(2.0), Inches(0.38), label,
               size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    return slide, n


def s27_comprehension():
    slide, n = _comprehension_slide(COMPREHENSION[:3], "🧩 Story Questions", "78–84 min", 1)
    notes(slide,
          "Let's think about Ben's story. Who is the story about? Point to your answer.",
          "Read the question and all three choices aloud. The student points — they do not "
          "need to read the options. One question at a time.",
          "1. Bear   2. Ball   3. Under the tree",
          "Turn back to the story slide and read the line that holds the answer, "
          "then ask again.",
          "Offer two choices instead of three.",
          "Ask them to say the answer in a full sentence: 'The story is about a bear.'",
          "That is exactly right! You remembered the story so well.",
          "Note literal recall of who, what and where.",
          "78–84 min")


def s28_picture_game():
    slide, n = new_slide("🖼️ Picture Question Game", "GAME", "78–84 min", "Picture game", TEAL)
    tb(slide, Inches(0.45), Inches(1.34), Inches(11.5), Inches(0.4),
       "I read a sentence — you point to the right picture!", size=16, bold=True, color=CORAL)
    for i, (sentence, options, _answer) in enumerate(PICTURE_MATCH):
        top = Inches(1.8 + i * 1.34)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.18), WHITE)
        tb(slide, Inches(0.8), top + Inches(0.36), Inches(6.2), Inches(0.5), sentence, size=20,
           bold=True, color=NAVY)
        for j, emoji in enumerate(options):
            left = Inches(7.5 + j * 1.8)
            add_round(slide, left, top + Inches(0.22), Inches(1.6), Inches(0.75), L_TEAL)
            tb(slide, left, top + Inches(0.3), Inches(1.6), Inches(0.55), emoji, size=24,
               align=PP_ALIGN.CENTER)
    remaining = COMPREHENSION[3:]
    top = Inches(5.9)
    add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.0), L_AMBER)
    tb(slide, Inches(0.8), top + Inches(0.12), Inches(5.5), Inches(0.4),
       "Two more story questions:", size=14, bold=True, color=NAVY)
    for i, (question, _options, _answer) in enumerate(remaining):
        tb(slide, Inches(0.8 + i * 6.0), top + Inches(0.55), Inches(5.7), Inches(0.38),
           f"{4 + i}. {question}", size=14, bold=True, color=DARK)
    notes(slide,
          "Listen to my sentence and point to the picture that matches. "
          "Ben has a red ball. Which picture?",
          "Read each sentence twice. The student points only. Then ask the two remaining "
          "story questions at the bottom out loud.",
          "🔴 ball  •  🐱 cat  •  🌳 tree.  Question 4: It was gone.  "
          "Question 5: Happy.",
          "Reread the sentence and stress the key noun: 'Ben has a red BALL.'",
          "Show only two pictures to choose between.",
          "Ask them to read the sentence themselves before pointing.",
          "You are listening so carefully. Perfect matching!",
          "Note listening comprehension separate from decoding skill — often much stronger.",
          "78–84 min")


def s29_final_challenge():
    slide, n = new_slide("🎯 Final Reading Challenge", "CHALLENGE", "84–87 min",
                         "Final challenge", CORAL)
    tb(slide, Inches(0.45), Inches(1.36), Inches(11.5), Inches(0.4),
       "Read as many as you can. Every single one counts! 🌟", size=16, bold=True, color=CORAL)
    add_round(slide, Inches(0.5), Inches(1.85), Inches(12.35), Inches(1.75), L_AMBER)
    tb(slide, Inches(0.8), Inches(2.0), Inches(3.0), Inches(0.4), "5 WORDS", size=15,
       bold=True, color=NAVY)
    for i, word in enumerate(FINAL_WORDS):
        left = Inches(0.8 + i * 2.4)
        add_round(slide, left, Inches(2.5), Inches(2.2), Inches(0.85), WHITE)
        tb(slide, left, Inches(2.65), Inches(2.2), Inches(0.6), word, size=28, bold=True,
           color=CORAL, align=PP_ALIGN.CENTER, font="Arial Black")
    add_round(slide, Inches(0.5), Inches(3.8), Inches(12.35), Inches(2.55), L_TEAL)
    tb(slide, Inches(0.8), Inches(3.95), Inches(4.0), Inches(0.4), "3 SENTENCES", size=15,
       bold=True, color=NAVY)
    for i, sentence in enumerate(FINAL_SENTENCES):
        top = Inches(4.45 + i * 0.62)
        add_round(slide, Inches(0.8), top, Inches(11.75), Inches(0.52), WHITE)
        tb(slide, Inches(1.05), top + Inches(0.06), Inches(11.2), Inches(0.4), sentence,
           size=22, bold=True, color=NAVY)
    star_badge(slide, Inches(10.3), Inches(6.45), "⭐ Every one counts!")
    notes(slide,
          "Last challenge! Read as many as you can. There is no pass or fail here — every "
          "word you read is a win. Ready?",
          "Count successes out loud as they go: 'That's two! That's three!' Never mention "
          "misses. If they stall on a word, supply it warmly and move on immediately.",
          "CAT, SUN, BIG, RUN, RED. I see a cat. The sun is big. Ben can run.",
          "Sound out the first word together to get momentum, then step back.",
          "Do the five words only and skip the sentences.",
          "Ask them to read one sentence in a silly voice.",
          "You read four words! Look how much you learned in one class.",
          "This is your key progress record — note exactly how many words and sentences "
          "were read independently, and compare next lesson.",
          "84–87 min")


def s30_recap():
    slide, n = new_slide("🎓 What Did We Learn Today?", "RECAP", "87–90 min", "Recap", FOREST)
    wins = [("🔤", "Letter Sounds", "A B C D E S M R"), ("🎵", "Blending", "C-A-T = CAT"),
            ("🧩", "New Words", "cat, sun, dog, big"), ("👀", "Sight Words", "the, can, my, see"),
            ("📖", "Sentences", "I see a cat."), ("📚", "A Whole Story", "The Little Bear")]
    for i, (icon, title, detail) in enumerate(wins):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.6 + row * 2.45)
        add_round(slide, left, top, Inches(3.9), Inches(2.2), WHITE)
        add_oval(slide, left + Inches(1.5), top + Inches(0.2), Inches(0.9), Inches(0.9), L_FOREST)
        tb(slide, left + Inches(1.5), top + Inches(0.34), Inches(0.9), Inches(0.6), icon,
           size=22, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(1.18), Inches(3.6), Inches(0.45), title,
           size=17, bold=True, color=FOREST, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(1.63), Inches(3.6), Inches(0.42), detail,
           size=13, color=SOFT, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(6.5), Inches(9.5), Inches(0.42),
       "🗣️ What was your favorite game? What word do you remember?", size=16, bold=True,
       color=CORAL)
    notes(slide,
          "Look at everything you did today! Sounds, blending, new words, sight words, "
          "sentences, and a whole story. That is a lot for one class.",
          "Point to each card and let the student say what they remember. Then ask the two "
          "questions at the bottom. Keep it celebratory, about one minute.",
          "The student names a favorite game and recalls at least two words.",
          "Prompt with the first sound: 'We read about a b... bear!'",
          "Ask only 'What was your favorite game today?'",
          "Ask them to teach YOU one thing they learned.",
          "You should be really proud of yourself today. I certainly am.",
          "Note what the student retains without prompting — a good measure of the lesson.",
          "87–90 min")


def s31_star_celebration():
    slide, n = new_slide("⭐ Your Reading Stars", "REWARD", "87–90 min", "Stars", GOLD)
    tb(slide, Inches(0.45), Inches(1.4), Inches(11.5), Inches(0.45),
       "Count your stars out loud with me!", size=18, bold=True, color=CORAL)
    stars = ["Letter Sounds", "Mystery Letter", "Sound Train", "Feed the Bear",
             "Sight Words", "Sentences", "Ben's Story", "Big Challenge"]
    for i, label in enumerate(stars):
        col, row = i % 4, i // 4
        left = Inches(0.5 + col * 3.15)
        top = Inches(2.0 + row * 2.3)
        add_round(slide, left, top, Inches(2.95), Inches(2.05), L_AMBER)
        add_oval(slide, left + Inches(1.05), top + Inches(0.22), Inches(0.85), Inches(0.85), GOLD)
        tb(slide, left + Inches(1.05), top + Inches(0.34), Inches(0.85), Inches(0.6), "⭐",
           size=22, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.22), Inches(2.71), Inches(0.65), label,
           size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(slide,
          "Look at all your stars! Let's count them together. One, two, three... "
          "You earned every single one.",
          "Count aloud together, touching each star. Make it a genuine celebration — this "
          "is the emotional high point of the lesson.",
          "The student counts to eight along with you.",
          "Count for them and let them just touch each star.",
          "Count only the stars for activities they clearly enjoyed.",
          "Ask which star they worked hardest for.",
          "Eight stars! That is a full adventure completed.",
          "Note pride and engagement — key indicators for a struggling reader's motivation.",
          "87–90 min")


def s32_badge():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.7, 0.8, SKY), (11.9, 0.85, CORAL), (0.9, 5.8, TEAL), (11.85, 5.75, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(1.05), Inches(12), Inches(0.8),
       "🏆 READING EXPLORER 🏆", size=40, bold=True, color=GOLD, align=PP_ALIGN.CENTER,
       font="Georgia")
    add_oval(slide, Inches(5.42), Inches(2.05), Inches(2.5), Inches(2.5), GOLD)
    tb(slide, Inches(5.42), Inches(2.6), Inches(2.5), Inches(1.3), "🐻", size=58,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(4.8), Inches(12), Inches(0.6), "Great Job!", size=34,
       bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Georgia")
    tb(slide, Inches(0.7), Inches(5.45), Inches(12), Inches(0.55),
       "You are becoming a stronger reader!", size=22, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(6.05), Inches(12), Inches(0.5),
       "Sounds ⭐ Words ⭐ Sentences ⭐ A Whole Story", size=16, color=L_SKY,
       align=PP_ALIGN.CENTER)
    footer(slide, n, "87–90 min", "Badge")
    fade(slide)
    notes(slide,
          "This badge is for you. You are officially a Reading Explorer. You read sounds, "
          "words, sentences and a whole story today, and Ben the Bear says thank you!",
          "Screenshot this slide and send it to the parent after class. End on maximum "
          "warmth and energy.",
          "The student smiles, says thank you, or asks about the next class.",
          "If they seem shy about praise, name one concrete fact: 'You read eight words "
          "all by yourself.'",
          "Simply say 'You did it!' and give a high five.",
          "Ask what adventure Ben should go on next time.",
          "I am so proud of you. See you next time, Reading Explorer!",
          "Note the overall emotional tone at the end — did the student leave feeling capable?",
          "87–90 min")


def s33_homework():
    slide, n = new_slide("🏡 Fun Practice at Home (Optional)", "HOME", "", "Homework", SKY)
    tb(slide, Inches(0.45), Inches(1.38), Inches(11.5), Inches(0.42),
       "Pick just ONE. Ten minutes is plenty!", size=16, bold=True, color=CORAL)
    for i, (icon, title, detail, mins) in enumerate(HOMEWORK):
        col, row = i % 2, i // 2
        left = Inches(0.5 + col * 6.25)
        top = Inches(1.9 + row * 2.4)
        add_round(slide, left, top, Inches(5.95), Inches(2.2), WHITE)
        add_oval(slide, left + Inches(0.28), top + Inches(0.6), Inches(1.0), Inches(1.0), L_SKY)
        tb(slide, left + Inches(0.28), top + Inches(0.78), Inches(1.0), Inches(0.65), icon,
           size=26, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(1.55), top + Inches(0.5), Inches(4.2), Inches(0.5), title,
           size=19, bold=True, color=SKY)
        tb(slide, left + Inches(1.55), top + Inches(1.02), Inches(4.2), Inches(0.6), detail,
           size=14, color=DARK)
        add_round(slide, left + Inches(1.55), top + Inches(1.62), Inches(1.5), Inches(0.42),
                  L_AMBER)
        tb(slide, left + Inches(1.55), top + Inches(1.69), Inches(1.5), Inches(0.34), mins,
           size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(slide,
          "This is not homework you have to do — it is just something fun if you want to. "
          "My favorite is reading to a grown-up, because they will be so impressed!",
          "Let the student pick one out loud so they own the choice. Tell the parent that "
          "daily practice should stay under ten minutes for a struggling reader.",
          "The student picks one activity and says when they will do it.",
          "Say clearly: 'Doing none of these is also completely fine.'",
          "Suggest only the drawing task, which carries no reading pressure.",
          "Invite them to bring their drawing to the next class.",
          "Great pick! I cannot wait to hear about it next time.",
          "Note parental support available at home, if the parent is present.",
          "After class")


def s34_assessment():
    slide, n = new_slide("📋 Teacher Assessment — Progress Record", "TEACHER ONLY", "",
                         "Assessment", NAVY)
    tb(slide, Inches(0.45), Inches(1.3), Inches(12.4), Inches(0.35),
       "TEACHER ONLY — do not show this slide to the student. Record progress, not failure.",
       size=13, bold=True, color=CORAL)
    header_y = Inches(1.75)
    add_round(slide, Inches(0.45), header_y, Inches(4.3), Inches(0.48), NAVY)
    tb(slide, Inches(0.65), header_y + Inches(0.08), Inches(4.0), Inches(0.34), "SKILL",
       size=13, bold=True, color=WHITE)
    add_round(slide, Inches(4.9), header_y, Inches(5.4), Inches(0.48), NAVY)
    tb(slide, Inches(4.9), header_y + Inches(0.09), Inches(5.4), Inches(0.32),
       "⭐ 1 Beginning  ·  2 Developing  ·  3 Progressing  ·  4 Strong",
       size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(10.45), header_y, Inches(2.4), Inches(0.48), NAVY)
    tb(slide, Inches(10.65), header_y + Inches(0.08), Inches(2.1), Inches(0.34), "NOTES",
       size=13, bold=True, color=WHITE)
    for i, skill in enumerate(RUBRIC_SKILLS):
        top = header_y + Inches(0.56 + i * 0.52)
        band = WHITE if i % 2 == 0 else L_SKY
        add_round(slide, Inches(0.45), top, Inches(4.3), Inches(0.46), band)
        tb(slide, Inches(0.65), top + Inches(0.08), Inches(4.0), Inches(0.32), skill, size=13,
           bold=True, color=NAVY)
        add_round(slide, Inches(4.9), top, Inches(5.4), Inches(0.46), band)
        tb(slide, Inches(4.9), top + Inches(0.06), Inches(5.4), Inches(0.34),
           "☆      ☆      ☆      ☆", size=14, color=SOFT, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(10.45), top, Inches(2.4), Inches(0.46), band)
    notes(slide,
          "TEACHER SLIDE — do not display to the student.",
          "Complete this within five minutes of finishing, while details are fresh. "
          "Record the Final Reading Challenge score (words and sentences read "
          "independently) in the notes column — that is the clearest progress measure "
          "across lessons.",
          "A struggling Grade 2 reader typically lands at 1–2 stars for sentence reading "
          "and 2–3 for comprehension, since listening comprehension usually outpaces "
          "decoding.",
          "Avoid rating harshly on a first session — it discourages both student and parent.",
          "If unsure between two levels, choose the higher one and note what to watch next.",
          "Add one specific next-step goal per skill for the following lesson.",
          "Frame all parent feedback as strengths first, then one growth area, then the plan.",
          "Compare against the previous lesson's record to show measurable progress.",
          "After class")


def s35_answer_key():
    slide, n = new_slide("🔑 Answer Key", "TEACHER ONLY", "", "Answer key", NAVY)
    tb(slide, Inches(0.45), Inches(1.28), Inches(12.4), Inches(0.35),
       "TEACHER ONLY — hide this slide before presenting.", size=13, bold=True, color=CORAL)
    add_round(slide, Inches(0.45), Inches(1.72), Inches(6.1), Inches(4.95), WHITE)
    tb(slide, Inches(0.7), Inches(1.9), Inches(5.6), Inches(0.4), "🔤 Phonics & Word Games",
       size=15, bold=True, color=CORAL)
    bullets(slide, Inches(0.7), Inches(2.35), Inches(5.6), Inches(4.2), [
        "Mystery Letter: M, B, S, T",
        "First sounds: /b/, /k/, /s/, /m/",
        "Sound Train: CAT, SUN, DOG, MAP",
        "Feed the Bear: CAT, SUN, DOG, BIG",
        "Word Hunt: each word appears 2 times",
        "Word Detective: the, can, my, see",
        "Memory Match: cat=CAT, sun=SUN,",
        "     dog=DOG, bear=BEAR",
    ], size=12, sp=6)
    add_round(slide, Inches(6.75), Inches(1.72), Inches(6.1), Inches(4.95), WHITE)
    tb(slide, Inches(7.0), Inches(1.9), Inches(5.6), Inches(0.4), "📖 Sentences & Story",
       size=15, bold=True, color=FOREST)
    bullets(slide, Inches(7.0), Inches(2.35), Inches(5.6), Inches(4.2), [
        "Build 1: I see a cat.",
        "Build 2: The cat is big.",
        "Build 3: I can run.",
        "Build 4: We see the sun.",
        "Build 5: My dog can run.",
        "Build 6: Ben has a red ball.",
        "Story Qs: Bear · Ball · Under the tree ·",
        "     It was gone · Happy",
        "Picture match: ball · cat · tree",
    ], size=12, sp=6)
    notes(slide,
          "TEACHER SLIDE — full answer key for every game in this lesson.",
          "Hide this slide in PowerPoint (right-click the thumbnail, then Hide Slide) before "
          "presenting, or keep it open on a second screen.",
          "See the slide content for all answers.",
          "Accidentally displaying this slide to the student.",
          "Accept any reasonable phrasing for the spoken comprehension answers.",
          "For a fast finisher, ask them to point to the line in the story that proves "
          "each answer.",
          "Reference only.",
          "Use the Word Hunt and Word Detective results to choose next lesson's review words.",
          "Reference")


BUILDERS = [
    s01_welcome, s02_mission, s03_feelings, s04_letter_sounds, s05_mystery_letter,
    s06_sound_game, s07_blending_intro, s08_sound_train, s09_cvc_practice,
    s10_feed_bear, s11_word_hunt, s12_sight_words, s13_word_detective,
    s14_memory_match, s15_brain_break, s16_sentence_intro, s17_read_with_me,
    s18_sentence_practice, s19_build_a, s20_build_b, s21_story_words,
    s22_story_1, s23_story_2, s24_story_3, s25_story_4, s26_story_challenge,
    s27_comprehension, s28_picture_game, s29_final_challenge, s30_recap,
    s31_star_celebration, s32_badge, s33_homework, s34_assessment, s35_answer_key,
]

for build in BUILDERS:
    build()

assert len(prs.slides) == TOTAL, f"expected {TOTAL} slides, built {len(prs.slides)}"

out = r"C:\Users\bhushaja\Downloads\shaip\Grade2_Bear_Reading_Adventure_90min.pptx"
try:
    prs.save(out)
except PermissionError:
    out = out.replace(".pptx", "_new.pptx")
    prs.save(out)

story_words = sum(len(line.split()) for _p, _e, lines, _c, _l in STORY for line in lines)
missing = [i + 1 for i, s in enumerate(prs.slides)
           if not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip()]
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
print(f"Story word count: {story_words}")
print(f"Slides missing teacher notes: {missing if missing else 'none'}")

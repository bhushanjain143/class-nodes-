"""Grade 2 reading intervention - 90 minutes, 37 slides, no speaker notes.

"The Great Reading Adventure" - the student is a Reading Detective helping Puppy
find a lost backpack. Difficulty climbs through seven levels: sounds, blending,
CVC words, sight words, sentences, story, comprehension.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

INDIGO = RGBColor(0x1E, 0x2A, 0x5A)
SLATE = RGBColor(0x33, 0x41, 0x6B)
TEAL = RGBColor(0x00, 0x9B, 0x95)
SKY = RGBColor(0x2C, 0xA5, 0xE6)
GRASS = RGBColor(0x3D, 0x9B, 0x5A)
AMBER = RGBColor(0xEF, 0xA4, 0x1C)
GOLD = RGBColor(0xFF, 0xC5, 0x30)
CORAL = RGBColor(0xE8, 0x58, 0x4F)
PLUM = RGBColor(0x82, 0x56, 0xC4)
ROSE = RGBColor(0xE0, 0x54, 0x86)
CREAM = RGBColor(0xFF, 0xFC, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x23, 0x2B, 0x38)
SOFT = RGBColor(0x6B, 0x78, 0x8A)
L_SKY = RGBColor(0xE3, 0xF3, 0xFD)
L_TEAL = RGBColor(0xDD, 0xF4, 0xF3)
L_AMBER = RGBColor(0xFF, 0xF3, 0xD9)
L_CORAL = RGBColor(0xFD, 0xEA, 0xE8)
L_PLUM = RGBColor(0xEF, 0xE9, 0xFB)
L_GRASS = RGBColor(0xE5, 0xF3, 0xE9)
L_ROSE = RGBColor(0xFD, 0xE8, 0xEF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TOTAL = 37
_counter = {"n": 0}

# ------------------------------------------------------------------ content

LEVELS = [
    ("1", "Sounds", "🔤", CORAL), ("2", "Blending", "🎵", AMBER),
    ("3", "Words", "🧩", TEAL), ("4", "Sight Words", "💎", PLUM),
    ("5", "Sentences", "📖", SKY), ("6", "Story", "📚", GRASS),
    ("7", "Detective", "🕵️", ROSE),
]

WARMUP = [
    ("🐾", "What is your favorite animal?"),
    ("🎨", "What is your favorite color?"),
    ("⚽", "What do you like to play?"),
    ("👀", "What do you see in the picture?"),
]

SPY_SCENE = [("☀️", 0.45, 0.25), ("🌳", 4.45, 2.05), ("🐶", 0.80, 2.45),
             ("⚽", 2.15, 2.75), ("🐱", 3.35, 2.55)]
SPY_ROUNDS = [("I spy a dog.", "🐶", "DOG"), ("I spy a cat.", "🐱", "CAT"),
              ("I spy the sun.", "☀️", "SUN")]

SOUNDS = [
    ("m", "/m/", "🗺️", "map", ROSE, L_ROSE),
    ("s", "/s/", "☀️", "sun", AMBER, L_AMBER),
    ("t", "/t/", "🌳", "tree", GRASS, L_GRASS),
    ("p", "/p/", "🐷", "pig", CORAL, L_CORAL),
    ("b", "/b/", "🎒", "bag", PLUM, L_PLUM),
    ("c", "/k/", "🐱", "cat", SKY, L_SKY),
    ("d", "/d/", "🐶", "dog", TEAL, L_TEAL),
    ("f", "/f/", "🦊", "fox", CORAL, L_CORAL),
    ("n", "/n/", "👃", "nose", ROSE, L_ROSE),
    ("r", "/r/", "🔴", "red", GRASS, L_GRASS),
]

DETECTIVE_A = [
    ("🐶", "DOG", ["/m/", "/d/", "/s/"]),
    ("🐱", "CAT", ["/k/", "/t/", "/n/"]),
    ("☀️", "SUN", ["/p/", "/f/", "/s/"]),
]
DETECTIVE_B = [
    ("🎒", "BAG", ["/b/", "/d/", "/r/"]),
    ("🦊", "FOX", ["/s/", "/f/", "/m/"]),
    ("🌳", "TREE", ["/t/", "/p/", "/n/"]),
    ("🗺️", "MAP", ["/n/", "/r/", "/m/"]),
]

TRAIN_WORDS = [
    ("CAT", ["C", "A", "T"], ["/k/", "/a/", "/t/"], "🐱", SKY, L_SKY),
    ("MAT", ["M", "A", "T"], ["/m/", "/a/", "/t/"], "🟫", ROSE, L_ROSE),
    ("SUN", ["S", "U", "N"], ["/s/", "/u/", "/n/"], "☀️", AMBER, L_AMBER),
]
BUILD_WORDS = [
    ("PIG", ["P", "I", "G"], ["/p/", "/i/", "/g/"], "🐷", CORAL, L_CORAL),
    ("DOG", ["D", "O", "G"], ["/d/", "/o/", "/g/"], "🐶", TEAL, L_TEAL),
    ("RUN", ["R", "U", "N"], ["/r/", "/u/", "/n/"], "🏃", GRASS, L_GRASS),
]

FAMILIES = [
    ("- at", ["cat", "bat", "mat", "sat"], "🐱", SKY, L_SKY),
    ("- og", ["dog", "log"], "🐶", TEAL, L_TEAL),
    ("- un", ["sun", "run"], "☀️", AMBER, L_AMBER),
    ("- ig", ["pig", "big"], "🐷", CORAL, L_CORAL),
]

MONSTER = [
    ("🐱", "CAT", ["CAT", "DOG", "SUN"]),
    ("🐶", "DOG", ["MAT", "DOG", "PIG"]),
    ("☀️", "SUN", ["SUN", "BAT", "RUN"]),
    ("🐷", "PIG", ["BIG", "PIG", "LOG"]),
    ("🎒", "BAG", ["BAG", "BIG", "BAT"]),
    ("🏃", "RUN", ["SUN", "FUN", "RUN"]),
]

HUNT_GRID = [
    ["CAT", "SUN", "DOG", "BAG", "RUN"],
    ["BIG", "CAT", "MAT", "SUN", "PIG"],
    ["DOG", "RED", "BAG", "RUN", "BIG"],
]
HUNT_TARGETS = ["CAT", "SUN", "DOG", "BAG", "RUN", "BIG"]

READ_AND_MOVE = [
    ("CAT", "clap 2 times", "👏"), ("DOG", "touch your head", "🙌"),
    ("SUN", "stand up", "🧍"), ("PIG", "give yourself a high five", "🙏"),
    ("BAG", "tap the table 3 times", "✋"), ("RUN", "run in place", "🏃"),
]

SIGHT_SETS = [
    ("Group 1", ["I", "am", "the", "is"], TEAL, L_TEAL),
    ("Group 2", ["a", "my", "can"], CORAL, L_CORAL),
    ("Group 3", ["see", "you", "we"], PLUM, L_PLUM),
]

TREASURE = [
    ("the", ["the", "dog", "sun"]),
    ("can", ["cat", "can", "run"]),
    ("I", ["is", "I", "it"]),
    ("my", ["me", "my", "may"]),
    ("see", ["see", "sea", "so"]),
]

ROLL_READ = [("1", "the"), ("2", "can"), ("3", "am"), ("4", "you"), ("5", "we"), ("6", "is")]

SENTENCES = [
    ("I see a dog.", "🐶", SKY, L_SKY),
    ("The dog can run.", "🐶", TEAL, L_TEAL),
    ("The cat is big.", "🐱", CORAL, L_CORAL),
    ("I can run.", "🏃", GRASS, L_GRASS),
    ("The sun is hot.", "☀️", AMBER, L_AMBER),
]

SCRAMBLES = [
    (["dog", "I", "a", "see"], "I see a dog.", "🐶"),
    (["can", "The", "run", "dog"], "The dog can run.", "🐶"),
    (["big", "is", "The", "cat"], "The cat is big.", "🐱"),
    (["run", "I", "can"], "I can run.", "🏃"),
    (["hot", "The", "is", "sun"], "The sun is hot.", "☀️"),
]

PARK_SCENE = [("☀️", 0.40, 0.20), ("🌳", 4.40, 1.95), ("🐶", 0.70, 2.40),
              ("⚽", 2.05, 2.80), ("🐱", 3.20, 2.55), ("👦", 5.75, 2.45)]
PARK_LABELS = ["🐶 dog", "🐱 cat", "⚽ ball", "🌳 tree", "☀️ sun", "👦 boy"]

CLUES = [
    ("The dog is big.", ["🐶", "🐱", "⚽"], "🐶"),
    ("I see a ball.", ["🌳", "⚽", "☀️"], "⚽"),
    ("The boy can run.", ["👦", "🐶", "🌳"], "👦"),
    ("The sun is hot.", ["⚽", "🐱", "☀️"], "☀️"),
]

VOCAB = [
    ("🐶", "dog", "/d/ /o/ /g/"), ("🐱", "cat", "/k/ /a/ /t/"),
    ("🎒", "bag", "/b/ /a/ /g/"), ("🌳", "tree", "t - ree"),
    ("🏞️", "park", "p - ar - k"), ("☀️", "sun", "/s/ /u/ /n/"),
    ("🏃", "run", "/r/ /u/ /n/"), ("🐘", "big", "/b/ /i/ /g/"),
    ("🔴", "red", "/r/ /e/ /d/"), ("😀", "happy", "hap - py"),
]

STORY = [
    ("Part 1", "🐶", ["Puppy goes to the park.", "Puppy has a red bag.",
                      "The bag is big.", "Puppy can run.",
                      "Puppy is happy."], GRASS, L_GRASS),
    ("Part 2", "🌳", ["Puppy sees a big tree.", "Puppy can run and play.",
                      "Puppy sees a cat.", "The cat can run too."], SKY, L_SKY),
    ("Part 3", "❓", ["Look! The bag is gone!", "Puppy can not see the bag.",
                      "Where is the red bag?", "Puppy looks and looks.",
                      "Puppy is sad."], PLUM, L_PLUM),
    ("Part 4", "🐱", ["Puppy looks at the tree.", "The bag is under the tree!",
                      "The cat sits on the bag.", "Look at that cat!"], CORAL, L_CORAL),
    ("Part 5", "😀", ["Puppy sees the red bag.", "Puppy and the cat play.",
                      "The sun is big and hot.", "It is a fun day!",
                      "Puppy is happy!"], AMBER, L_AMBER),
]

STORY_ROUNDS = [
    ("Round 1", "Teacher Model", "I read it. You just listen.", L_AMBER),
    ("Round 2", "Echo Reading", "I read one line. You say it back.", L_TEAL),
    ("Round 3", "Choral Reading", "We read it together.", L_SKY),
    ("Round 4", "Independent Try", "You pick an easy line and read it. ⭐", L_CORAL),
]

QUESTIONS = [
    ("Where did Puppy go?", [("🏫", "School"), ("🏞️", "Park"), ("🏠", "Home")], "Park"),
    ("What color was the bag?", [("🔴", "Red"), ("🔵", "Blue"), ("🟢", "Green")], "Red"),
    ("Where was the bag?", [("🌳", "Under the tree"), ("🏠", "In the house"),
                            ("🏫", "At school")], "Under the tree"),
    ("How did Puppy feel at the end?", [("😀", "Happy"), ("😠", "Angry"),
                                        ("😢", "Sad")], "Happy"),
]

FINAL_WORDS = ["CAT", "SUN", "DOG"]
FINAL_SENTENCES = ["I see a dog.", "The cat can run."]

LEARNED = [
    ("🔤", "Sounds"), ("🎵", "Blending"), ("🧩", "Words"), ("💎", "Sight words"),
    ("📖", "Sentences"), ("📚", "Story"), ("🕵️", "Comprehension"),
]

STRATEGIES = [
    ("1", "Reduce the words", "Show 2 words instead of 5."),
    ("2", "Use the picture", "Let the picture give the clue first."),
    ("3", "Sound it out", "Break it: /d/ /o/ /g/ — then blend."),
    ("4", "Echo reading", "You read the line, child repeats it."),
    ("5", "Read together", "Say it at the same time, same speed."),
    ("6", "Two choices", "Give 2 answer options, not 3."),
    ("7", "Repeat the word", "Bring the same word back in another game."),
    ("8", "Praise the effort", "\"Great try!\" even when it is not correct."),
]

RUBRIC_SKILLS = ["Letter sounds", "Blending", "CVC words", "Sight words",
                 "Sentence reading", "Story reading", "Comprehension", "Confidence"]
RUBRIC_NOTES = ["Words read independently", "Words needing support", "Next lesson focus"]

# ------------------------------------------------------------------ helpers


def set_run(run, size=20, bold=False, color=DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    tree = slide.shapes._spTree
    el = s._element
    tree.remove(el)
    tree.insert(2, el)


def add_rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_round(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_oval(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def tb(slide, l, t, w, h, text, size=20, bold=False, color=DARK,
       align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size, bold, color, font)
    return box


def bullets(slide, l, t, w, h, items, size=12, color=DARK, sp=6):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(sp)
        r = p.add_run()
        r.text = "•  " + item
        set_run(r, size, False, color)
    return box


def fade(slide):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", "med")
    tr.set("advClick", "1")
    etree.SubElement(tr, qn("p:fade"))


def footer(slide, n, timing="", step=""):
    add_rect(slide, Inches(0), Inches(7.02), prs.slide_width, Inches(0.08),
             RGBColor(0xE4, 0xEA, 0xF0))
    add_rect(slide, Inches(0), Inches(7.02), int(prs.slide_width * n / TOTAL), Inches(0.08), GOLD)
    add_rect(slide, Inches(0), Inches(7.10), prs.slide_width, Inches(0.40), INDIGO)
    msg = "🕵️ The Great Reading Adventure  |  Grade 2  |  90 min"
    if step:
        msg += f"  |  {step}"
    if timing:
        msg += f"  |  {timing}"
    tb(slide, Inches(0.35), Inches(7.14), Inches(11.3), Inches(0.32), msg, size=10, color=WHITE)
    tb(slide, Inches(11.85), Inches(7.14), Inches(1.2), Inches(0.32), f"{n} / {TOTAL}",
       size=10, color=WHITE, align=PP_ALIGN.RIGHT)


def chip(slide, text, color):
    add_round(slide, Inches(0.4), Inches(0.26), Inches(3.4), Inches(0.42), color)
    tb(slide, Inches(0.4), Inches(0.31), Inches(3.4), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def time_chip(slide, text):
    add_round(slide, Inches(10.3), Inches(0.26), Inches(2.65), Inches(0.42), CORAL)
    tb(slide, Inches(10.3), Inches(0.31), Inches(2.65), Inches(0.34), text, size=12, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER)


def new_slide(title, tag, timing, step, accent=SKY, bg=CREAM):
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_rect(slide, Inches(0), Inches(0), Inches(0.16), prs.slide_height, accent)
    chip(slide, tag, accent)
    if timing:
        time_chip(slide, timing)
    tb(slide, Inches(0.4), Inches(0.8), Inches(12.4), Inches(0.5), title, size=28, bold=True,
       color=INDIGO, font="Georgia")
    footer(slide, n, timing, step)
    fade(slide)
    return slide, n


def subtitle(slide, text, color=CORAL, top=1.36, size=16):
    return tb(slide, Inches(0.45), Inches(top), Inches(11.5), Inches(0.4), text,
              size=size, bold=True, color=color)


def ido_wedo_youdo(slide, top, accent, labels=("I read it", "We read it", "YOU read it!")):
    steps = [("1️⃣", labels[0], "Listen and follow", L_AMBER, INDIGO),
             ("2️⃣", labels[1], "Say it with me", L_TEAL, INDIGO),
             ("3️⃣", labels[2], "Your turn ⭐", accent, WHITE)]
    for i, (num, title, sub, bg, fg) in enumerate(steps):
        left = Inches(0.55 + i * 4.15)
        add_round(slide, left, top, Inches(3.9), Inches(1.5), bg)
        tb(slide, left + Inches(0.15), top + Inches(0.14), Inches(3.6), Inches(0.45), num,
           size=22, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(0.62), Inches(3.6), Inches(0.45), title,
           size=18, bold=True, color=fg, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.15), top + Inches(1.06), Inches(3.6), Inches(0.36), sub,
           size=13, color=fg, align=PP_ALIGN.CENTER)


def scene_panel(slide, left, top, width, height, items, caption):
    add_round(slide, left, top, width, height, L_SKY)
    add_round(slide, left, top + height - Inches(1.75), width, Inches(1.75), L_GRASS)
    for emoji, x, y in items:
        tb(slide, left + Inches(x), top + Inches(y), Inches(1.3), Inches(1.1), emoji,
           size=40, align=PP_ALIGN.CENTER)
    tb(slide, left + Inches(0.2), top + height - Inches(0.5), width - Inches(0.4),
       Inches(0.4), caption, size=14, bold=True, color=INDIGO)


def sound_rows(slide, items, start_index, accent, light):
    for i, (emoji, word, options) in enumerate(items):
        top = Inches(1.9 + i * 1.24)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.08), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.28), Inches(0.55), Inches(0.55), accent)
        tb(slide, Inches(0.72), top + Inches(0.35), Inches(0.55), Inches(0.42),
           str(start_index + i), size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.5), top + Inches(0.26), Inches(0.9), Inches(0.6), emoji, size=24,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.55), top + Inches(0.28), Inches(2.3), Inches(0.55), word, size=26,
           bold=True, color=INDIGO, font="Arial Black")
        for j, opt in enumerate(options):
            left = Inches(5.3 + j * 2.55)
            add_round(slide, left, top + Inches(0.22), Inches(2.3), Inches(0.64), light)
            tb(slide, left, top + Inches(0.28), Inches(2.3), Inches(0.5),
               f"{chr(65 + j)}.  {opt}", size=19, bold=True, color=INDIGO,
               align=PP_ALIGN.CENTER)


def blend_rows(slide, items, top_start=1.95, gap=1.5):
    for i, (word, letters, sounds, emoji, color, light) in enumerate(items):
        top = Inches(top_start + i * gap)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.32), WHITE)
        for j, (letter, sound) in enumerate(zip(letters, sounds)):
            left = Inches(0.85 + j * 1.85)
            add_round(slide, left, top + Inches(0.2), Inches(1.55), Inches(0.92), light)
            tb(slide, left, top + Inches(0.26), Inches(1.55), Inches(0.5), letter, size=26,
               bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
            tb(slide, left, top + Inches(0.75), Inches(1.55), Inches(0.32), sound, size=12,
               bold=True, color=SOFT, align=PP_ALIGN.CENTER)
            if j < len(letters) - 1:
                tb(slide, left + Inches(1.57), top + Inches(0.38), Inches(0.26), Inches(0.5),
                   "→", size=18, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, Inches(6.4), top + Inches(0.36), Inches(0.8), Inches(0.55), "➡", size=22,
           bold=True, color=CORAL, align=PP_ALIGN.CENTER)
        add_round(slide, Inches(7.4), top + Inches(0.2), Inches(2.9), Inches(0.92), L_AMBER)
        tb(slide, Inches(7.4), top + Inches(0.3), Inches(2.9), Inches(0.7), word, size=30,
           bold=True, color=INDIGO, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, Inches(10.6), top + Inches(0.3), Inches(1.0), Inches(0.65), emoji, size=24,
           align=PP_ALIGN.CENTER)
        add_round(slide, Inches(11.75), top + Inches(0.38), Inches(1.1), Inches(0.56), L_TEAL)
        tb(slide, Inches(11.75), top + Inches(0.46), Inches(1.1), Inches(0.4), "⭐", size=14,
           align=PP_ALIGN.CENTER)


def choice_rows(slide, items, accent, light, label_fmt="Find  {0}", numbered=True,
                top_start=1.85, gap=0.84, height=0.78):
    for i, (emoji, target, options) in enumerate(items):
        top = Inches(top_start + i * gap)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(height), WHITE)
        if numbered:
            add_oval(slide, Inches(0.7), top + Inches(0.16), Inches(0.46), Inches(0.46), accent)
            tb(slide, Inches(0.7), top + Inches(0.21), Inches(0.46), Inches(0.36), str(i + 1),
               size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.35), top + Inches(0.14), Inches(0.7), Inches(0.5), emoji, size=18,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.2), top + Inches(0.18), Inches(2.8), Inches(0.44),
           label_fmt.format(target), size=17, bold=True, color=accent)
        for j, opt in enumerate(options):
            left = Inches(5.3 + j * 2.55)
            add_round(slide, left, top + Inches(0.11), Inches(2.3), Inches(0.56), light)
            tb(slide, left, top + Inches(0.15), Inches(2.3), Inches(0.46), opt, size=19,
               bold=True, color=INDIGO, align=PP_ALIGN.CENTER, font="Arial Black")


# ------------------------------------------------------------------ slides


def s01_title():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, INDIGO)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.65, 0.8, SKY), (11.95, 0.85, CORAL), (0.85, 5.8, TEAL), (11.9, 5.75, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(1.15), Inches(12), Inches(1.05), "🕵️", size=60,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(2.2), Inches(12), Inches(0.85),
       "The Great Reading Adventure", size=42, bold=True, color=WHITE,
       align=PP_ALIGN.CENTER, font="Georgia")
    tb(slide, Inches(0.7), Inches(3.12), Inches(12), Inches(0.5),
       "Welcome, Reading Detective!  •  Grade 2  •  90 Minutes",
       size=18, color=L_SKY, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(3.3), Inches(3.9), Inches(6.7), Inches(1.3), TEAL)
    tb(slide, Inches(3.5), Inches(4.15), Inches(6.3), Inches(0.85),
       "🐶 Puppy lost his backpack.\nRead the clues and help him find it!",
       size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(5.55), Inches(12), Inches(0.5),
       "🔤 Sounds → 🎵 Blending → 🧩 Words → 💎 Sight Words → 📖 Sentences → 📚 Story",
       size=15, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    footer(slide, n, "0–7 min", "Welcome")
    fade(slide)


def s02_meet_puppy():
    slide, n = new_slide("🐶 Meet Puppy, Your Partner", "MEET", "0–7 min", "Meet Puppy", TEAL)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(5.6), Inches(4.95), L_TEAL)
    tb(slide, Inches(0.5), Inches(2.3), Inches(5.6), Inches(1.8), "🐶", size=96,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(4.3), Inches(5.6), Inches(0.7), "PUPPY", size=40, bold=True,
       color=TEAL, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(0.8), Inches(5.1), Inches(5.0), Inches(0.6),
       "Your reading partner today", size=16, color=SOFT, align=PP_ALIGN.CENTER)
    facts = [("🎒", "Puppy lost his red backpack."),
             ("🏞️", "He lost it somewhere in the park."),
             ("🔎", "Every clue is a word to read."),
             ("⭐", "You are the Reading Detective!")]
    for i, (icon, text) in enumerate(facts):
        top = Inches(1.6 + i * 1.22)
        add_round(slide, Inches(6.45), top, Inches(6.4), Inches(1.05), WHITE)
        add_oval(slide, Inches(6.7), top + Inches(0.24), Inches(0.58), Inches(0.58), L_AMBER)
        tb(slide, Inches(6.7), top + Inches(0.3), Inches(0.58), Inches(0.45), icon, size=16,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(7.5), top + Inches(0.3), Inches(5.1), Inches(0.5), text, size=17,
           bold=True, color=INDIGO)


def s03_mission():
    slide, n = new_slide("🗺️ Today's Mission — 7 Levels", "MISSION", "0–7 min", "Mission", AMBER)
    subtitle(slide, "Finish a level, unlock the next one. 🚀")
    for i, (num, label, icon, color) in enumerate(LEVELS):
        col, row = i % 4, i // 4
        left = Inches(0.5 + col * 3.16)
        top = Inches(1.9 + row * 2.45)
        add_round(slide, left, top, Inches(2.96), Inches(2.2), WHITE)
        add_oval(slide, left + Inches(1.13), top + Inches(0.2), Inches(0.7), Inches(0.7), color)
        tb(slide, left + Inches(1.13), top + Inches(0.31), Inches(0.7), Inches(0.5), icon,
           size=18, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.02), Inches(2.72), Inches(0.45),
           f"LEVEL {num}", size=13, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.45), Inches(2.72), Inches(0.55), label,
           size=19, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(10.0), Inches(4.35), Inches(2.85), Inches(2.2), L_AMBER)
    tb(slide, Inches(10.0), Inches(4.75), Inches(2.85), Inches(0.8), "🏆", size=40,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(10.1), Inches(5.6), Inches(2.65), Inches(0.6), "READING\nDETECTIVE BADGE",
       size=13, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)


def s04_warmup():
    slide, n = new_slide("💬 Warm-Up — Let's Talk First", "CONNECT", "0–7 min", "Warm-up", ROSE)
    subtitle(slide, "No reading yet. Just tell me about you!")
    for i, (icon, question) in enumerate(WARMUP):
        top = Inches(1.9 + i * 1.1)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.95), WHITE)
        add_oval(slide, Inches(0.78), top + Inches(0.2), Inches(0.55), Inches(0.55), L_ROSE)
        tb(slide, Inches(0.78), top + Inches(0.26), Inches(0.55), Inches(0.42), icon, size=16,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.6), top + Inches(0.22), Inches(10.8), Inches(0.52), question,
           size=22, bold=True, color=INDIGO)
    add_round(slide, Inches(0.5), Inches(6.28), Inches(12.35), Inches(0.6), L_AMBER)
    tb(slide, Inches(0.8), Inches(6.38), Inches(11.7), Inches(0.42),
       "⭐ Teacher answers first, then the student — it makes talking feel safe.",
       size=16, bold=True, color=INDIGO)


def s05_i_spy():
    slide, n = new_slide("👀 Mini Game: I Spy", "GAME", "0–7 min", "I spy", SKY)
    subtitle(slide, "I spy something... can you find it? Then read its word!")
    scene_panel(slide, Inches(0.5), Inches(1.8), Inches(6.3), Inches(4.55), SPY_SCENE,
                "🏞️ The Park")
    for i, (clue, emoji, word) in enumerate(SPY_ROUNDS):
        top = Inches(1.9 + i * 1.42)
        add_round(slide, Inches(7.05), top, Inches(5.8), Inches(1.22), WHITE)
        tb(slide, Inches(7.3), top + Inches(0.18), Inches(4.0), Inches(0.45), clue, size=18,
           bold=True, color=INDIGO)
        add_round(slide, Inches(7.3), top + Inches(0.66), Inches(2.1), Inches(0.45), L_SKY)
        tb(slide, Inches(7.3), top + Inches(0.71), Inches(2.1), Inches(0.36), word, size=17,
           bold=True, color=SKY, align=PP_ALIGN.CENTER, font="Arial Black")
        add_round(slide, Inches(11.6), top + Inches(0.3), Inches(1.0), Inches(0.62), L_GRASS)
        tb(slide, Inches(11.6), top + Inches(0.36), Inches(1.0), Inches(0.5), emoji, size=18,
           align=PP_ALIGN.CENTER)
    add_round(slide, Inches(7.05), Inches(6.15), Inches(5.8), Inches(0.72), L_AMBER)
    tb(slide, Inches(7.3), Inches(6.3), Inches(5.4), Inches(0.45),
       "⭐ Point first, then say the word out loud.", size=15, bold=True, color=INDIGO)


def s06_phonics():
    slide, n = new_slide("🔤 LEVEL 1 — Phonics Mission", "PHONICS", "7–17 min",
                         "Letter sounds", CORAL)
    subtitle(slide, "I say the sound → You say it back. Then I point → You tell me!")
    for i, (letter, sound, emoji, word, color, light) in enumerate(SOUNDS):
        col, row = i % 5, i // 5
        left = Inches(0.45 + col * 2.5)
        top = Inches(1.9 + row * 2.42)
        add_round(slide, left, top, Inches(2.32), Inches(2.22), light)
        tb(slide, left, top + Inches(0.06), Inches(2.32), Inches(0.72), letter, size=40,
           bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left, top + Inches(0.8), Inches(2.32), Inches(0.42), sound, size=19,
           bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.16), top + Inches(1.24), Inches(2.0), Inches(0.86), WHITE)
        tb(slide, left + Inches(0.16), top + Inches(1.3), Inches(2.0), Inches(0.42), emoji,
           size=18, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.16), top + Inches(1.71), Inches(2.0), Inches(0.36), word,
           size=14, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)


def s07_detective_a():
    slide, n = new_slide("🔎 Game: Sound Detective", "GAME", "7–17 min", "Sound detective", PLUM)
    subtitle(slide, "What sound does the word START with?")
    sound_rows(slide, DETECTIVE_A, 1, PLUM, L_PLUM)
    add_round(slide, Inches(0.5), Inches(5.65), Inches(12.35), Inches(1.15), L_AMBER)
    tb(slide, Inches(0.85), Inches(5.82), Inches(11.6), Inches(0.4),
       "🔎 Detective tip: stretch the first sound — \"ddddog\" — then ask again.",
       size=15, bold=True, color=INDIGO)
    tb(slide, Inches(0.85), Inches(6.24), Inches(11.6), Inches(0.4),
       "Still tricky? Cover one wrong answer so only two are left.",
       size=15, color=INDIGO)


def s08_detective_b():
    slide, n = new_slide("🔎 Sound Detective — Round 2", "GAME", "7–17 min",
                         "Sound detective", PLUM)
    subtitle(slide, "These clues are a little harder, Detective!")
    sound_rows(slide, DETECTIVE_B, 4, PLUM, L_PLUM)


def s09_blending():
    slide, n = new_slide("🎵 LEVEL 2 — Sounds Hold Hands", "BLENDING", "17–27 min",
                         "Blending", AMBER)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(12.35), Inches(2.5), WHITE)
    parts = [("C", "/k/", SKY), ("A", "/a/", CORAL), ("T", "/t/", GRASS)]
    for i, (letter, sound, color) in enumerate(parts):
        left = Inches(1.35 + i * 2.5)
        add_round(slide, left, Inches(1.85), Inches(1.9), Inches(1.8), L_SKY)
        tb(slide, left, Inches(1.98), Inches(1.9), Inches(0.95), letter, size=48, bold=True,
           color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left, Inches(2.95), Inches(1.9), Inches(0.5), sound, size=20, bold=True,
           color=INDIGO, align=PP_ALIGN.CENTER)
        if i < 2:
            tb(slide, left + Inches(1.93), Inches(2.3), Inches(0.55), Inches(0.65), "→",
               size=26, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
    tb(slide, Inches(8.85), Inches(2.3), Inches(0.8), Inches(0.65), "➡", size=26, bold=True,
       color=CORAL, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(9.7), Inches(1.85), Inches(2.75), Inches(1.8), L_AMBER)
    tb(slide, Inches(9.7), Inches(1.98), Inches(2.75), Inches(0.95), "CAT", size=42, bold=True,
       color=INDIGO, align=PP_ALIGN.CENTER, font="Arial Black")
    tb(slide, Inches(9.7), Inches(2.95), Inches(2.75), Inches(0.55), "🐱", size=24,
       align=PP_ALIGN.CENTER)
    ido_wedo_youdo(slide, Inches(4.3), AMBER, ("I blend it", "We blend it", "YOU blend it!"))
    tb(slide, Inches(0.55), Inches(6.05), Inches(9.5), Inches(0.42),
       "🐢 Slow:  /k/ ... /a/ ... /t/        🐇 Fast:  CAT!", size=17, bold=True, color=CORAL)


def s10_sound_train():
    slide, n = new_slide("🚂 Game: Sound Train", "GAME", "17–27 min", "Sound train", SKY)
    subtitle(slide, "Each car holds one sound. Push them together and read the word!")
    blend_rows(slide, TRAIN_WORDS, top_start=1.95, gap=1.62)


def s11_build_word():
    slide, n = new_slide("🔤 Game: Build the Word", "GAME", "17–27 min", "Build the word", TEAL)
    subtitle(slide, "Say each sound, then blend them fast.")
    blend_rows(slide, BUILD_WORDS, top_start=1.95, gap=1.62)


def s12_word_families():
    slide, n = new_slide("🧩 LEVEL 3 — CVC Words & Word Families", "WORDS", "27–37 min",
                         "Word families", GRASS)
    subtitle(slide, "Change the first sound and you get a brand new word!")
    for i, (family, words, emoji, color, light) in enumerate(FAMILIES):
        left = Inches(0.5 + i * 3.16)
        add_round(slide, left, Inches(1.9), Inches(2.96), Inches(4.45), light)
        tb(slide, left, Inches(2.12), Inches(2.96), Inches(0.7), emoji, size=28,
           align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(0.78), Inches(2.88), Inches(1.4), Inches(0.58), color)
        tb(slide, left + Inches(0.78), Inches(2.98), Inches(1.4), Inches(0.42), family,
           size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, word in enumerate(words):
            top = Inches(3.62 + j * 0.68)
            add_round(slide, left + Inches(0.35), top, Inches(2.26), Inches(0.56), WHITE)
            tb(slide, left + Inches(0.35), top + Inches(0.06), Inches(2.26), Inches(0.44),
               word, size=22, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)


def s13_monster():
    slide, n = new_slide("👾 Game: Feed the Reading Monster", "GAME", "27–37 min",
                         "Feed the monster", CORAL)
    subtitle(slide, "The monster only eats the RIGHT word. Find it and read it!")
    choice_rows(slide, MONSTER, CORAL, L_CORAL, "Monster wants  {0}")


def s14_word_hunt():
    slide, n = new_slide("🐾 Game: Word Hunt", "GAME", "27–37 min", "Word hunt", PLUM)
    subtitle(slide, "Point to the word I say — then read it out loud!")
    for r, row in enumerate(HUNT_GRID):
        for c, word in enumerate(row):
            left = Inches(0.5 + c * 1.92)
            top = Inches(1.9 + r * 1.32)
            add_round(slide, left, top, Inches(1.78), Inches(1.15), WHITE)
            tb(slide, left, top + Inches(0.3), Inches(1.78), Inches(0.6), word, size=22,
               bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(10.25), Inches(1.9), Inches(2.6), Inches(3.79), L_PLUM)
    tb(slide, Inches(10.45), Inches(2.06), Inches(2.2), Inches(0.42), "Find these:", size=16,
       bold=True, color=INDIGO)
    for i, word in enumerate(HUNT_TARGETS):
        tb(slide, Inches(10.55), Inches(2.58 + i * 0.48), Inches(2.1), Inches(0.42),
           f"{i + 1}.  {word}", size=16, bold=True, color=DARK)
    add_round(slide, Inches(0.5), Inches(5.95), Inches(9.35), Inches(0.6), L_AMBER)
    tb(slide, Inches(0.75), Inches(6.05), Inches(8.9), Inches(0.42),
       "⭐ Every word is hidden TWICE. Can you find both?", size=15, bold=True, color=INDIGO)


def s15_brain_break():
    slide, n = new_slide("🎮 Brain Break: Read & Move", "BREAK", "37–42 min",
                         "Read & move", GOLD)
    subtitle(slide, "Read the word first — then do the move!")
    for i, (word, action, emoji) in enumerate(READ_AND_MOVE):
        col, row = i % 2, i // 2
        left = Inches(0.5 + col * 6.2)
        top = Inches(1.95 + row * 1.55)
        add_round(slide, left, top, Inches(5.95), Inches(1.35), WHITE)
        add_round(slide, left + Inches(0.25), top + Inches(0.3), Inches(1.85), Inches(0.75),
                  L_AMBER)
        tb(slide, left + Inches(0.25), top + Inches(0.4), Inches(1.85), Inches(0.55), word,
           size=24, bold=True, color=INDIGO, align=PP_ALIGN.CENTER, font="Arial Black")
        tb(slide, left + Inches(2.3), top + Inches(0.42), Inches(0.6), Inches(0.5), "➡",
           size=18, bold=True, color=SOFT, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(3.0), top + Inches(0.38), Inches(0.6), Inches(0.55), emoji,
           size=20, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(3.65), top + Inches(0.45), Inches(2.1), Inches(0.5), action,
           size=14, bold=True, color=INDIGO)


def s16_sight_words():
    slide, n = new_slide("💎 LEVEL 4 — Sight Words", "SIGHT WORDS", "42–52 min",
                         "Sight words", PLUM)
    subtitle(slide, "We do NOT sound these out. We just know them by sight!")
    top = Inches(1.9)
    for label, words, color, light in SIGHT_SETS:
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.5), light)
        add_round(slide, Inches(0.72), top + Inches(0.45), Inches(1.6), Inches(0.55), color)
        tb(slide, Inches(0.72), top + Inches(0.54), Inches(1.6), Inches(0.4), label, size=13,
           bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, word in enumerate(words):
            left = Inches(2.7 + j * 2.55)
            add_round(slide, left, top + Inches(0.3), Inches(2.3), Inches(0.88), WHITE)
            tb(slide, left, top + Inches(0.45), Inches(2.3), Inches(0.58), word, size=30,
               bold=True, color=color, align=PP_ALIGN.CENTER, font="Arial Black")
        top = top + Inches(1.62)
    tb(slide, Inches(0.5), Inches(6.78), Inches(11.0), Inches(0.3), "", size=8)


def s17_treasure():
    slide, n = new_slide("💎 Game: Sight Word Treasure Hunt", "GAME", "42–52 min",
                         "Treasure hunt", TEAL)
    subtitle(slide, "Find the word to unlock the next spot on the map!")
    for i, (target, options) in enumerate(TREASURE):
        top = Inches(1.85 + i * 0.92)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.8), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.17), Inches(0.46), Inches(0.46), TEAL)
        tb(slide, Inches(0.72), top + Inches(0.22), Inches(0.46), Inches(0.36), str(i + 1),
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.4), top + Inches(0.15), Inches(0.7), Inches(0.5), "🗺️", size=18,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.25), top + Inches(0.19), Inches(2.7), Inches(0.44),
           f"Find  {target}", size=19, bold=True, color=TEAL)
        for j, opt in enumerate(options):
            left = Inches(5.3 + j * 2.55)
            add_round(slide, left, top + Inches(0.12), Inches(2.3), Inches(0.56), L_TEAL)
            tb(slide, left, top + Inches(0.16), Inches(2.3), Inches(0.46), opt, size=20,
               bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.45), Inches(12.35), Inches(0.5), L_AMBER)
    tb(slide, Inches(0.8), Inches(6.54), Inches(11.7), Inches(0.36),
       "🏴 Five stops found — the treasure is the next game!", size=15, bold=True, color=INDIGO)


def s18_roll_read():
    slide, n = new_slide("🎲 Game: Roll & Read", "GAME", "42–52 min", "Roll & read", ROSE)
    subtitle(slide, "Pick a number, read that word. Read all six to win!")
    for i, (num, word) in enumerate(ROLL_READ):
        col, row = i % 3, i // 3
        left = Inches(0.5 + col * 4.15)
        top = Inches(1.95 + row * 2.35)
        add_round(slide, left, top, Inches(3.9), Inches(2.1), WHITE)
        add_oval(slide, left + Inches(0.25), top + Inches(0.25), Inches(0.7), Inches(0.7), ROSE)
        tb(slide, left + Inches(0.25), top + Inches(0.34), Inches(0.7), Inches(0.5), num,
           size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_round(slide, left + Inches(1.15), top + Inches(0.55), Inches(2.5), Inches(1.15),
                  L_ROSE)
        tb(slide, left + Inches(1.15), top + Inches(0.75), Inches(2.5), Inches(0.75), word,
           size=34, bold=True, color=ROSE, align=PP_ALIGN.CENTER, font="Arial Black")


def s19_sentence_intro():
    slide, n = new_slide("📖 LEVEL 5 — Now We Read Sentences", "READING", "52–62 min",
                         "Sentences", SKY)
    sentence, emoji, color, light = SENTENCES[0]
    add_round(slide, Inches(0.5), Inches(1.5), Inches(12.35), Inches(2.55), WHITE)
    tb(slide, Inches(0.85), Inches(1.95), Inches(1.5), Inches(1.7), emoji, size=54,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(2.6), Inches(2.15), Inches(9.9), Inches(1.3), sentence, size=54,
       bold=True, color=INDIGO, align=PP_ALIGN.CENTER, font="Arial Black")
    ido_wedo_youdo(slide, Inches(4.35), SKY)
    tb(slide, Inches(0.55), Inches(6.1), Inches(9.2), Inches(0.42),
       "👉 Point under each word as you read it.", size=17, bold=True, color=CORAL)


def s20_build_sentence():
    slide, n = new_slide("🧩 Game: Build the Sentence", "GAME", "52–62 min", "Build it", PLUM)
    subtitle(slide, "The words fell out of order! Put them back, then read it.")
    for i, (tiles, _answer, emoji) in enumerate(SCRAMBLES):
        top = Inches(1.85 + i * 0.92)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(0.8), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.17), Inches(0.46), Inches(0.46), PLUM)
        tb(slide, Inches(0.72), top + Inches(0.22), Inches(0.46), Inches(0.36), str(i + 1),
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.4), top + Inches(0.15), Inches(0.7), Inches(0.5), emoji, size=18,
           align=PP_ALIGN.CENTER)
        for j, tile in enumerate(tiles):
            left = Inches(2.3 + j * 2.55)
            add_round(slide, left, top + Inches(0.12), Inches(2.3), Inches(0.56), L_PLUM)
            tb(slide, left, top + Inches(0.16), Inches(2.3), Inches(0.46), tile, size=19,
               bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.45), Inches(12.35), Inches(0.5), L_AMBER)
    tb(slide, Inches(0.8), Inches(6.54), Inches(11.7), Inches(0.36),
       "🔎 Clue: the first word has a CAPITAL letter. The last word has the period.",
       size=15, bold=True, color=INDIGO)


def s21_sentence_reading():
    slide, n = new_slide("⭐ Sentence Reading — Your Turn", "READING", "52–62 min",
                         "Sentences", TEAL)
    subtitle(slide, "I read → We read → YOU read. Take your time!")
    for i, (sentence, emoji, color, light) in enumerate(SENTENCES[1:5]):
        top = Inches(1.9 + i * 1.25)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.1), light)
        tb(slide, Inches(0.8), top + Inches(0.24), Inches(1.0), Inches(0.65), emoji, size=26,
           align=PP_ALIGN.CENTER)
        tb(slide, Inches(2.0), top + Inches(0.22), Inches(8.3), Inches(0.68), sentence,
           size=32, bold=True, color=INDIGO, font="Arial Black")
        add_round(slide, Inches(10.5), top + Inches(0.28), Inches(2.1), Inches(0.55), color)
        tb(slide, Inches(10.5), top + Inches(0.37), Inches(2.1), Inches(0.4), "YOU read ⭐",
           size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(0.5), Inches(6.95), Inches(12.35), Inches(0.02), CREAM)


def s22_picture_scene():
    slide, n = new_slide("🏞️ Picture Scene — What Do You See?", "READING", "62–70 min",
                         "Picture scene", GRASS)
    subtitle(slide, "Point to each thing and say its name out loud.")
    scene_panel(slide, Inches(0.5), Inches(1.85), Inches(7.6), Inches(4.6), PARK_SCENE,
                "🏞️ Puppy's Park")
    add_round(slide, Inches(8.35), Inches(1.85), Inches(4.5), Inches(4.6), L_GRASS)
    tb(slide, Inches(8.6), Inches(2.05), Inches(4.0), Inches(0.45), "I can see:", size=18,
       bold=True, color=INDIGO)
    for i, label in enumerate(PARK_LABELS):
        top = Inches(2.6 + i * 0.6)
        add_round(slide, Inches(8.6), top, Inches(4.0), Inches(0.5), WHITE)
        tb(slide, Inches(8.85), top + Inches(0.06), Inches(3.6), Inches(0.4), label, size=17,
           bold=True, color=INDIGO)


def s23_read_the_clue():
    slide, n = new_slide("🕵️ Game: Read the Clue", "GAME", "62–70 min", "Read the clue", ROSE)
    subtitle(slide, "Read the clue, then point to the right picture!")
    for i, (clue, options, _answer) in enumerate(CLUES):
        top = Inches(1.9 + i * 1.28)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(1.12), WHITE)
        add_oval(slide, Inches(0.72), top + Inches(0.3), Inches(0.52), Inches(0.52), ROSE)
        tb(slide, Inches(0.72), top + Inches(0.36), Inches(0.52), Inches(0.4), str(i + 1),
           size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.5), top + Inches(0.3), Inches(5.4), Inches(0.55), clue, size=24,
           bold=True, color=INDIGO)
        for j, emoji in enumerate(options):
            left = Inches(7.4 + j * 1.85)
            add_round(slide, left, top + Inches(0.2), Inches(1.6), Inches(0.72), L_ROSE)
            tb(slide, left, top + Inches(0.26), Inches(1.6), Inches(0.58), emoji, size=24,
               align=PP_ALIGN.CENTER)


def s24_story_words():
    slide, n = new_slide("📚 Story Words — Get Ready", "VOCAB", "70–80 min",
                         "Story words", AMBER)
    subtitle(slide, "Picture → Word → Sounds. These words are all in Puppy's story!")
    for i, (emoji, word, sounds) in enumerate(VOCAB):
        col, row = i % 5, i // 5
        left = Inches(0.45 + col * 2.5)
        top = Inches(1.9 + row * 2.42)
        add_round(slide, left, top, Inches(2.32), Inches(2.22), L_AMBER)
        tb(slide, left, top + Inches(0.14), Inches(2.32), Inches(0.7), emoji, size=30,
           align=PP_ALIGN.CENTER)
        tb(slide, left, top + Inches(0.92), Inches(2.32), Inches(0.62), word, size=26,
           bold=True, color=AMBER, align=PP_ALIGN.CENTER, font="Arial Black")
        add_round(slide, left + Inches(0.16), top + Inches(1.58), Inches(2.0), Inches(0.5), WHITE)
        tb(slide, left + Inches(0.16), top + Inches(1.65), Inches(2.0), Inches(0.38), sounds,
           size=12, bold=True, color=SOFT, align=PP_ALIGN.CENTER)


def s25_story_method():
    slide, n = new_slide("📖 How We Read Puppy's Story", "STORY", "70–80 min",
                         "Reading method", SKY)
    subtitle(slide, "Four rounds. You never have to read it alone the first time.")
    for i, (round_name, title, detail, light) in enumerate(STORY_ROUNDS):
        col, row = i % 2, i // 2
        left = Inches(0.5 + col * 6.25)
        top = Inches(1.95 + row * 2.4)
        add_round(slide, left, top, Inches(5.95), Inches(2.2), light)
        add_round(slide, left + Inches(0.3), top + Inches(0.28), Inches(1.55), Inches(0.5),
                  INDIGO)
        tb(slide, left + Inches(0.3), top + Inches(0.36), Inches(1.55), Inches(0.38),
           round_name, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.3), top + Inches(0.95), Inches(5.3), Inches(0.55), title,
           size=24, bold=True, color=INDIGO)
        tb(slide, left + Inches(0.3), top + Inches(1.55), Inches(5.3), Inches(0.5), detail,
           size=15, color=DARK)


def _story_slide(index, timing):
    part, emoji, lines, color, light = STORY[index]
    slide, n = new_slide(f"📚 Puppy's Lost Backpack — {part}", "STORY", timing, part, color)
    add_round(slide, Inches(0.5), Inches(1.5), Inches(8.9), Inches(4.95), WHITE)
    gap, first, size = (1.12, 1.95, 32) if len(lines) <= 4 else (0.92, 1.78, 28)
    for i, line in enumerate(lines):
        tb(slide, Inches(0.95), Inches(first + i * gap), Inches(8.1), Inches(0.85), line,
           size=size, bold=True, color=INDIGO)
    add_round(slide, Inches(9.65), Inches(1.5), Inches(3.2), Inches(4.95), light)
    tb(slide, Inches(9.65), Inches(2.1), Inches(3.2), Inches(1.4), emoji, size=64,
       align=PP_ALIGN.CENTER)
    steps = ["1️⃣  I read it", "2️⃣  You echo it", "3️⃣  We read together",
             "4️⃣  You pick ONE line"]
    for i, step in enumerate(steps):
        top = Inches(3.75 + i * 0.66)
        add_round(slide, Inches(9.9), top, Inches(2.7), Inches(0.54), WHITE)
        tb(slide, Inches(9.9), top + Inches(0.09), Inches(2.7), Inches(0.4), step, size=13,
           bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    return slide, n


def s26_story_1():
    _story_slide(0, "70–80 min")


def s27_story_2():
    _story_slide(1, "70–80 min")


def s28_story_3():
    _story_slide(2, "70–80 min")


def s29_story_4():
    _story_slide(3, "70–80 min")


def s30_story_5():
    _story_slide(4, "70–80 min")


def _question_slide(items, title, timing, start_index, accent, light):
    slide, n = new_slide(title, "DETECTIVE", timing, "Comprehension", accent)
    for i, (question, options, _answer) in enumerate(items):
        top = Inches(1.6 + i * 2.65)
        add_round(slide, Inches(0.5), top, Inches(12.35), Inches(2.4), WHITE)
        add_oval(slide, Inches(0.75), top + Inches(0.9), Inches(0.65), Inches(0.65), accent)
        tb(slide, Inches(0.75), top + Inches(1.0), Inches(0.65), Inches(0.48),
           str(start_index + i), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, Inches(1.65), top + Inches(0.92), Inches(4.0), Inches(0.7), question,
           size=20, bold=True, color=INDIGO)
        for j, (emoji, label) in enumerate(options):
            left = Inches(5.9 + j * 2.35)
            add_round(slide, left, top + Inches(0.35), Inches(2.2), Inches(1.7), light)
            tb(slide, left, top + Inches(0.55), Inches(2.2), Inches(0.7), emoji, size=30,
               align=PP_ALIGN.CENTER)
            tb(slide, left + Inches(0.1), top + Inches(1.35), Inches(2.0), Inches(0.5), label,
               size=15, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    return slide, n


def s31_story_detective():
    _question_slide(QUESTIONS[:2], "🕵️ LEVEL 7 — Story Detective", "80–86 min", 1,
                    PLUM, L_PLUM)


def s32_comprehension():
    _question_slide(QUESTIONS[2:], "🧩 Comprehension Game", "80–86 min", 3, TEAL, L_TEAL)


def s33_final_mission():
    slide, n = new_slide("⭐ Reading Detective — Final Mission", "CHALLENGE", "86–88 min",
                         "Final mission", CORAL)
    subtitle(slide, "Read as many as you can. Every single one counts! 🌟")
    add_round(slide, Inches(0.5), Inches(1.85), Inches(12.35), Inches(2.1), L_AMBER)
    tb(slide, Inches(0.85), Inches(2.02), Inches(3.0), Inches(0.4), "3 WORDS", size=15,
       bold=True, color=INDIGO)
    for i, word in enumerate(FINAL_WORDS):
        left = Inches(0.85 + i * 4.0)
        add_round(slide, left, Inches(2.55), Inches(3.7), Inches(1.15), WHITE)
        tb(slide, left, Inches(2.78), Inches(3.7), Inches(0.75), word, size=36, bold=True,
           color=CORAL, align=PP_ALIGN.CENTER, font="Arial Black")
    add_round(slide, Inches(0.5), Inches(4.15), Inches(12.35), Inches(2.1), L_TEAL)
    tb(slide, Inches(0.85), Inches(4.32), Inches(4.0), Inches(0.4), "2 SENTENCES", size=15,
       bold=True, color=INDIGO)
    for i, sentence in enumerate(FINAL_SENTENCES):
        top = Inches(4.85 + i * 0.68)
        add_round(slide, Inches(0.85), top, Inches(11.65), Inches(0.58), WHITE)
        tb(slide, Inches(1.15), top + Inches(0.06), Inches(11.1), Inches(0.45), sentence,
           size=24, bold=True, color=INDIGO)
    add_round(slide, Inches(0.5), Inches(6.4), Inches(12.35), Inches(0.5), GOLD)
    tb(slide, Inches(0.8), Inches(6.48), Inches(11.7), Inches(0.36),
       "⭐ Not a test — a celebration. Count every success out loud!", size=14, bold=True,
       color=INDIGO)


def s34_learned():
    slide, n = new_slide("✅ What I Learned Today", "RECAP", "88–90 min", "Recap", GRASS)
    subtitle(slide, "Look at everything you can do now!")
    for i, (icon, text) in enumerate(LEARNED):
        col, row = i % 4, i // 4
        left = Inches(0.5 + col * 3.16)
        top = Inches(1.9 + row * 2.45)
        add_round(slide, left, top, Inches(2.96), Inches(2.2), WHITE)
        add_oval(slide, left + Inches(1.13), top + Inches(0.22), Inches(0.7), Inches(0.7),
                 L_GRASS)
        tb(slide, left + Inches(1.13), top + Inches(0.33), Inches(0.7), Inches(0.5), icon,
           size=18, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.1), Inches(2.72), Inches(0.5), "☑",
           size=17, bold=True, color=GRASS, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.12), top + Inches(1.58), Inches(2.72), Inches(0.5), text,
           size=17, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    add_round(slide, Inches(10.0), Inches(4.35), Inches(2.85), Inches(2.2), L_AMBER)
    tb(slide, Inches(10.15), Inches(4.6), Inches(2.55), Inches(1.7),
       "🗣️ What was your favorite game today?", size=16, bold=True, color=INDIGO,
       align=PP_ALIGN.CENTER)


def s35_badge():
    _counter["n"] += 1
    n = _counter["n"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, INDIGO)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.26), GOLD)
    for x, y, c in [(0.7, 0.85, SKY), (11.9, 0.85, CORAL), (0.9, 5.75, TEAL), (11.85, 5.7, AMBER)]:
        add_oval(slide, Inches(x), Inches(y), Inches(0.8), Inches(0.8), c)
    tb(slide, Inches(0.7), Inches(0.95), Inches(12), Inches(0.8),
       "🏅 I AM A READING DETECTIVE! 🏅", size=36, bold=True, color=GOLD,
       align=PP_ALIGN.CENTER, font="Georgia")
    add_oval(slide, Inches(5.42), Inches(1.9), Inches(2.5), Inches(2.5), GOLD)
    tb(slide, Inches(5.42), Inches(2.4), Inches(2.5), Inches(1.4), "🐶", size=58,
       align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(4.65), Inches(12), Inches(0.6), "You worked hard today!",
       size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Georgia")
    tb(slide, Inches(0.7), Inches(5.32), Inches(12), Inches(0.55),
       "Every word you read makes you stronger!", size=21, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.7), Inches(5.95), Inches(12), Inches(0.5),
       "You helped Puppy find his backpack. Mission complete! 🎒",
       size=16, color=L_SKY, align=PP_ALIGN.CENTER)
    footer(slide, n, "88–90 min", "Badge")
    fade(slide)


def s36_struggling():
    slide, n = new_slide("🧰 If the Student Is Struggling", "TEACHER ONLY", "",
                         "Support", INDIGO)
    tb(slide, Inches(0.45), Inches(1.34), Inches(12.4), Inches(0.35),
       "TEACHER ONLY — eight quick moves. Use them instead of giving the answer.",
       size=13, bold=True, color=CORAL)
    for i, (num, title, detail) in enumerate(STRATEGIES):
        col, row = i % 4, i // 4
        left = Inches(0.45 + col * 3.18)
        top = Inches(1.85 + row * 2.5)
        add_round(slide, left, top, Inches(3.0), Inches(2.3), WHITE)
        add_oval(slide, left + Inches(0.2), top + Inches(0.2), Inches(0.56), Inches(0.56), TEAL)
        tb(slide, left + Inches(0.2), top + Inches(0.26), Inches(0.56), Inches(0.42), num,
           size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        tb(slide, left + Inches(0.2), top + Inches(0.88), Inches(2.6), Inches(0.5), title,
           size=17, bold=True, color=INDIGO)
        tb(slide, left + Inches(0.2), top + Inches(1.42), Inches(2.6), Inches(0.75), detail,
           size=13, color=DARK)


def s37_assessment_key():
    slide, n = new_slide("🔒 Teacher Assessment & Answer Key", "TEACHER ONLY", "",
                         "Assessment", INDIGO)
    tb(slide, Inches(0.45), Inches(1.34), Inches(12.4), Inches(0.35),
       "TEACHER ONLY — ANSWER KEY. Hide this slide before presenting.",
       size=13, bold=True, color=CORAL)
    header_y = Inches(1.78)
    cols = [(0.45, 3.0, "SKILL"), (3.6, 1.5, "Needs"), (5.25, 1.5, "Dev."), (6.9, 1.4, "Strong")]
    for left, width, label in cols:
        add_round(slide, Inches(left), header_y, Inches(width), Inches(0.44), INDIGO)
        tb(slide, Inches(left), header_y + Inches(0.07), Inches(width), Inches(0.32), label,
           size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, skill in enumerate(RUBRIC_SKILLS):
        top = header_y + Inches(0.5 + i * 0.42)
        band = WHITE if i % 2 == 0 else L_SKY
        add_round(slide, Inches(0.45), top, Inches(3.0), Inches(0.37), band)
        tb(slide, Inches(0.6), top + Inches(0.04), Inches(2.8), Inches(0.29), skill, size=12,
           bold=True, color=INDIGO)
        for left, width, _l in cols[1:]:
            add_round(slide, Inches(left), top, Inches(width), Inches(0.37), band)
            tb(slide, Inches(left), top + Inches(0.02), Inches(width), Inches(0.32), "☐",
               size=14, color=SOFT, align=PP_ALIGN.CENTER)
    for i, label in enumerate(RUBRIC_NOTES):
        top = Inches(5.72 + i * 0.42)
        add_round(slide, Inches(0.45), top, Inches(7.85), Inches(0.36), L_AMBER)
        tb(slide, Inches(0.6), top + Inches(0.04), Inches(7.6), Inches(0.28), f"{label}:",
           size=12, bold=True, color=INDIGO)
    add_round(slide, Inches(8.5), Inches(1.78), Inches(4.35), Inches(5.05), WHITE)
    tb(slide, Inches(8.75), Inches(1.94), Inches(3.9), Inches(0.38), "🔑 ANSWER KEY", size=14,
       bold=True, color=CORAL)
    bullets(slide, Inches(8.75), Inches(2.38), Inches(3.9), Inches(4.3), [
        "Sound Detective: d, k, s, b, f, t, m",
        "Sound Train: CAT, MAT, SUN",
        "Build the Word: PIG, DOG, RUN",
        "Families: -at -og -un -ig",
        "Monster: CAT DOG SUN PIG BAG RUN",
        "Word Hunt: each word appears 2x",
        "Treasure: the, can, I, my, see",
        "Sentences: I see a dog. / The dog can",
        "     run. / The cat is big. / I can run. /",
        "     The sun is hot.",
        "Clues: dog, ball, boy, sun",
        "Story Q: Park / Red / Under the tree /",
        "     Happy",
        "Final: CAT SUN DOG + 2 sentences",
    ], size=11, sp=3)


BUILDERS = [
    s01_title, s02_meet_puppy, s03_mission, s04_warmup, s05_i_spy, s06_phonics,
    s07_detective_a, s08_detective_b, s09_blending, s10_sound_train, s11_build_word,
    s12_word_families, s13_monster, s14_word_hunt, s15_brain_break, s16_sight_words,
    s17_treasure, s18_roll_read, s19_sentence_intro, s20_build_sentence,
    s21_sentence_reading, s22_picture_scene, s23_read_the_clue, s24_story_words,
    s25_story_method, s26_story_1, s27_story_2, s28_story_3, s29_story_4, s30_story_5,
    s31_story_detective, s32_comprehension, s33_final_mission, s34_learned, s35_badge,
    s36_struggling, s37_assessment_key,
]

for build in BUILDERS:
    build()

assert len(prs.slides) == TOTAL, f"expected {TOTAL} slides, built {len(prs.slides)}"

# This deck ships without speaker notes.
for slide in prs.slides:
    if slide.has_notes_slide:
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype.endswith("/notesSlide"):
                slide.part.drop_rel(rid)

out = r"C:\Users\bhushaja\Downloads\shaip\Grade2_Great_Reading_Adventure_90min.pptx"
try:
    prs.save(out)
except PermissionError:
    out = out.replace(".pptx", "_new.pptx")
    prs.save(out)

story_words = sum(len(line.split()) for _p, _e, lines, _c, _l in STORY for line in lines)
with_notes = [i + 1 for i, s in enumerate(prs.slides) if s.has_notes_slide]
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
print(f"Story word count: {story_words}")
print(f"Slides carrying notes: {with_notes if with_notes else 'none'}")
